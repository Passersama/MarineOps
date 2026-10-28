import copy
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from cargo_track.domain import Command, AppError
from cargo_track.service import ApplicationService
from cargo_track.personnel_cargo import read_cargo
from cargo_track.personnel_import import read_personnel_excel, PersonnelError
from cargo_track.personnel_service import PersonnelService, clean_payload
from cargo_track.text_case import normalize, personnel_record, legacy_personnel
from cargo_track.berth_plan import generate_plan, export_plan


class TextCaseRules(unittest.TestCase):
    def test_fields_units_models_punctuation_and_stability(self):
        cases = [('name', '  jORDAN   lEE ', 'Jordan Lee'),
                 ('name', "ANNE-MARIE O’NEILL", 'Anne-Marie O’Neill'),
                 ('name', 'MCDONALD J. P.', 'McDonald J. P.'),
                 ('position', 'nav qc / hse OFFICER', 'NAV QC / HSE Officer'),
                 ('cargo', 'PVC PIPE DN50 12V 10 KG 2 mA\r\n GPS ANTENNA', 'PVC Pipe DN50 12V 10 kg 2 mA\nGPS Antenna'),
                 ('cargo', 'eBay iPhone X-500 USB-C 中文', 'eBay iPhone X-500 USB-C 中文'),
                 ('cargo', 'GPS电缆 220V-24v电源转换器 sAfEtY heLMET', 'GPS电缆 220V-24v电源转换器 Safety Helmet'),
                 ('cargo', 'fuel/oil PIPE-HOSE XYZ UNIT', 'Fuel/Oil Pipe-Hose XYZ UNIT'),
                 ('cargo', '4 mA 4 MA 5 mW 5 MW', '4 mA 4 MA 5 mW 5 MW'),
                 ('cargo', '=SUM(A1:A2)', '=SUM(A1:A2)')]
        for kind, raw, expected in cases:
            with self.subTest(kind=kind, raw=raw):
                result = normalize(raw, kind)[0]
                self.assertEqual(expected, result)
                self.assertEqual(result, normalize(result, kind)[0])

    def test_preserved_personal_spelling_and_legacy_correction(self):
        data = personnel_record({'name':'de la CRUZ','position':'qC'}, preserve=True)
        self.assertEqual('de la CRUZ', personnel_record(data)['name'])
        old = {'name':'Alex Morgan','position':'Nav Qc','original_values':{'name':'ALEX MORGAN','position':'NAV QC'}}
        self.assertEqual('NAV QC', legacy_personnel(old)['position'])
        old['name']='Special McArthur'
        before=copy.deepcopy(old)
        self.assertEqual('Special McArthur',legacy_personnel(old)['name'])
        self.assertEqual(before,old)


class TextCaseFlow(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.cargo=ApplicationService(self.root/'cargo.sqlite3');self.cargo.open();self.counter=0
        self.vessel=self.command('create_vessel',{'name':'CHA'})['id']
        self.command('create_voyage',{'name':'Case fixture'})
        self.command('add_stop',{'vessel_id':self.vessel})
        self.snapshot=read_cargo(self.cargo.path)
        self.shore=next(x['id'] for x in self.snapshot['locations'] if x['kind']=='SHORE')
        self.service=PersonnelService(self.root/'personnel.sqlite3')
        self.plan=self.service.create_plan('Case crew','2026-10-26')

    def tearDown(self):
        self.service.close();self.cargo.close();self.temp.cleanup()

    def command(self,operation,payload):
        self.counter+=1
        return self.cargo.handle(Command(f'case-{self.counter}',operation,payload)).result

    def create_cargo(self,description,**kwargs):
        return self.command('create_transfer',{'description':description,'from_location_id':self.shore,'to_location_id':self.vessel,**kwargs})

    def record(self,**kwargs):
        return self.service.bind_record({'name':'jORDAN lEE','position':'nav qc',
            'date':'2026-10-26','from_name':'xx port','to_name':'CHA',**kwargs},self.snapshot)

    def legacy(self):
        identifier=self.service.add_passenger(self.plan,self.record())
        data=clean_payload(self.service.passenger(identifier));data.pop('text_case',None)
        data.update(name='JORDAN LEE',position='Nav Qc',original_values={'name':'JORDAN LEE','position':'NAV QC'})
        self.service.db.execute('UPDATE passengers SET payload=? WHERE id=?',(json.dumps(data),identifier))
        return identifier

    def test_import_preview_originals_preserve_duplicates_and_export(self):
        path=self.root/'crew.xlsx';book=Workbook();sheet=book.active
        sheet.append(['Name','Position','From','To','Date'])
        sheet.append(['JORDAN LEE','NAV QC','xx port','CHA','2026-10-26']);book.save(path);book.close()
        original=path.read_bytes();parsed=read_personnel_excel(path)
        self.assertEqual('NAV QC',parsed['records'][0]['position'])
        self.assertEqual('JORDAN LEE',parsed['records'][0]['original_values']['name'])
        preview=self.service.import_preview(parsed,self.snapshot,self.plan)
        self.service.commit_import(preview)
        with self.assertRaises(PersonnelError):self.service.import_preview(parsed,self.snapshot,self.plan)
        crew=self.service.crew_records(self.plan,self.snapshot)
        self.assertEqual('Jordan Lee',crew[0]['name'])
        self.assertEqual('NAV QC',crew[0]['position'])
        saved=self.service.save_berth(generate_plan(self.snapshot,crew,[],'2026-10-26','xx port',mode='Final'))
        export=self.root/'plan.xlsx';export_plan(saved,export)
        workbook=load_workbook(export,read_only=True)
        cells=' '.join(str(c.value) for s in workbook for row in s for c in row)
        workbook.close();self.assertIn('NAV QC',cells);self.assertIn('Jordan Lee',cells)
        self.assertEqual(original,path.read_bytes())

    def test_manual_preserve_survives_refresh_reopen_and_backup(self):
        data=personnel_record(self.record(name='de la CRUZ',position='hSe'),preserve=True)
        identifier=self.service.add_passenger(self.plan,data)
        self.assertEqual('de la CRUZ',self.service.crew_records(self.plan,self.snapshot)[0]['name'])
        row=self.service.passenger(identifier);row['comment']='No case change'
        self.service.edit_crew_entry(identifier,row,row['version'])
        self.service.close();self.service=PersonnelService(self.root/'personnel.sqlite3')
        self.assertEqual('hSe',self.service.passenger(identifier)['position'])
        self.assertEqual([],self.service.text_case_preview(self.plan))
        PersonnelService.validate_file(self.service.path)

    def test_bulk_personnel_is_selective_atomic_and_keeps_saved_snapshots(self):
        identifier=self.legacy()
        draft=generate_plan(self.snapshot,self.service.crew_records(self.plan,self.snapshot),[],'2026-10-26','xx port',mode='Final')
        saved=self.service.save_berth(draft)
        before=list(self.service.db.execute('SELECT payload FROM berth_versions'))[0][0]
        proposals=self.service.text_case_preview(self.plan)
        position=next(p for p in proposals if p['field']=='position')
        self.assertEqual(1,self.service.apply_text_case(self.plan,[position]))
        row=self.service.passenger(identifier)
        self.assertEqual('JORDAN LEE',row['name']);self.assertEqual('NAV QC',row['position'])
        with self.assertRaises(PersonnelError):self.service.apply_text_case(self.plan,proposals)
        self.assertEqual(before,list(self.service.db.execute('SELECT payload FROM berth_versions'))[0][0])
        self.assertEqual(identifier,row['person_id'])
        self.assertEqual('Name casing normalized',self.service.history(identifier)[-1]['event'])

    def test_cargo_creation_edit_correction_next_leg_and_preserve(self):
        created=self.create_cargo('gps ANTENNA 12V')
        row=self.cargo.transfer_detail(created['no'])['row'];self.assertEqual('GPS Antenna 12V',row['description'])
        self.command('edit_transfer',{'no':row['no'],'expected_version':row['version'],'description':'PVC PIPE DN50'})
        self.assertEqual('PVC Pipe DN50',self.cargo.db.execute('SELECT default_description FROM cargo').fetchone()[0])
        row=self.cargo.transfer_detail(created['no'])['row']
        self.command('correct_transfer',{'no':row['no'],'expected_version':row['version'],'description':'special ITEM',
            'status':'Done','reason':'fixture','confirmed_location':True,'preserve_text_case':True})
        choice=self.cargo.cargo_choices()[0];self.assertEqual('special ITEM',choice['default_description'])
        next_leg=self.command('create_transfer',{'cargo_id':created['cargo_id'],'description':choice['default_description'],
            'from_location_id':self.vessel,'to_location_id':self.shore,'preserve_text_case':choice['preserve_text_case']})
        self.assertEqual('special ITEM',self.cargo.transfer_detail(next_leg['no'])['row']['description'])
        self.assertFalse(any(p['id']==next_leg['no'] and p['kind']=='transfer' for p in self.cargo.text_case_preview()))

    def test_cargo_bulk_preview_stale_rollback_and_export(self):
        created=self.create_cargo('safety helmet')
        self.cargo.db.execute('UPDATE transfer_legs SET description=? WHERE no=?',('SAFETY HELMET',created['no']))
        self.cargo.db.execute('UPDATE cargo SET default_description=? WHERE id=?',('SAFETY HELMET',created['cargo_id']))
        proposals=self.cargo.text_case_preview();self.assertEqual(2,len(proposals))
        invalid=copy.deepcopy(proposals);invalid[-1]['version']='stale'
        with self.assertRaises(AppError):self.command('normalize_text_case',{'changes':invalid})
        self.assertEqual('SAFETY HELMET',self.cargo.transfer_detail(created['no'])['row']['description'])
        self.assertEqual(2,self.command('normalize_text_case',{'changes':proposals})['changed'])
        self.assertEqual([],self.cargo.text_case_preview())
        export=self.root/'cargo.xlsx';self.cargo.export_to(export,'all',{'range':'all'},'xlsx',True)
        book=load_workbook(export,read_only=True);self.assertEqual('Safety Helmet',book['Transfers']['C2'].value);book.close()
        self.assertEqual('Waiting',self.cargo.transfer_detail(created['no'])['row']['status'])

    def test_cargo_preserve_survives_unrelated_edit_and_implicit_next_leg(self):
        created=self.create_cargo('dEliBERATE spelling',preserve_text_case=True)
        self.command('edit_transfer',{'no':created['no'],'expected_version':0,'description':'dEliBERATE spelling','comment':'Edited note'})
        row=self.cargo.transfer_detail(created['no'])['row']
        self.assertTrue(row['preserve_text_case'])
        self.command('correct_transfer',{'no':created['no'],'expected_version':row['version'],'status':'Done','reason':'fixture'})
        choice=self.cargo.cargo_choices()[0]
        next_leg=self.command('create_transfer',{'cargo_id':created['cargo_id'],'description':choice['default_description'],
            'from_location_id':self.vessel,'to_location_id':self.shore})
        self.assertTrue(self.cargo.transfer_detail(next_leg['no'])['row']['preserve_text_case'])
        self.assertEqual([],self.cargo.text_case_preview())

    def test_legacy_later_manual_spelling_is_excluded_from_bulk(self):
        identifier=self.legacy();data=clean_payload(self.service.passenger(identifier))
        data['name']='de la CRUZ'
        self.service.db.execute('UPDATE passengers SET payload=? WHERE id=?',(json.dumps(data),identifier))
        self.assertEqual(['position'],[p['field'] for p in self.service.text_case_preview(self.plan)])

    def test_empty_name_is_not_converted_to_a_literal_none(self):
        with self.assertRaises(PersonnelError):
            self.service.add_passenger(self.plan,self.record(name=None))

    def test_personnel_casing_keeps_linked_identity_and_route(self):
        first=self.service.add_passenger(self.plan,self.record())
        original=self.service.passenger(first)
        second=self.service.add_next_leg(first,self.record(name='JORDAN LEE',from_name='CHA',to_name='xx port'),self.snapshot,original['version'])
        row=self.service.passenger(second)
        self.assertEqual('Jordan Lee',row['name'])
        self.assertEqual(original['person_id'],row['person_id'])
        self.assertEqual(first,row['previous_task_id'])
        self.assertEqual((self.vessel,self.shore),(row['from_id'],row['to_id']))

    def test_unchanged_protected_legacy_name_stays_protected_when_other_field_is_cleaned(self):
        identifier=self.legacy();data=clean_payload(self.service.passenger(identifier));data['name']='de la CRUZ'
        self.service.db.execute('UPDATE passengers SET payload=? WHERE id=?',(json.dumps(data),identifier))
        proposals=self.service.text_case_preview(self.plan);self.service.apply_text_case(self.plan,proposals)
        self.assertEqual('de la CRUZ',self.service.passenger(identifier)['name'])
        self.assertTrue(self.service.passenger(identifier)['text_case']['name']['preserve'])
        PersonnelService.validate_file(self.service.path)


if __name__=='__main__':unittest.main()
