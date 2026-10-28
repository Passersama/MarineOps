"""Port boundaries and persisted automatic crew-plan updates."""
import copy
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
import test_personnel as fixtures
import test_frontend as frontend
from cargo_track.berth_plan import generate_plan, plan_assignments, route_options, export_plan
from cargo_track.crew_plan_sync import refresh_crew_plan
from cargo_track.personnel_cargo import read_cargo
from cargo_track.personnel_import import read_personnel_excel
from cargo_track.personnel_ui import CrewPlanBuilder


class CrewPlanSyncTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.PersonnelTests();self.f.setUp()
    def tearDown(self):self.f.tearDown()
    def add(self,name,direction,ship):
        data=self.f.service.bind_record(dict(name=name,position='Engineer',date='2026-10-06',
            direction=direction,vessel=ship),self.f.snapshot)
        return self.f.service.add_passenger(self.f.plan,data)
    def source(self):
        source=read_cargo(self.f.cargo_path)
        source['stops']=[dict(id='before',vessel_id=self.f.a,vessel_name_snapshot='CHA'),
            dict(id='port',vessel_id=self.f.shore,vessel_name_snapshot='xx port'),
            dict(id='after',vessel_id=self.f.a,vessel_name_snapshot='CHA'),
            dict(id='destination',vessel_id=self.f.c,vessel_name_snapshot='INN')]
        return source
    def save(self):
        source=self.source();people=self.f.service.crew_records(self.f.plan,source)
        plan=generate_plan(source,people,[],'2026-10-06','xx port',mode='Final',personnel_plan_id=self.f.plan)
        plan['personnel_context']=self.f.service.context(self.f.plan)
        self.assertFalse(plan['issues'],plan['issues'])
        return self.f.service.save_berth(plan,True)
    def test_same_ship_before_and_after_port_is_separated_by_import_direction(self):
        off=self.add('Off Crew','Sign-off','CHA');on=self.add('On Crew','Sign-on','CHA')
        plan=self.save();assignments=plan_assignments(plan)
        self.assertEqual(assignments['Passengers:'+off],dict(pick='before',drop='port'))
        self.assertEqual(assignments['Passengers:'+on],dict(pick='port',drop='after'))
        self.assertEqual(1,len(plan['stops'][1]['drop_passengers']))
        self.assertEqual(1,len(plan['stops'][1]['pick_passengers']))
    def test_sign_on_has_no_delivery_before_the_only_port(self):
        identifier=self.add('On Crew','Sign-on','CHA')
        source=self.source();source['stops']=source['stops'][:2]
        person=self.f.service.crew_records(self.f.plan,source)[0]
        self.assertEqual([],route_options(source,person,'Final'))
        plan=generate_plan(source,[person],[],'2026-10-06','xx port',mode='Final')
        self.assertTrue(plan['issues']);self.assertNotIn('Passengers:'+identifier,plan['included'])
    def test_master_edit_updates_saved_from_to_names_and_export_without_generate(self):
        identifier=self.add('On Crew','Sign-on','CHA');original=self.save()
        cargo_before=self.f.cargo_path.read_bytes()
        record=self.f.service.passenger(identifier)
        edited=self.f.service.bind_record({**record,'name':'Updated Crew','to_name':'INN','vessel':'INN'},self.f.snapshot)
        self.f.service.edit_crew_entry(identifier,edited,record['version'])
        latest=self.f.service.berth_versions()[0]
        self.assertNotEqual(original['id'],latest['id']);self.assertEqual(original['family_id'],latest['family_id'])
        self.assertFalse(latest['issues']);self.assertEqual('Published',latest['status'])
        self.assertEqual(dict(pick='port',drop='destination'),plan_assignments(latest)['Passengers:'+identifier])
        self.assertEqual('Updated Crew',latest['stops'][-1]['drop_passengers'][0]['name'])
        self.assertEqual('INN',latest['stops'][-1]['drop_passengers'][0]['to_name'])
        self.assertEqual(original,next(p for p in self.f.service.berth_versions() if p['id']==original['id']))
        path=self.f.root/'updated.xlsx';export_plan(latest,path)
        book=load_workbook(path)
        try:
            values=' '.join(str(cell.value) for row in book['Berthing Plan'] for cell in row if cell.value)
            self.assertIn('Updated Crew',values)
        finally:book.close()
        self.assertEqual(cargo_before,self.f.cargo_path.read_bytes())
    def test_add_and_remove_refresh_current_version_and_context(self):
        first=self.add('On Crew','Sign-on','CHA');self.save()
        second=self.add('Second Crew','Sign-off','CHA')
        latest=self.f.service.berth_versions()[0]
        self.assertIn('Passengers:'+second,latest['included'])
        self.f.service.remove_crew_entries({first:0})
        latest=self.f.service.berth_versions()[0]
        self.assertNotIn('Passengers:'+first,latest['included'])
        self.assertIn('Passengers:'+second,latest['included'])
        self.assertEqual(self.f.service.context(self.f.plan),latest['personnel_context'])
    def test_import_replace_updates_current_plan_in_same_transaction(self):
        from openpyxl import Workbook
        identifier=self.add('On Crew','Sign-on','CHA');self.save()
        book=Workbook();sheet=book.active
        sheet.append(['Name','Position','From','To','Date','Task ID'])
        sheet.append(['On Crew','Chief Engineer','xxPort','INN','2026-10-06',identifier])
        path=self.f.root/'replacement.xlsx';book.save(path);book.close()
        parsed=read_personnel_excel(path)
        preview=self.f.service.import_preview(parsed,self.f.snapshot,self.f.plan,mode='replace',crew_list=True)
        preview['items'][0]['match_id']=identifier
        self.f.service.commit_import(preview,plan_date='2026-10-06',cancel_missing=True)
        latest=self.f.service.berth_versions()[0]
        self.assertFalse(latest['issues'],latest['issues'])
        task=next(t for s in latest['stops'] for t in s['drop_passengers'] if t['id']==identifier)
        self.assertEqual(('Chief Engineer','INN'),(task['position'],task['to_name']))
    def test_sign_off_source_change_updates_pickup_and_adds_missing_call_before_port(self):
        identifier=self.add('Off Crew','Sign-off','CHA');self.save()
        record=self.f.service.passenger(identifier)
        data=self.f.service.bind_record({**record,'from_name':'Niger','vessel':'Niger'},self.f.snapshot)
        self.f.service.edit_crew_entry(identifier,data,record['version'])
        latest=self.f.service.berth_versions()[0];assignment=plan_assignments(latest)['Passengers:'+identifier]
        source=next(s for s in latest['stops'] if s['id']==assignment['pick'])
        self.assertEqual(self.f.b,source['vessel_id'])
        self.assertEqual(assignment['drop'],'port');self.assertFalse(latest['issues'])
    def test_failed_sync_rolls_back_roster_and_versions_together(self):
        identifier=self.add('On Crew','Sign-on','CHA');self.save()
        before=self.f.service.passenger(identifier);versions=self.f.service.berth_versions()
        with patch('cargo_track.crew_plan_sync.refresh_crew_plan',side_effect=RuntimeError('Unexpected sync failure')):
            with self.assertRaises(RuntimeError):
                self.f.service.edit_crew_entry(identifier,{**before,'name':'Changed'},before['version'])
        self.assertEqual(before,self.f.service.passenger(identifier))
        self.assertEqual(versions,self.f.service.berth_versions())
    def test_crew_sync_does_not_select_other_available_materials(self):
        identifier=self.add('On Crew','Sign-on','CHA')
        numbers=[self.f.command('create_transfer',dict(description='Material '+str(i),from_location_id=self.f.a,
                 to_location_id=self.f.c))['no'] for i in range(2)]
        source=self.source();people=self.f.service.crew_records(self.f.plan,source)
        plan=generate_plan(source,people,[numbers[0]],'2026-10-06','xx port',mode='Final',personnel_plan_id=self.f.plan)
        plan['available_tasks']=['Passengers:'+identifier]+['Materials:'+str(n) for n in numbers]
        plan['personnel_context']=self.f.service.context(self.f.plan)
        saved=self.f.service.save_berth(plan,True)
        row=self.f.service.passenger(identifier)
        self.f.service.edit_crew_entry(identifier,{**row,'position':'Chief Engineer'},row['version'])
        current=self.f.service.berth_versions()[0]
        self.assertEqual(saved['source_dependencies']['cargo_ids'],current['source_dependencies']['cargo_ids'])
        self.assertIn('Materials:'+str(numbers[0]),current['included'])
        self.assertNotIn('Materials:'+str(numbers[1]),current['selected_tasks'])


class CrewPlanSyncEditorTests(unittest.TestCase):
    setUpClass=classmethod(frontend.FrontendTests.setUpClass.__func__)
    setUp=frontend.FrontendTests.setUp
    tearDown=frontend.FrontendTests.tearDown
    wait=frontend.FrontendTests.wait
    def test_new_plan_cannot_deliver_sign_on_after_departure_port_is_removed(self):
        builder=CrewPlanBuilder(self.page)
        try:
            self.assertEqual(self.f.shore,builder.route.item(0).data(256)['vessel_id'])
            builder.route.setCurrentRow(0);builder.remove_stop()
            self.assertTrue(builder.build()['issues'])
        finally:builder.deleteLater()
    def test_saved_plan_updates_on_form_edit_without_generate(self):
        saved=self.page.service.save_berth(self.plan,True)
        self.page.saved_plan=saved;self.page.show_plan(saved);self.page.refresh_rows()
        self.page.passenger_table.selectRow(0)
        record=self.page.selected()[0]
        edited=self.page.service.bind_record({**record,'name':'Updated In Form','to_name':'INN'},self.page.source)
        with patch('cargo_track.personnel_ui.PassengerForm.exec',return_value=1),patch('cargo_track.personnel_ui.PassengerForm.data',return_value=edited):
            self.page.edit_passenger();self.wait(lambda:not self.page._write_pending)
        plan=self.page.displayed_plan
        self.assertFalse(plan['issues'],plan['issues'])
        self.assertNotEqual(saved['id'],plan['id'])
        task=next(t for s in plan['stops'] for t in s['drop_passengers'] if t['id']==record['id'])
        self.assertEqual(('Updated In Form','INN'),(task['name'],task['to_name']))
        self.assertEqual([],self.errors)
    def test_live_draft_keeps_comment_and_updates_endpoints(self):
        self.page.draft=copy.deepcopy(self.plan);self.page.show_plan(self.plan)
        builder=self.page.live_builder;stop_id=builder.route.item(0).data(256)['id']
        builder.stop_comments[stop_id]='Keep this instruction'
        self.page.refresh_rows();self.page.passenger_table.selectRow(0)
        record=self.page.selected()[0]
        edited=self.page.service.bind_record({**record,'to_name':'INN','name':'Changed Draft Crew'},self.page.source)
        with patch('cargo_track.personnel_ui.PassengerForm.exec',return_value=1),patch('cargo_track.personnel_ui.PassengerForm.data',return_value=edited):
            self.page.edit_passenger();self.wait(lambda:not self.page._write_pending)
        self.assertIs(builder,self.page.live_builder)
        plan=builder.build();self.assertFalse(plan['issues'],plan['issues'])
        self.assertEqual('Keep this instruction',next(s['comment'] for s in plan['stops'] if s['id']==stop_id))
        self.assertEqual('INN',next(t['to_name'] for s in plan['stops'] for t in s['drop_passengers'] if t['id']==record['id']))
        self.assertFalse(self.page.view_state.source_invalid)
