"""Personnel multi-leg, optional Port, transfer-origin and screenshot export regressions."""
import copy
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
import test_personnel as fixtures
from cargo_track.berth_plan import generate_plan, route_options, plan_assignments, plan_text, export_plan
from cargo_track.personnel_cargo import read_cargo
from cargo_track.personnel_import import PersonnelError
from cargo_track.personnel_service import PersonnelService


class PersonnelRouteTests(unittest.TestCase):
    def test_legacy_missing_and_text_leg_numbers_have_numeric_order(self):
        from cargo_track.personnel_routes import leg_order
        rows=[{'leg_index':2},{'leg_index':''},{},{'leg_index':'10'},{'leg_index':None}]
        self.assertEqual([1,1,1,2,10],sorted(leg_order(r) for r in rows))
        for value in (True,0,-1,'invalid',1.5):
            with self.assertRaises(PersonnelError):leg_order({'leg_index':value})

    def test_new_snapshot_of_legacy_root_and_linked_legs_keeps_integer_indexes(self):
        from cargo_track.personnel_routes import leg_order
        first,second,source=self.chain();legacy=self.add('CHA','INN','Legacy Fictional')
        people=self.service.crew_records(self.f.plan,source)
        row=next(r for r in people if r['id']==legacy);row.pop('leg_index');row.pop('person_id')
        before=self.service.path.read_bytes()
        plan=generate_plan(source,people,[],'2026-10-12',self.service.port_name(source),mode='Final',personnel_plan_id=self.f.plan)
        grouped={t['id']:t for s in plan['stops'] for g in ('pick_passengers','drop_passengers') for t in s[g]}
        self.assertEqual(1,grouped[legacy]['leg_index']);self.assertEqual(2,grouped[second]['leg_index'])
        self.assertTrue(all(type(t['leg_index']) is int for t in grouped.values()))
        self.assertEqual([1,1,2],[leg_order(t) for t in sorted(grouped.values(),key=leg_order)])
        self.assertEqual(before,self.service.path.read_bytes())

    def setUp(self):
        self.f = fixtures.PersonnelTests();self.f.setUp()
        self.service = self.f.service

    def tearDown(self):
        self.f.tearDown()

    def source(self):
        return self.service.planning_source(read_cargo(self.f.cargo_path))

    def record(self, source, target, name='Alex Fictional'):
        return self.service.bind_record({'name':name, 'position':'OBS', 'date':'2026-10-06',
            'direction':'Transfer', 'from_name':source, 'to_name':target}, self.source())

    def add(self, source, target, name='Alex Fictional'):
        return self.service.add_passenger(self.f.plan, self.record(source, target, name))

    def plan(self, source=None, assignments=None, selections=None):
        source = source or self.source()
        return generate_plan(source, self.service.crew_records(self.f.plan, source), [], '2026-10-06',
            self.service.port_name(source), mode='Final', assignments=assignments, selections=selections,
            personnel_plan_id=self.f.plan)

    def chain(self):
        first = self.add('CHA', 'Niger')
        old = self.service.passenger(first)
        second = self.service.add_next_leg(first, self.record('Niger', 'INN'), self.source(), old['version'])
        source = self.source()
        repeated = dict(source['stops'][1], id='plan:second-niger')
        source['stops'].insert(2, repeated)
        return first, second, source

    def test_no_automatic_port_and_cargo_database_is_unchanged(self):
        self.add('CHA', 'Niger')
        before = self.f.cargo_path.read_bytes()
        plan = self.plan()
        self.assertEqual([], plan['issues'])
        self.assertEqual(['CHA','Niger','INN'], [s['name'] for s in plan['stops']])
        self.service.save_berth(plan, True)
        self.assertEqual(before, self.f.cargo_path.read_bytes())

    def test_port_delivery_requires_explicit_stop_and_can_be_placed_in_middle(self):
        person = self.add('CHA', 'xxPort')
        missing = self.plan()
        self.assertTrue(any('xx port' in issue['message'].casefold() or 'xxport' in issue['message'].casefold() for issue in missing['issues']))
        source = self.source()
        source['stops'].insert(1, {'id':'plan:port', 'vessel_id':self.f.shore, 'vessel_name_snapshot':'xx port'})
        plan = self.plan(source)
        self.assertEqual([], plan['issues'])
        self.assertEqual(person, plan['stops'][1]['drop_passengers'][0]['id'])
        self.assertEqual('INN', plan['stops'][-1]['name'])

    def test_transfer_vessel_from_matches_only_destination_without_source_visit(self):
        self.service.set_setting('transfer_vessel_name', 'Niger')
        person = self.add('Niger', 'INN')
        source = self.source();source['stops'] = [source['stops'][-1]]
        plan = self.plan(source)
        self.assertEqual([], plan['issues'])
        self.assertEqual(person, plan['stops'][0]['drop_passengers'][0]['id'])
        self.assertFalse(any(s['pick_passengers'] for s in plan['stops']))
        self.assertEqual('Niger', plan['stops'][0]['drop_passengers'][0]['from_name'])

    def test_virtual_transfer_vessel_import_save_and_backup_are_valid(self):
        self.service.set_setting('transfer_vessel_name', 'Plan Boat')
        data = self.record('  PLAN Boat ', 'INN')
        self.assertEqual([], data['issues']);self.assertIsNone(data['from_id'])
        self.service.add_passenger(self.f.plan, data)
        plan = self.plan();self.assertEqual([], plan['issues'])
        self.service.save_berth(plan, True)
        backup = self.f.root / 'virtual-origin.sqlite3'
        self.service.backup(backup);PersonnelService.validate_file(backup)
        self.assertFalse(any(s['pick_passengers'] for s in plan['stops']))

    def test_empty_or_other_from_is_not_assumed_to_be_transfer_vessel(self):
        self.service.set_setting('transfer_vessel_name', 'Plan Boat')
        self.assertTrue(self.record('', 'INN')['issues'])
        self.assertTrue(self.record('Plan Boat Two', 'INN')['issues'])
        person = self.add('CHA', 'INN')
        self.assertEqual(self.f.s1, plan_assignments(self.plan())['Passengers:' + person]['pick'])

    def test_personnel_override_takes_precedence_over_cargo_transfer_setting(self):
        self.f.command('set_setting', {'key':'transfer_vessel_name', 'value':'CHA'})
        self.service.set_setting('transfer_vessel_name', 'Niger')
        a = self.add('CHA', 'INN', 'First Fictional')
        b = self.add('Niger', 'INN', 'Second Fictional')
        plan = self.plan();choices = plan_assignments(plan)
        self.assertEqual(self.f.s1, choices['Passengers:' + a]['pick'])
        self.assertIsNone(choices['Passengers:' + b]['pick'])
        self.assertEqual('Niger', plan['transfer_vessel'])

    def test_material_source_rule_remains_normal_pickup(self):
        source = self.source();source['plan_transfer_vessel'] = 'CHA'
        task = {'kind':'Materials','from_id':self.f.a,'to_id':self.f.c,'status':'Waiting'}
        self.assertEqual([{'pick':self.f.s1,'drop':self.f.s3}], route_options(source, task, 'Final'))

    def test_next_leg_reboards_at_later_repeated_visit(self):
        first, second, source = self.chain()
        plan = self.plan(source);self.assertEqual([], plan['issues'])
        assignments = plan_assignments(plan)
        self.assertEqual({'pick':self.f.s1,'drop':self.f.s2}, assignments['Passengers:' + first])
        self.assertEqual({'pick':'plan:second-niger','drop':self.f.s3}, assignments['Passengers:' + second])
        self.assertEqual(self.service.passenger(first)['person_id'], self.service.passenger(second)['person_id'])
        self.assertEqual(2, self.service.passenger(second)['leg_index'])

    def test_linked_leg_cannot_pickup_on_the_preceding_drop_visit(self):
        first, second, source = self.chain()
        invalid = self.plan(source, {'Passengers:' + second:{'pick':self.f.s2,'drop':self.f.s3}})
        self.assertTrue(any('no longer valid' in issue['message'] for issue in invalid['issues']))
        with self.assertRaises(PersonnelError):self.service.save_berth(invalid, True)

    def test_linked_leg_requires_a_later_source_call(self):
        first, second, source = self.chain()
        source['stops'] = [s for s in source['stops'] if s['id'] != 'plan:second-niger']
        invalid = self.plan(source)
        self.assertTrue(any('later visit' in issue['message'] for issue in invalid['issues']))

    def test_next_leg_from_port_requires_a_later_real_pickup(self):
        first = self.add('CHA', 'xxPort')
        second = self.service.add_next_leg(first, self.record('xxPort', 'INN'), self.source(), 0)
        source = self.source()
        ports = [{'id': 'plan:port-' + str(i), 'vessel_id': self.f.shore,
                  'vessel_name_snapshot': 'xxPort'} for i in range(2)]
        source['stops'][1:1] = ports
        plan = self.plan(source)
        self.assertEqual([], plan['issues'])
        choices = plan_assignments(plan)
        self.assertEqual('plan:port-0', choices['Passengers:' + first]['drop'])
        self.assertEqual({'pick': 'plan:port-1', 'drop': self.f.s3}, choices['Passengers:' + second])
        saved = self.service.save_berth(plan, True)
        export_plan(saved, self.f.root / 'port-next-leg.xlsx')
        source['stops'].remove(ports[1])
        self.assertTrue(self.plan(source)['issues'])

    def test_linked_leg_from_transfer_vessel_requires_a_real_return_visit(self):
        first, second, source = self.chain()
        self.service.set_setting('transfer_vessel_name', 'Niger')
        source = self.service.planning_source(source)
        plan = self.plan(source)
        self.assertEqual([], plan['issues'])
        self.assertEqual('plan:second-niger', plan_assignments(plan)['Passengers:' + second]['pick'])

    def test_excluded_preceding_leg_cannot_silently_detach_next_leg(self):
        first, second, source = self.chain()
        plan = self.plan(source, selections={'Passengers:' + first: {'state': 'user_excluded', 'reason': 'Test'}})
        self.assertTrue(any(p['task'] == 'Passengers:' + second for p in plan['issues']))
        self.assertNotIn('Passengers:' + second, plan['included'])
        with self.assertRaises(PersonnelError):self.service.save_berth(plan, True)

    def test_same_name_rows_are_distinct_until_user_links_them(self):
        first = self.add('CHA', 'Niger');second = self.add('Niger', 'INN')
        self.assertNotEqual(self.service.passenger(first)['person_id'], self.service.passenger(second)['person_id'])
        self.service.link_crew_legs([first, second], {first:0,second:0})
        self.assertEqual(self.service.passenger(first)['person_id'], self.service.passenger(second)['person_id'])
        self.assertEqual(first, self.service.passenger(second)['previous_task_id'])

    def test_stale_or_discontinuous_link_is_atomic(self):
        first = self.add('CHA', 'Niger');second = self.add('CHA', 'INN')
        before = self.service.passenger(second)
        with self.assertRaises(PersonnelError):self.service.link_crew_legs([first,second], {first:0,second:0})
        self.assertEqual(before, self.service.passenger(second))
        with self.assertRaises(PersonnelError):self.service.link_crew_legs([first,second], {first:0,second:99})
        self.assertEqual(before, self.service.passenger(second))

    def test_next_leg_does_not_create_a_second_branch(self):
        first, second, source = self.chain()
        before = self.service.passengers(self.f.plan)
        with self.assertRaises(PersonnelError):
            self.service.add_next_leg(first, self.record('Niger','INN'), source, 0)
        self.assertEqual(before, self.service.passengers(self.f.plan))

    def test_transfer_setting_change_invalidates_plan_context(self):
        self.add('CHA','INN')
        plan = self.plan();plan['personnel_context'] = self.service.context(self.f.plan)
        self.service.set_setting('transfer_vessel_name','CHA')
        with self.assertRaises(PersonnelError):self.service.save_berth(plan, True)

    def test_changed_preceding_destination_requires_review_not_substitution(self):
        first, second, source = self.chain()
        old = self.service.passenger(first)
        self.service.edit_crew_entry(first, self.record('CHA','INN'), old['version'])
        plan = self.plan(source)
        self.assertTrue(any('preceding destination' in issue['message'] for issue in plan['issues']))

    def test_export_matches_four_row_merged_stop_layout_and_inline_counts(self):
        first, second, source = self.chain()
        plan = self.service.save_berth(self.plan(source), True)
        path = self.f.root / 'four-row-plan.xlsx';export_plan(plan, path)
        book = load_workbook(path)
        try:
            sheet = book['Berthing Plan']
            self.assertEqual('Comments', sheet['D1'].value)
            self.assertTrue(sheet['B1'].is_date)
            self.assertEqual('F58220', sheet['C1'].fill.fgColor.rgb[-6:])
            self.assertEqual(1 + 4 * len(plan['stops']), sheet.max_row)
            self.assertEqual({'A2:A5','B2:B5','D2:D5'}, {str(r) for r in sheet.merged_cells.ranges if r.min_row == 2})
            self.assertEqual('Full speed to CHA', sheet['B2'].value)
            self.assertEqual('Pick up 1 passengers: 1. Alex Fictional(OBS)', sheet['C3'].value)
            self.assertEqual('Drop 0 passengers:', sheet['C2'].value)
            self.assertEqual(0, len(sheet.row_breaks.brk))
            self.assertEqual('$1:$1', sheet.print_title_rows)
            self.assertFalse(any(cell.data_type == 'f' for ws in book for row in ws for cell in row))
            self.assertEqual(['Person ID','Leg','Previous task ID'], [book['Transfer Details'].cell(1,c).value for c in (17,18,19)])
            self.assertIn('Pick up 1 passengers: 1. Alex Fictional(OBS)', plan_text(plan))
        finally:book.close()

    def test_saved_version_and_database_survive_export_failure(self):
        self.add('CHA', 'Niger')
        saved = self.service.save_berth(self.plan(), True)
        path = self.f.root / 'keep.xlsx';path.write_bytes(b'previous export')
        with patch('cargo_track.berth_plan.os.replace', side_effect=OSError('Write denied')):
            with self.assertRaises(OSError):export_plan(saved,path)
        self.assertEqual(b'previous export', path.read_bytes())
        self.assertEqual(saved, self.service.berth_versions()[0])

    def test_portrait_export_wraps_long_english_and_chinese_rows(self):
        self.add('CHA', 'INN', 'Fictional Passenger ' + 'Long Name ' * 12)
        plan = self.service.save_berth(self.plan(), True)
        before = copy.deepcopy(plan)
        path = self.f.root / 'portrait.xlsx'
        export_plan(plan, path)
        book = load_workbook(path)
        try:
            sheet = book['Berthing Plan']
            self.assertEqual('portrait', sheet.page_setup.orientation)
            self.assertEqual(int(sheet.PAPERSIZE_A4), sheet.page_setup.paperSize)
            self.assertEqual(1, sheet.page_setup.fitToWidth)
            self.assertEqual(0, sheet.page_setup.fitToHeight)
            self.assertEqual(60, sheet.column_dimensions['C'].width)
            self.assertGreater(sheet.row_dimensions[3].height, sheet.row_dimensions[2].height)
            self.assertTrue(sheet['C3'].alignment.wrap_text)
        finally:book.close()
        self.assertEqual(before, plan)

    def test_legacy_complete_snapshot_remains_immutable_and_exportable(self):
        self.add('CHA','Niger')
        saved = self.service.save_berth(self.f.generate(), True);frozen=copy.deepcopy(saved)
        export_plan(saved, self.f.root / 'legacy.xlsx')
        self.assertEqual('PORT', saved['stops'][-1]['id'])
        self.assertEqual(frozen, self.service.berth_versions()[0])


if __name__ == '__main__':
    unittest.main()
