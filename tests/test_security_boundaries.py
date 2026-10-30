"""External-file regressions with bounded fictional inputs and no network use."""
from pathlib import Path, PureWindowsPath
from types import SimpleNamespace
from unittest.mock import patch
import hashlib
import io
import json
import os
import sqlite3
import stat
import struct
import tempfile
import unittest
from uuid import uuid4
import warnings
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook
import cargo_track.personnel_import as imports
import cargo_track.source_paths as paths
import cargo_track.xlsx_safety as safety
from cargo_track.berth_plan import export_plan
from cargo_track.coordination import WorkspaceFiles, ProtectedFileError
from cargo_track.personnel_service import PersonnelService
import test_dashboard as fixtures


BAD_PATHS = (r'\\audit.invalid\share\ledger.sqlite3', '//audit.invalid/share/ledger.sqlite3',
             r'\/audit.invalid/share/ledger.sqlite3', r'\\?\UNC\audit.invalid\share\ledger.sqlite3',
             r'\\?\C:\ledger.sqlite3', r'\\.\pipe\probe', 'file://audit.invalid/share/ledger.sqlite3',
             'FILE:/C:/ledger.sqlite3', 'https://audit.invalid/ledger.sqlite3',
             r'C:ledger.sqlite3', r'\ledger.sqlite3', r'C:\NUL\ledger.sqlite3', 'bad\0path')


class LocalSourceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def test_unsafe_encodings_rejected_before_probe(self):
        for value in BAD_PATHS:
            with self.subTest(value=value), patch.object(paths.os,'lstat',side_effect=AssertionError('File probe')):
                with self.assertRaises(paths.SourcePathError):paths.local_source_path(value,allow_relative=True)
    def test_metadata_matching_has_no_filesystem_effect(self):
        with patch.object(paths.os,'lstat',side_effect=AssertionError('stat')),patch.object(Path,'resolve',side_effect=AssertionError('resolve')):
            self.assertEqual(paths.source_path_key(self.root/'ledger.sqlite3'),paths.source_path_key(self.root/'child'/'..'/'ledger.sqlite3'))
    def test_remote_drive_is_rejected_before_stat(self):
        with patch.object(paths,'_drive_is_local',return_value=False),patch.object(paths.os,'lstat',side_effect=AssertionError('stat')):
            with self.assertRaises(paths.SourcePathError):paths.local_source_path(self.root/'ledger.sqlite3')
    def link_probe(self,target):
        link=self.root/'alias';original=os.lstat
        def probe(value,*args,**kwargs):
            if Path(value)==link:return SimpleNamespace(st_mode=stat.S_IFLNK,st_file_attributes=stat.FILE_ATTRIBUTE_REPARSE_POINT,st_reparse_tag=stat.IO_REPARSE_TAG_SYMLINK)
            return original(value,*args,**kwargs)
        return link,patch.object(paths.os,'lstat',probe),patch.object(paths.os,'readlink',return_value=target)
    def test_remote_junction_target_rejected_without_following(self):
        for target in (r'\\audit.invalid\share',r'\\?\UNC\audit.invalid\share'):
            link,probe,readlink=self.link_probe(target)
            with self.subTest(target=target),probe,readlink,patch.object(Path,'resolve',side_effect=AssertionError('resolve')):
                with self.assertRaises(paths.SourcePathError):paths.local_source_path(link/'ledger.sqlite3')
    def test_verified_local_alias_is_kept(self):
        local=self.root/'actual';local.mkdir();(local/'ledger.sqlite3').write_text('fictional')
        link,probe,readlink=self.link_probe(str(local))
        with probe,readlink:self.assertEqual(local/'ledger.sqlite3',paths.local_source_path(link/'ledger.sqlite3'))
    def test_alias_cycle_is_bounded(self):
        link,probe,readlink=self.link_probe('alias')
        with probe,readlink as read:
            with self.assertRaises(paths.SourcePathError):paths.local_source_path(link/'ledger.sqlite3')
            self.assertEqual(40,read.call_count)
    def test_source_hardlink_is_protected(self):
        source=self.root/'ledger.sqlite3';source.write_text('fictional');alias=self.root/'alias.xlsx'
        try:os.link(source,alias)
        except OSError:self.skipTest('Hard links unavailable')
        with self.assertRaises(ProtectedFileError):WorkspaceFiles().check_output(alias,source_references=(source,))
    def test_explicit_network_outputs_and_trusted_paths_keep_their_authority(self):
        with patch.object(Path,'resolve',lambda p,*a,**k:p),patch.object(os.path,'samefile',return_value=False),patch('cargo_track.coordination.source_path_key',side_effect=AssertionError('saved-reference restriction')):
            WorkspaceFiles().check_output(Path(r'\\chosen.invalid\exports\file.xlsx'),r'\\chosen.invalid\data\ledger.sqlite3')


class SavedReferenceTests(unittest.TestCase):
    def setUp(self):self.f=fixtures.DashboardTests();self.f.setUp()
    def tearDown(self):self.f.tearDown()
    def insert(self,database,field):
        value=dict(self.f.saved);identifier=str(uuid4())
        value.update(id=identifier,family_id=identifier,version=1)
        value[field]=BAD_PATHS[0]
        database.execute('INSERT INTO berth_versions VALUES(?,?,?,?,?,?)',(identifier,identifier,1,value['status'],json.dumps(value),'2099-01-01T00:00:00+00:00'))
    def test_dashboard_rejects_unsafe_metadata_without_resolving_it(self):
        self.insert(self.f.personnel.db,'source_path');original=Path.resolve
        def resolve(value,*args,**kwargs):
            if PureWindowsPath(str(value)).drive.startswith('\\\\'):raise AssertionError('UNC filesystem access')
            return original(value,*args,**kwargs)
        before=self.f.personnel.path.read_bytes()
        with patch.object(Path,'resolve',resolve):self.assertEqual('error',self.f.read()['personnel']['status'])
        self.assertEqual(before,self.f.personnel.path.read_bytes())
    def test_both_export_reference_fields_rejected_before_writing(self):
        for field in ('source_path','personnel_source_path'):
            value=dict(self.f.saved);value[field]=BAD_PATHS[0];destination=self.f.root/(field+'.xlsx')
            with self.subTest(field=field),self.assertRaises(imports.PersonnelError):export_plan(value,destination)
            self.assertFalse(destination.exists())
    def test_restore_rejects_unsafe_plan_reference_and_preserves_live(self):
        for field in ('source_path','personnel_source_path'):
            candidate=self.f.root/(field+'.sqlite3');self.f.personnel.backup(candidate)
            db=sqlite3.connect(candidate)
            try:self.insert(db,field);db.commit()
            finally:db.close()
            before=self.f.personnel.path.read_bytes()
            with self.subTest(field=field),self.assertRaises(imports.PersonnelError):self.f.personnel.restore(candidate)
            self.assertEqual(before,self.f.personnel.path.read_bytes())


class WorkbookResourceTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.path=self.root/'crew.xlsx'
        self.original=self.workbook('Original Person');self.path.write_bytes(self.original)
    def tearDown(self):self.temp.cleanup()
    @staticmethod
    def workbook(name):
        book=Workbook();book.active.append(['Name','Position','From','To']);book.active.append([name,'Engineer','Ship A','Ship B'])
        stream=io.BytesIO();book.save(stream);book.close();return stream.getvalue()
    def rewrite(self,extra=(),override=''):
        with ZipFile(io.BytesIO(self.original)) as source,ZipFile(self.path,'w',ZIP_DEFLATED) as target:
            for entry in source.infolist():
                data=source.read(entry.filename)
                if entry.filename=='[Content_Types].xml' and override:data=data.replace(b'</Types>',override.encode()+b'</Types>')
                target.writestr(entry,data)
            with warnings.catch_warnings():
                warnings.simplefilter('ignore',UserWarning)
                for name,data in extra:target.writestr(name,data)
    def rejects_before_openpyxl(self):
        operations=(lambda:imports.workbook_sheets(self.path),lambda:imports.read_personnel_excel(self.path),
                    lambda:imports.read_personnel_excel(self.path,mapping={'name':1,'position':2,'from_name':3,'to_name':4},header_row=1))
        with patch.object(imports,'load_workbook',side_effect=AssertionError('Unsafe input reached parser')):
            for operation in operations:
                with self.assertRaises(imports.PersonnelError):operation()
    def test_original_shared_string_attack_is_bounded_by_object_limit(self):
        xml=('<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'+'<si><t>A</t></si>'*(safety.MAX_SHARED_STRINGS+1)+'</sst>').encode()
        self.rewrite([('xl/strings.data',xml)],'<Override PartName="/xl/strings.data" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>')
        self.rejects_before_openpyxl()
    def test_expanded_member_budget(self):
        self.rewrite([('extra.bin',b'A'*4097)])
        with patch.object(safety,'MAX_MEMBER',4096):self.rejects_before_openpyxl()
    def test_total_actual_expansion_budget(self):
        with ZipFile(io.BytesIO(self.original)) as source:base=sum(i.file_size for i in source.infolist())
        self.rewrite([('one.bin',b'A'*2000),('two.bin',b'B'*2000)])
        with patch.object(safety,'MAX_EXPANDED',base+3000):self.rejects_before_openpyxl()
    def test_expansion_ratio_budget(self):
        self.rewrite([('extra.bin',b'A'*20000)])
        with patch.object(safety,'MAX_RATIO',100):self.rejects_before_openpyxl()
    def test_declared_central_count_rejected_before_zip_member_allocation(self):
        data=bytearray(self.original);index=data.rfind(b'PK\x05\x06')
        struct.pack_into('<HH',data,index+8,safety.MAX_MEMBERS+1,safety.MAX_MEMBERS+1);self.path.write_bytes(data)
        with patch.object(safety,'ZipFile',side_effect=AssertionError('central directory allocation')):self.rejects_before_openpyxl()
    def test_duplicate_parts_rejected(self):
        self.rewrite([('xl/workbook.xml',b'<workbook/>')]);self.rejects_before_openpyxl()
    def test_utf16_dtd_in_custom_xml_part_rejected(self):
        xml='<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE r [<!ENTITY x "probe">]><r>&x;</r>'.encode('utf-16')
        self.rewrite([('xl/custom.data',xml)],'<Override PartName="/xl/custom.data" ContentType="application/xml"/>');self.rejects_before_openpyxl()
    def test_ambiguous_manifest_part_reference_rejected(self):
        self.rewrite([('/xl/custom.data',b'<r/>')],'<Override PartName="//xl/custom.data" ContentType="application/xml"/>');self.rejects_before_openpyxl()
    def test_original_bytes_parser_and_hash_remain_bound_after_replacement(self):
        original_loader=imports.load_workbook
        def replace_after_check(frozen,**kwargs):
            self.path.write_bytes(self.workbook('Replacement Person'))
            return original_loader(frozen,**kwargs)
        with patch.object(imports,'load_workbook',replace_after_check):parsed=imports.read_personnel_excel(self.path)
        self.assertEqual('Original Person',parsed['records'][0]['name'])
        self.assertEqual(hashlib.sha256(self.original).hexdigest(),parsed['file_hash'])
        self.assertNotEqual(parsed['file_hash'],hashlib.sha256(self.path.read_bytes()).hexdigest())
    def test_existing_row_column_limits_and_unselected_sheet_semantics_remain(self):
        book=Workbook();sheet=book.active;sheet.append(['Name','Position','From','To']);sheet.append(['Valid Person','Engineer','A','B'])
        sheet.cell(20000,200,'Ignored');book.create_sheet('Unselected').cell(1000000,1,'Ignored')
        book.save(self.path);book.close()
        self.assertEqual(1,len(imports.read_personnel_excel(self.path)['records']))
        self.assertEqual(['Sheet','Unselected'],imports.workbook_sheets(self.path))
