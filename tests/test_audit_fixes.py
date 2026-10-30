"""Regression cases from the 2026-10-07 audit, with isolated business fixtures."""
import copy
import hashlib
import os
import sqlite3
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from PySide6.QtWidgets import QApplication, QWidget
from openpyxl import Workbook, load_workbook
import test_personnel as fixtures
from cargo_track.berth_plan import generate_plan, export_plan, plan_assignments, task_key, plan_text
from cargo_track.coordination import workspace_files
from cargo_track.domain import AppError
from cargo_track.personnel_cargo import read_cargo, dependency_hash
from cargo_track.personnel_import import PersonnelError, read_personnel_excel
from cargo_track.personnel_service import PersonnelService
from cargo_track.personnel_read_model import PersonnelReadModel
from cargo_track.personnel_ui import CrewPlanBuilder


class Owner(QWidget):
    def __init__(self,fixture):
        super().__init__()
        self.service=fixture.service
        self.host=SimpleNamespace(database_path=fixture.cargo_path)
        self.plan_combo=SimpleNamespace(currentData=lambda:fixture.plan)
        self.plan_date=lambda:'2026-10-06'

    @property
    def read_model(self):
        return PersonnelReadModel(self.service.path,self.service.ui_state(self.plan_combo.currentData()))

    @property
    def source(self):return read_cargo(self.host.database_path,active_only=True)


class AuditFixTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app=QApplication.instance() or QApplication([])

    def setUp(self):
        self.f=fixtures.PersonnelTests()
        self.f.setUp()
        self.owner=Owner(self.f)
        self.builders=[]

    def tearDown(self):
        for builder in self.builders:builder.deleteLater()
        self.owner.deleteLater()
        self.app.processEvents()
        self.f.tearDown()

    def builder(self,base=None,inline=False):
        builder=CrewPlanBuilder(self.owner,base,inline)
        self.builders.append(builder)
        return builder

    def test_f01_multisheet_mapping_and_idempotent_commit(self):
        book=Workbook()
        for sheet,name in ((book.active,'Alpha'),(book.create_sheet(),'Bravo')):
            sheet.title=name
            sheet.append(['Name','Position','From','To','Date','Comment'])
            sheet.append([name,'Engineer','xxPort','CHA','2026-10-06','Second mapping'])
        path=self.f.root/'multiple.xlsx'
        book.save(path);book.close()
        first,second=[read_personnel_excel(path,name) for name in ('Alpha','Bravo')]
        preview=self.f.service.import_preview(first,self.f.snapshot,self.f.plan,crew_list=True)
        self.f.service.commit_import(preview)
        with self.assertRaises(PersonnelError):self.f.service.commit_import(preview)
        self.f.service.commit_import(self.f.service.import_preview(second,self.f.snapshot,self.f.plan,crew_list=True))
        with self.assertRaises(PersonnelError):self.f.service.import_preview(first,self.f.snapshot,self.f.plan)
        mapped=read_personnel_excel(path,'Alpha',mapping={'name':1,'position':2,'from_name':3,'to_name':4,'date':5},header_row=1)
        self.f.service.commit_import(self.f.service.import_preview(mapped,self.f.snapshot,self.f.plan))
        self.assertEqual(3,len(self.f.service.passengers(self.f.plan)))

    def test_f02_invalid_repeated_assignment_requires_review_and_new_choice(self):
        identifier=self.f.passenger(source='CHA',target='xxPort')
        last=self.f.command('add_stop',{'vessel_id':self.f.a})['stop_id']
        source=read_cargo(self.f.cargo_path)
        source['stops'].append({'id':'PORT','vessel_id':self.f.shore,'vessel_name_snapshot':'xxPort'})
        draft=generate_plan(source,self.f.service.crew_records(self.f.plan,self.f.snapshot),[],
                            '2026-10-06','xxPort',mode='Final',personnel_plan_id=self.f.plan,
                            assignments={'Passengers:'+identifier:{'pick':last,'drop':'PORT'}})
        builder=self.builder(draft,inline=True)
        index=next(i for i in range(builder.route.count()) if builder.route.item(i).data(256)['id']==last)
        builder.route.takeItem(index)
        builder.refresh_assignments()
        changed=builder.build()
        self.assertEqual(last,builder.assignments[0].currentData()['pick'])
        self.assertTrue(any('no longer valid' in p['message'] for p in changed['issues']))
        with self.assertRaises(PersonnelError):self.f.service.save_berth(changed,True)
        builder.assignments[0].setCurrentIndex(1)
        valid=builder.build()
        self.assertFalse(valid['issues'])
        self.assertEqual(self.f.s1,plan_assignments(valid)['Passengers:'+identifier]['pick'])

    def test_f01_replayed_new_list_submission_cannot_create_another_list(self):
        parsed=read_personnel_excel(self.f.workbook())
        preview=self.f.service.import_preview(parsed,self.f.snapshot,None,mode='new',crew_list=True)
        saved=self.f.service.commit_import(preview,'New List','2026-10-06')
        before=len(self.f.service.plans())
        with self.assertRaises(PersonnelError):self.f.service.commit_import(preview,'New List','2026-10-06')
        self.assertEqual(before,len(self.f.service.plans()))
        fresh=self.f.service.import_preview(parsed,self.f.snapshot,None,mode='new',crew_list=True)
        other=self.f.service.commit_import(fresh,'Another deliberate list','2026-10-06')
        self.assertNotEqual(saved['plan_id'],other['plan_id'])

    def test_f02_inline_material_assignment_cannot_silently_fallback(self):
        self.f.passenger()
        number=self.f.command('create_transfer',{'description':'Box','from_location_id':self.f.a,'to_location_id':self.f.b})['no']
        last=self.f.command('add_stop',{'vessel_id':self.f.a})['stop_id']
        draft=generate_plan(read_cargo(self.f.cargo_path),self.f.service.crew_records(self.f.plan,self.f.snapshot),[number],
                            '2026-10-06','xxPort',mode='Final',personnel_plan_id=self.f.plan,
                            assignments={'Materials:'+str(number):{'pick':last,'drop':None}})
        # A collection-only route is valid when its destination is not called.
        draft['stops']=[s for s in draft['stops'] if s['vessel_id']!=self.f.b]
        draft['assignments']={'Materials:'+str(number):{'pick':last,'drop':None}}
        builder=self.builder(draft,inline=True)
        index=next(i for i in range(builder.route.count()) if builder.route.item(i).data(256)['id']==last)
        builder.route.takeItem(index);builder.refresh_assignments()
        self.assertTrue(any(p['task']=='Materials:'+str(number) for p in builder.build()['issues']))

    def test_f03_automatic_omissions_and_user_exclusions_are_distinct(self):
        self.f.passenger(name='Matched')
        bad=self.f.service.add_passenger(self.f.plan,self.f.service.bind_record({'name':'Unmatched','position':'Engineer',
             'date':'2026-10-06','direction':'Sign-on','vessel':'Unknown'},self.f.snapshot),validate=False)
        builder=self.builder()
        draft=builder.build()
        self.assertTrue(any(p['task']=='Passengers:'+bad for p in draft['issues']))
        self.assertEqual('unavailable',draft['task_selections']['Passengers:'+bad]['state'])
        self.assertEqual(2,len(draft['personnel_versions']))
        with self.assertRaises(PersonnelError):self.f.service.save_berth(draft,True)
        row=next(i for i,t in enumerate(builder.tasks) if t['id']==bad)
        builder.rows.selectRow(row);builder.remove_from_plan()
        valid=builder.build()
        self.assertFalse(valid['issues'])
        self.assertEqual('user_excluded',valid['task_selections']['Passengers:'+bad]['state'])
        saved=self.f.service.save_berth(valid,True)
        self.assertIn('Unmatched',plan_text(saved))
        path=self.f.root/'exclusions.xlsx';export_plan(saved,path)
        book=load_workbook(path)
        try:self.assertIn('user_excluded',' '.join(str(c.value) for s in book for row in s for c in row))
        finally:book.close()
        reopened=self.builder(saved)
        self.assertEqual('user_excluded',reopened.build()['task_selections']['Passengers:'+bad]['state'])

    def test_f04_restore_invalidates_builder_and_precomputed_draft(self):
        self.f.passenger(name='Retained')
        second=self.f.passenger(name='Restored')
        backup=self.f.root/'two.sqlite3';self.f.service.backup(backup)
        self.f.service.remove_crew_entries({second:self.f.service.passenger(second)['version']})
        builder=self.builder(inline=True);draft=builder.build()
        self.f.service.restore(backup)
        self.assertEqual(2,len(self.f.service.crew_records(self.f.plan,self.f.snapshot)))
        with self.assertRaises(PersonnelError):builder.build()
        with self.assertRaises(PersonnelError):self.f.service.save_berth(draft,True)
        with self.assertRaises(PersonnelError):self.f.service.save_berth(draft,False)
        legacy=copy.deepcopy(draft);legacy.pop('personnel_context')
        with self.assertRaisesRegex(PersonnelError,'scope changed'):self.f.service.save_berth(legacy,True)

    def test_f04_full_scope_detects_new_or_removed_unselected_people(self):
        self.f.passenger()
        builder=self.builder();draft=builder.build()
        new=self.f.passenger(name='Added later')
        with self.assertRaises(PersonnelError):builder.build()
        with self.assertRaises(PersonnelError):self.f.service.save_berth(draft,True)
        builder.reload_tasks()
        self.assertIn(new,builder.build()['personnel_versions'])

    def test_f05_both_modules_protect_closed_companion_and_exports(self):
        self.f.passenger()
        before=self.f.service.path.read_bytes()
        self.f.service.close()
        with self.assertRaises(AppError):self.f.cargo.create_backup(self.f.service.path)
        self.assertEqual(before,self.f.service.path.read_bytes())
        self.f.service=PersonnelService(self.f.service.path);self.owner.service=self.f.service
        cargo_before=self.f.cargo_path.read_bytes()
        with self.assertRaises(PersonnelError):self.f.service.backup(self.f.cargo_path)
        with self.assertRaises(AppError):self.f.cargo.export_to(self.f.service.path,'all',{},'csv')
        self.assertEqual(cargo_before,self.f.cargo_path.read_bytes())

    def test_f05_samefile_identity_guard_rejects_differently_named_alias(self):
        alias=self.f.root/'alias.db'
        with patch('cargo_track.coordination.os.path.samefile',return_value=True):
            with self.assertRaises(PersonnelError):self.f.service.backup(alias)

    def test_f05_export_raw_database_path_is_rejected_before_extension_change(self):
        self.f.passenger()
        saved=self.f.service.save_berth(self.builder().build(),True)
        before=self.f.service.path.read_bytes()
        with self.assertRaises(PersonnelError):export_plan(saved,self.f.service.path)
        self.assertEqual(before,self.f.service.path.read_bytes())

    def test_f06_field_overrides_clear_only_when_master_field_changes(self):
        identifier=self.f.passenger(name='Original')
        builder=self.builder(inline=True)
        builder.rows.item(0,2).setText('Transient')
        builder.rows.item(0,3).setText('Temporary position')
        old=self.f.service.passenger(identifier)
        route=self.f.record(name='Original',target='Niger')
        self.f.service.edit_crew_entry(identifier,route,old['version'])
        builder.reload_tasks()
        self.assertEqual('Transient',builder.tasks[0]['name'])
        route['name']='Permanent'
        self.f.service.edit_crew_entry(identifier,route,self.f.service.passenger(identifier)['version'])
        builder.reload_tasks()
        self.assertEqual('Permanent',builder.tasks[0]['name'])
        self.assertEqual('Temporary position',builder.tasks[0]['position'])
        self.assertNotIn('name',builder.local_edits['Passengers:'+identifier])

    def test_f07_missing_objects_and_malformed_payloads_leave_live_data_unchanged(self):
        self.f.passenger()
        before=self.f.service.path.read_bytes()
        for sql in ('DROP TABLE module_settings','DROP TRIGGER history_no_update',
                    "DELETE FROM module_settings WHERE key='port_name'",
                    "UPDATE passengers SET payload='{\"issues\":{},\"name\":\"Bad\"}'"):
            candidate=self.f.root/(hashlib.sha256(sql.encode()).hexdigest()+'.sqlite3')
            self.f.service.backup(candidate)
            conn=sqlite3.connect(candidate);conn.execute(sql);conn.commit();conn.close()
            with self.assertRaises(PersonnelError):self.f.service.restore(candidate)
            self.assertEqual(before,self.f.service.path.read_bytes())

    def test_f07_legacy_backup_migrates_without_changing_history(self):
        self.f.passenger()
        candidate=self.f.root/'legacy.sqlite3';self.f.service.backup(candidate)
        conn=sqlite3.connect(candidate);conn.execute('DROP TABLE import_keys');conn.commit();conn.close()
        self.f.service.restore(candidate)
        self.assertEqual(1,len(self.f.service.passengers(self.f.plan)))
        self.f.service.validate_file(self.f.service.path)
        self.assertTrue(self.f.service.db.execute("SELECT 1 FROM sqlite_master WHERE name='import_keys'").fetchone())

    def test_f07_double_restore_failure_enters_maintenance_with_safety_copy(self):
        self.f.passenger();candidate=self.f.root/'valid.sqlite3';self.f.service.backup(candidate)
        record=self.f.record()
        real=self.f.service.validate_file
        def fail_live(path):
            if Path(path)==self.f.service.path:raise PersonnelError('Injected live verification failure')
            return real(path)
        with patch.object(self.f.service,'validate_file',side_effect=fail_live):
            with self.assertRaisesRegex(PersonnelError,'Writes are blocked'):self.f.service.restore(candidate)
        self.assertTrue(self.f.service.maintenance)
        self.assertTrue(Path(self.f.service.recovery_files['safety_copy']).is_file())
        self.assertTrue(self.f.service.recovery_marker.is_file())
        with self.assertRaisesRegex(PersonnelError,'recovery is required'):PersonnelService(self.f.service.path)
        with self.assertRaises(PersonnelError):self.f.service.add_passenger(self.f.plan,record)

    def test_atomic_cargo_and_personnel_settings_roll_back_all_fields(self):
        before=self.f.cargo.workspace_snapshot()
        with self.assertRaises(AppError):self.f.command('update_settings',{'values':{'business_timezone':'UTC','transfer_vessel_name':'Changed','port_name':'CHA'}})
        self.assertEqual(before,self.f.cargo.workspace_snapshot())
        old=self.f.service.settings()
        with self.assertRaises(PersonnelError):self.f.service.update_settings({'transfer_vessel_name':'Changed','unknown':'Invalid'})
        self.assertEqual(old,self.f.service.settings())

    def test_cargo_restore_rejects_missing_immutable_trigger_and_wrong_module_open_is_readonly(self):
        candidate=self.f.root/'cargo-broken.sqlite3';self.f.cargo.create_backup(candidate)
        conn=sqlite3.connect(candidate);conn.execute('DROP TRIGGER audit_events_no_delete');conn.commit();conn.close()
        before=self.f.cargo_path.read_bytes()
        with self.assertRaises(AppError):self.f.cargo.restore_backup(candidate)
        self.assertEqual(before,self.f.cargo_path.read_bytes())
        from cargo_track.service import ApplicationService
        personnel_before=self.f.service.path.read_bytes()
        wrong=ApplicationService(self.f.service.path)
        with self.assertRaises(RuntimeError):wrong.open()
        self.assertEqual(personnel_before,self.f.service.path.read_bytes())

    def test_legacy_alias_signatures_and_uuid_aliases_survive_storage_move(self):
        location=next(r for r in self.f.snapshot['locations'] if r['id']==self.f.a)
        from cargo_track.personnel_cargo import location_signature
        self.f.service.db.execute('INSERT INTO vessel_aliases VALUES(?,?,?,?)',
            (self.f.snapshot['source_path'],'legacy code',self.f.a,location_signature(location)))
        self.assertEqual(self.f.a,self.f.service.resolve_location(self.f.snapshot,'legacy code')['id'])
        self.f.service.set_alias(self.f.snapshot,'new code',self.f.a)
        moved=self.f.cargo.relocated_service(self.f.root/'new-source'/'cargo-track.sqlite3')
        try:self.assertEqual(self.f.a,self.f.service.resolve_location(read_cargo(moved.path),'new code')['id'])
        finally:moved.close()

    def test_dependency_signature_ignores_unrelated_cargo_but_detects_selected_changes(self):
        self.f.passenger()
        selected=self.f.command('create_transfer',{'description':'Selected','from_location_id':self.f.a,'to_location_id':self.f.b})['no']
        builder=self.builder()
        draft=builder.build()
        self.f.command('create_transfer',{'description':'Unrelated','from_location_id':self.f.b,'to_location_id':self.f.c})
        self.f.service.save_berth(draft,True)
        fresh=builder.build()
        self.f.command('edit_transfer',{'no':selected,'expected_version':0,'description':'Changed','comment':''})
        with self.assertRaises(PersonnelError):self.f.service.save_berth(fresh,True)

    def test_permanent_cargo_identity_survives_relocation_and_restore_generation_changes(self):
        source=read_cargo(self.f.cargo_path)
        backup=self.f.root/'cargo-backup.sqlite3';self.f.cargo.create_backup(backup)
        replacement=self.f.cargo.relocated_service(self.f.root/'relocated'/'cargo-track.sqlite3')
        try:self.assertEqual(source['ledger_identity'],read_cargo(replacement.path)['ledger_identity'])
        finally:replacement.close()
        generation=source['source_generation']
        self.f.cargo.restore_backup(backup)
        self.assertNotEqual(generation,read_cargo(self.f.cargo_path)['source_generation'])

    def test_streaming_import_preserves_rows_after_header_preview_and_bounds_mapping(self):
        book=Workbook();sheet=book.active
        for _ in range(45):sheet.append(['Intro'])
        sheet.append(['Name','Position','From','To','Date'])
        for i in range(150):sheet.append([f'Person {i}','Engineer','xxPort','CHA','2026-10-06'])
        path=self.f.root/'long.xlsx';book.save(path);book.close()
        parsed=read_personnel_excel(path,mapping={'name':1,'position':2,'from_name':3,'to_name':4,'date':5},header_row=46)
        self.assertEqual(150,len(parsed['records']))
        self.assertEqual(47,parsed['records'][0]['source_row'])
        with self.assertRaises(PersonnelError):read_personnel_excel(path,mapping={'name':201},header_row=46)


if __name__=='__main__':unittest.main()
