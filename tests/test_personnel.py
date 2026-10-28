from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from cargo_track.domain import Command
from cargo_track.service import ApplicationService
from cargo_track.personnel_cargo import read_cargo, location_signature
from cargo_track.personnel_import import PersonnelError, read_personnel_excel
from cargo_track.personnel_service import PersonnelService
from cargo_track.berth_plan import (generate_plan, revise_complete, export_plan, plan_text, route_options, task_key,
                                  plan_review_groups, plan_task_selection, plan_assignments)


class PersonnelTests(unittest.TestCase):
    def test_final_material_pickup_does_not_require_destination_call(self):
        destinations = [self.command("create_vessel", {"name": name})["id"] for name in ("HB6", "BGP Victor")]
        numbers = [self.command("create_transfer", {"description": description, "from_location_id": self.a,
            "to_location_id": target, "physical_id": "", "comment": "Keep"})["no"]
            for description, target in zip(("220V-24v电源转换器", "GPS电缆"), destinations)]
        before = self.cargo_path.read_bytes()
        snapshot = read_cargo(self.cargo_path)
        plan = self.generate(personnel=[], cargo=numbers, mode="Final")
        self.assertEqual([], plan["issues"])
        self.assertEqual(2, len(plan["stops"][0]["pick_materials"]))
        self.assertFalse(any(s["drop_materials"] for s in plan["stops"]))
        self.assertEqual(["CHA", "Niger", "INN"], [s["name"] for s in plan["stops"]])
        self.assertIn("220V-24v电源转换器 to HB6", plan_text(plan))
        self.assertIn("GPS电缆 to BGP Victor", plan_text(plan))
        saved = self.service.save_berth(plan, published=True)
        assignments = plan_assignments(saved)
        reopened = self.generate(personnel=[], cargo=numbers, mode="Final", assignments=assignments)
        self.assertEqual([], reopened["issues"])
        self.assertEqual(saved["stops"], reopened["stops"])
        export = self.root / "pickup-only.xlsx"
        export_plan(saved, export)
        book = load_workbook(export, read_only=True)
        try:
            text = " ".join(str(c.value) for row in book["Berthing Plan"] for c in row if c.value is not None)
            self.assertIn("Pick up 2 materials", text)
            self.assertIn("GPS电缆 to BGP Victor", text)
        finally:
            book.close()
        self.assertEqual(before, self.cargo_path.read_bytes())
        task = {"kind": "Materials", "from_id": self.a, "to_id": destinations[0], "status": "Waiting"}
        self.assertEqual([{ "pick": self.s1, "drop": None}], route_options(snapshot, task, "Final"))
        self.assertEqual([], route_options(snapshot, task, "Complete"))
        self.assertEqual([], route_options(snapshot, {**task, "kind": "Passengers"}, "Final"))
        self.assertEqual([], route_options(snapshot, {**task, "status": "On Board"}, "Final"))
        snapshot["stops"].append({**snapshot["stops"][0], "id": "repeat"})
        self.assertEqual([{ "pick": self.s1, "drop": None}, {"pick": "repeat", "drop": None}],
                         route_options(snapshot, task, "Final"))
        snapshot["stops"].append({"id": "delivery", "vessel_id": destinations[0]})
        self.assertEqual([{ "pick": self.s1, "drop": "delivery"}, {"pick": "repeat", "drop": "delivery"}],
                         route_options(snapshot, task, "Final"))

    def test_import_title_case_preserves_source_and_other_fields(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Name", "Position", "From", "To", "Date", "Comment"])
        for name, position in (("ALEX MORGAN", "NAV QC"), ("jordan lEE", "GUNNER"),
                               ("ANNE-MARIE O'NEILL", "senior engineer")):
            sheet.append([name, position, "xxPort", "CHA", "2026-10-06", "Keep UPPER Comment"])
        path = self.root / "case.xlsx"
        book.save(path)
        book.close()
        before = path.read_bytes()
        parsed = read_personnel_excel(path)
        self.assertEqual([("Alex Morgan", "NAV QC"), ("Jordan Lee", "Gunner"),
                          ("Anne-Marie O'Neill", "Senior Engineer")],
                         [(r["name"], r["position"]) for r in parsed["records"]])
        self.assertEqual("ALEX MORGAN", parsed["records"][0]["original_values"]["name"])
        self.assertEqual("NAV QC", parsed["records"][0]["original_values"]["position"])
        mapped = read_personnel_excel(path, mapping={"name": 1, "position": 2, "from_name": 3,
                                     "to_name": 4, "date": 5, "comment": 6}, header_row=1)
        self.assertEqual(parsed["records"], mapped["records"])
        self.service.commit_import(self.service.import_preview(parsed, self.snapshot, self.plan))
        crew = self.service.crew_records(self.plan, self.snapshot)
        self.assertTrue(all(r["to_name"] == "CHA" and r["comment"] == "Keep UPPER Comment" for r in crew))
        export = self.root / "title-export.xlsx"
        saved = self.service.save_berth(self.generate(personnel=crew, mode="Final"))
        export_plan(saved, export)
        exported = load_workbook(export, read_only=True)
        try:
            cells = " ".join(str(c.value) for s in exported for row in s for c in row if c.value is not None)
            self.assertIn("Alex Morgan", cells)
            self.assertIn("NAV QC", cells)
            self.assertNotIn("ALEX MORGAN", cells)
        finally:
            exported.close()
        self.assertEqual(before, path.read_bytes())

    def test_existing_imports_use_title_case_without_rewriting_database(self):
        record = self.record(name="ALEX MORGAN")
        record.update(position="NAV QC", batch_id="legacy-import",
                      original_values={"name": "ALEX MORGAN", "position": "NAV QC"})
        identifier = self.service.add_passenger(self.plan, record)
        legacy = {k:v for k,v in self.service.passenger(identifier).items() if k not in {'text_case','id','plan_id','version','status','created_utc'}}
        legacy.update(name='ALEX MORGAN',position='NAV QC')
        self.service.db.execute('UPDATE passengers SET payload=? WHERE id=?',(json.dumps(legacy),identifier))
        before = self.service.path.read_bytes()
        crew = self.service.crew_records(self.plan, self.snapshot)
        self.assertEqual(("Alex Morgan", "NAV QC"), (crew[0]["name"], crew[0]["position"]))
        self.assertEqual("ALEX MORGAN", self.service.passenger(identifier)["name"])
        self.assertEqual(before, self.service.path.read_bytes())

    def test_final_route_review_names_the_missing_vessel(self):
        identifier = self.passenger(source="CHA", target="Niger")
        source = json.loads(json.dumps(self.snapshot))
        source["stops"] = [s for s in source["stops"] if s["vessel_id"] != self.b]
        plan = generate_plan(source, self.service.passengers(self.plan), [], "2026-10-06", "xxPort", mode="Final")
        self.assertEqual([{ "task": "Passengers:" + identifier, "message": "Add destination vessel Niger to this plan.", "name": "Person A"}], plan["issues"])

    def test_review_groups_keep_distinct_people_and_show_each_reason_once(self):
        saved = {"issues": [{"task": "Passengers:a", "name": "Same Name", "message": "Missing vessel"},
            {"task": "Passengers:a", "name": "Same Name", "message": "Missing date"},
            {"task": "Passengers:a", "name": "Same Name", "message": "Missing vessel"},
            {"task": "Passengers:b", "name": "Same Name", "message": "Outside route"}]}
        original = json.dumps(saved)
        groups = plan_review_groups(saved)
        self.assertEqual(2, len(groups))
        self.assertEqual(["Missing vessel", "Missing date"], groups[0]["messages"])
        self.assertEqual(original, json.dumps(saved))

    def test_legacy_selection_retains_unassigned_materials_but_explicit_removal_wins(self):
        saved = {"included": ["Passengers:a"], "issues": [{"task": "Materials:3"}, {"task": ""}]}
        self.assertEqual({"Passengers:a", "Materials:3"}, plan_task_selection(saved))
        saved["selected_tasks"] = ["Materials:3"]
        self.assertEqual({"Materials:3"}, plan_task_selection(saved))

    def test_assignments_recovered_from_saved_visits_when_not_explicit(self):
        person = {"kind": "Passengers", "id": "a"}
        material = {"kind": "Materials", "id": "3"}
        def stop(identifier, **groups):
            return {"id": identifier, **{g: groups.get(g, []) for g in
                ("pick_passengers", "drop_passengers", "pick_materials", "drop_materials")}}
        saved = {"stops": [stop("second-a", pick_passengers=[person], pick_materials=[material]),
                           stop("b", drop_materials=[material]), stop("PORT", drop_passengers=[person])]}
        self.assertEqual({"Passengers:a": {"pick": "second-a", "drop": "PORT"},
                          "Materials:3": {"pick": "second-a", "drop": "b"}}, plan_assignments(saved))

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cargo_path = self.root / "cargo.sqlite3"
        self.cargo = ApplicationService(self.cargo_path)
        self.cargo.open()
        self.seq = 0
        self.a = self.command("create_vessel", {"name": "CHA"})["id"]
        self.b = self.command("create_vessel", {"name": "Niger"})["id"]
        self.c = self.command("create_vessel", {"name": "INN"})["id"]
        self.command("create_voyage", {"name": "Test route"})
        self.s1 = self.command("add_stop", {"vessel_id": self.a})["stop_id"]
        self.s2 = self.command("add_stop", {"vessel_id": self.b})["stop_id"]
        self.s3 = self.command("add_stop", {"vessel_id": self.c})["stop_id"]
        self.snapshot = read_cargo(self.cargo_path)
        self.shore = next(r["id"] for r in self.snapshot["locations"] if r["kind"] == "SHORE")
        self.service = PersonnelService(self.root / "personnel.sqlite3")
        self.plan = self.service.create_plan("Crew change", "2026-10-06")

    def tearDown(self):
        self.service.close()
        self.cargo.close()
        self.temp.cleanup()

    def command(self, operation, payload):
        self.seq += 1
        return self.cargo.handle(Command(f"personnel-test-{self.seq}", operation, payload)).result

    def record(self, name="Person A", source="xxPort", target="CHA"):
        return self.service.bind_record({"name": name, "position": "Engineer", "nationality": "Test",
            "from_name": source, "to_name": target, "date": "2026-10-06", "direction": "Transfer", "comment": ""}, self.snapshot)

    def passenger(self, **kwargs):
        return self.service.add_passenger(self.plan, self.record(**kwargs))

    def workbook(self):
        book = Workbook()
        sheet = book.active
        sheet.title = "20261006-SF NIGER"
        sheet.merge_cells("A1:W5")
        sheet["A1"] = "Crew Change Plan"
        headers = ["SN", "Sign-off", "Position", "Nationality", "Passport No./ID No.", "IQAMA", "Vessel", "Company",
                   "Laptop Details", "", "Date of departure", "Mark", "SN", "Sign-on", "Position", "Nationality",
                   "Passport No./ID No.", "IQAMA", "Vessel", "Company", "Laptop Details", "", "Date of arrival", "Mark"]
        for c, label in enumerate(headers, 1):
            sheet.cell(6, c, label)
        sheet["I7"], sheet["J7"], sheet["U7"], sheet["V7"] = "Brand", "SR NO.", "Brand", "SR NO."
        for r, vessel in enumerate(["CHA", "Niger", "INN", "INN"], 8):
            for c, value in {14: f"Person {r}", 15: "Engineer", 16: "Test", 17: "SECRET-PASSPORT", 18: "SECRET-IQAMA",
                             19: vessel, 20: "SECRET-COMPANY", 21: "SECRET-LAPTOP", 22: "SECRET-SERIAL",
                             23: "2026-10-06", 24: "SECRET-MARK"}.items():
                sheet.cell(r, c, value)
        for c, value in {2: "Returning", 3: "QC", 4: "Test", 5: "SECRET-PASSPORT", 7: "CHA", 11: "2026-10-06"}.items():
            sheet.cell(8, c, value)
        sheet["M17"] = "   "
        path = self.root / "crew.xlsx"
        book.save(path)
        book.close()
        return path

    def imported(self):
        parsed = read_personnel_excel(self.workbook())
        preview = self.service.import_preview(parsed, self.snapshot, self.plan)
        return self.service.commit_import(preview), parsed

    def generate(self, personnel=None, cargo=None, **kwargs):
        return generate_plan(read_cargo(self.cargo_path), personnel if personnel is not None else self.service.passengers(self.plan),
                             cargo or [], "2026-10-06", self.service.settings()["port_name"], personnel_plan_id=self.plan, **kwargs)

    def test_double_column_import_ignores_private_data_and_preserves_source(self):
        path = self.workbook()
        before = path.read_bytes()
        parsed = read_personnel_excel(path)
        self.assertEqual(5, len(parsed["records"]))
        self.assertEqual(1, sum(r["direction"] == "Sign-off" for r in parsed["records"]))
        self.assertEqual([8, 8, 9, 10, 11], [r["source_row"] for r in parsed["records"]])
        self.assertNotIn("SECRET-", json.dumps(parsed))
        preview = self.service.import_preview(parsed, self.snapshot, self.plan)
        self.service.commit_import(preview)
        self.assertNotIn("SECRET-", json.dumps(self.service.passengers(self.plan)))
        self.assertEqual(before, path.read_bytes())

    def test_sample_plan_counts_and_terminal_only_port(self):
        self.imported()
        plan = self.generate()
        self.assertEqual([], plan["issues"])
        self.assertEqual(["CHA", "Niger", "INN", "xxPort"], [s["name"] for s in plan["stops"]])
        self.assertEqual([1, 1, 2, 1], [len(s["drop_passengers"]) for s in plan["stops"]])
        self.assertEqual(1, len(plan["stops"][0]["pick_passengers"]))
        self.assertFalse(plan["stops"][-1]["pick_passengers"])
        self.assertNotIn("departure", plan_text(plan).lower())

    def test_duplicate_file_and_renamed_copy_do_not_duplicate_tasks(self):
        result, parsed = self.imported()
        parsed["file_name"] = "renamed.xlsx"
        with self.assertRaises(PersonnelError):
            self.service.import_preview(parsed, self.snapshot, self.plan)
        self.assertEqual(5, len(self.service.passengers(self.plan)))

    def test_append_keeps_missing_tasks_and_complete_replace_protects_executed(self):
        first = self.passenger(name="Existing")
        on = self.passenger(name="Onboard")
        self.service.transition([on], "On Board")
        done = self.passenger(name="Delivered")
        self.service.transition([done], "On Board")
        self.service.transition([done], "Done")
        parsed = read_personnel_excel(self.workbook())
        append = self.service.import_preview(parsed, self.snapshot, self.plan, "append")
        self.service.commit_import(append)
        self.assertEqual("Waiting", self.service.passenger(first)["status"])
        parsed["file_hash"] = "new-file"
        replace = self.service.import_preview(parsed, self.snapshot, self.plan, "replace")
        result = self.service.commit_import(replace, cancel_missing=True)
        self.assertGreater(result["cancelled"], 0)
        self.assertEqual("Cancelled", self.service.passenger(first)["status"])
        self.assertEqual("On Board", self.service.passenger(on)["status"])
        self.assertEqual("Done", self.service.passenger(done)["status"])
        self.assertEqual(13, len(self.service.passengers(self.plan)))

    def test_same_name_requires_manual_match_and_never_uses_passport(self):
        self.passenger(name="Person 8")
        preview = self.service.import_preview(read_personnel_excel(self.workbook()), self.snapshot, self.plan)
        candidate = next(r for r in preview["items"] if r["data"]["name"] == "Person 8")
        self.assertIsNone(candidate["match_id"])
        self.assertEqual(1, len(candidate["candidates"]))

    def test_stale_import_is_atomic(self):
        parsed = read_personnel_excel(self.workbook())
        preview = self.service.import_preview(parsed, self.snapshot, self.plan)
        self.passenger()
        with self.assertRaises(PersonnelError):
            self.service.commit_import(preview)
        self.assertEqual(1, len(self.service.passengers(self.plan)))
        self.assertEqual(0, self.service.db.execute("SELECT count(*) FROM import_batches").fetchone()[0])

    def test_duplicate_row_matches_rollback_whole_batch(self):
        identifier = self.passenger()
        preview = self.service.import_preview(read_personnel_excel(self.workbook()), self.snapshot, self.plan)
        preview["items"][0]["match_id"] = identifier
        preview["items"][1]["match_id"] = identifier
        with self.assertRaises(PersonnelError):
            self.service.commit_import(preview)
        self.assertEqual(0, self.service.db.execute("SELECT count(*) FROM import_batches").fetchone()[0])

    def test_onboard_cannot_be_cancelled_even_via_correction(self):
        identifier = self.passenger()
        self.service.transition([identifier], "On Board")
        for correction in (False, True):
            with self.assertRaises(PersonnelError):
                self.service.transition([identifier], "Cancelled", correction=correction, reason="wrong")
        self.assertEqual("On Board", self.service.passenger(identifier)["status"])

    def test_destination_adjustment_keeps_onboard_and_preserves_history(self):
        identifier = self.passenger()
        self.service.transition([identifier], "On Board")
        old = self.service.passenger(identifier)
        changed = self.record(target="Niger")
        with self.assertRaises(PersonnelError):
            self.service.edit_passenger(identifier, changed, old["version"])
        self.service.edit_passenger(identifier, changed, old["version"], "Change rotation")
        self.assertEqual("On Board", self.service.passenger(identifier)["status"])
        self.assertEqual(self.b, self.service.passenger(identifier)["to_id"])
        self.assertIn("Destination changed", [e["event"] for e in self.service.history(identifier)])

    def test_batch_transition_rolls_back_if_one_task_is_invalid(self):
        one, two = self.passenger(name="One"), self.passenger(name="Two")
        self.service.transition([two], "On Board")
        with self.assertRaises(PersonnelError):
            self.service.transition([one, two], "On Board")
        self.assertEqual("Waiting", self.service.passenger(one)["status"])

    def test_correction_requires_reason_and_history_is_immutable(self):
        identifier = self.passenger()
        self.service.transition([identifier], "On Board")
        with self.assertRaises(PersonnelError):
            self.service.transition([identifier], "Waiting", correction=True)
        self.service.transition([identifier], "Waiting", correction=True, reason="Pickup was not performed")
        self.assertIsNone(self.service.passenger(identifier)["pickup_utc"])
        with self.assertRaises(Exception):
            self.service.db.execute("UPDATE personnel_history SET reason='changed'")

    def test_duplicate_stop_needs_specific_assignment(self):
        self.command("add_stop", {"vessel_id": self.a})
        identifier = self.passenger()
        unresolved = self.generate()
        self.assertIn("Select a specific repeated stop", [e["message"] for e in unresolved["issues"]])
        last = read_cargo(self.cargo_path)["stops"][-1]["id"]
        plan = self.generate(assignments={f"Passengers:{identifier}": {"pick": None, "drop": last}})
        self.assertFalse(plan["issues"])
        self.assertFalse(plan["stops"][0]["drop_passengers"])
        self.assertEqual(1, len(plan["stops"][-2]["drop_passengers"]))

    def test_reverse_route_is_not_reordered(self):
        self.passenger(source="INN", target="CHA")
        plan = self.generate()
        self.assertTrue(plan["issues"])
        self.assertEqual([self.s1, self.s2, self.s3, "PORT"], [s["id"] for s in plan["stops"]])

    def test_remaining_does_not_use_a_passed_pickup_stop(self):
        self.passenger(source="CHA", target="INN")
        plan = self.generate(mode="Remaining", start_stop_id=self.s2)
        self.assertTrue(plan["issues"])
        self.assertEqual([2, 3, 4], [s["number"] for s in plan["stops"]])

    def test_remaining_waiting_port_origin_requires_pickup(self):
        identifier = self.passenger()
        waiting = self.generate(mode="Remaining", start_stop_id=self.s1)
        self.assertTrue(waiting["issues"])
        self.service.transition([identifier], "On Board")
        loaded = self.generate(mode="Remaining", start_stop_id=self.s1)
        self.assertFalse(loaded["issues"])
        self.assertEqual(1, len(loaded["stops"][0]["drop_passengers"]))

    def test_followup_port_pickup_never_delivers_to_earlier_ships(self):
        identifier = self.passenger()
        plan = self.generate(followups=[f"Passengers:{identifier}"])
        self.assertFalse(plan["issues"])
        self.assertEqual(1, len(plan["stops"][-1]["pick_passengers"]))
        self.assertEqual(0, sum(len(s["drop_passengers"]) for s in plan["stops"]))
        self.assertIn("Next voyage", plan_text(plan))

    def test_cargo_counts_are_ids_not_descriptions_or_display_lines(self):
        numbers = []
        for _ in range(2):
            numbers.append(self.command("create_transfer", {"description": "2 Boxes", "from_location_id": self.a,
                            "to_location_id": self.b, "physical_id": "", "comment": ""})["no"])
        plan = self.generate(personnel=[], cargo=numbers)
        self.assertEqual(2, len(plan["stops"][0]["pick_materials"]))
        self.assertEqual(2, len(plan["stops"][1]["drop_materials"]))
        self.assertIn("1. 2 Boxes to Niger; 2. 2 Boxes to Niger", plan_text(plan))

    def test_unnamed_populated_row_retained_and_replacement_does_not_cancel(self):
        first, second = self.passenger(name="Alice"), self.passenger(name="Bob")
        book = Workbook()
        sheet = book.active
        sheet.append(["Name", "Position", "From", "To", "Date"])
        sheet.append(["Alice", "Engineer", "xxPort", "CHA", "2026-10-06"])
        sheet.append([None, "Engineer", "xxPort", "CHA", "2026-10-06"])
        path = self.root / "partial.xlsx"
        book.save(path)
        book.close()
        parsed = read_personnel_excel(path)
        self.assertEqual(2, len(parsed["records"]))
        self.assertIn("Missing name", parsed["records"][1]["issues"])
        preview = self.service.import_preview(parsed, self.snapshot, self.plan, mode="replace")
        preview["items"][0]["match_id"] = first
        with self.assertRaises(PersonnelError):
            self.service.commit_import(preview, cancel_missing=True)
        self.assertEqual("Waiting", self.service.passenger(second)["status"])
        self.assertEqual(2, len(self.service.passengers(self.plan)))

    def test_cargo_pickup_requires_original_complete_version(self):
        number = self.command("create_transfer", {"description": "Box", "from_location_id": self.a,
            "to_location_id": self.b, "physical_id": "", "comment": ""})["no"]
        self.command("pick_up", {"no": number, "expected_version": 0, "confirmed_location": True})
        draft = self.generate(personnel=[], cargo=[number])
        self.assertTrue(any("already picked up" in issue["message"] for issue in draft["issues"]))
        with self.assertRaises(PersonnelError):
            self.service.save_berth(draft, published=True)
        remaining = self.generate(personnel=[], cargo=[number], mode="Remaining", start_stop_id=self.s1)
        self.assertFalse(remaining["issues"])
        self.assertEqual(1, len(remaining["stops"][1]["drop_materials"]))

    def test_cargo_correction_to_onboard_also_requires_original_complete(self):
        number = self.command("create_transfer", {"description": "Box", "from_location_id": self.a,
            "to_location_id": self.b, "physical_id": "", "comment": ""})["no"]
        self.command("correct_transfer", {"no": number, "expected_version": 0, "status": "On Board", "reason": "Pickup was recorded late"})
        self.assertTrue(any("already picked up" in issue["message"] for issue in self.generate(personnel=[], cargo=[number])["issues"]))

    def test_complete_revision_preserves_pickup_changes_destination_and_adds_task(self):
        identifier = self.passenger(source="CHA", target="Niger")
        saved = self.service.save_berth(self.generate(), published=True)
        self.service.transition([identifier], "On Board")
        current = self.service.passenger(identifier)
        self.service.edit_passenger(identifier, self.record(source="CHA", target="INN"), current["version"], "Rotation changed")
        extra = self.passenger(name="New person", source="Niger", target="INN")
        snapshot = read_cargo(self.cargo_path)
        live = self.service.passengers(self.plan)
        pending = self.generate(mode="Remaining", start_stop_id=self.s2)
        revised = revise_complete(saved, pending, snapshot, live)
        self.assertFalse(revised["issues"])
        self.assertEqual("Complete", revised["mode"])
        self.assertEqual([identifier], [i["id"] for i in revised["stops"][0]["pick_passengers"]])
        self.assertEqual([], revised["stops"][1]["drop_passengers"])
        self.assertEqual([extra], [i["id"] for i in revised["stops"][1]["pick_passengers"]])
        self.assertEqual({identifier, extra}, {i["id"] for i in revised["stops"][2]["drop_passengers"]})
        published = self.service.save_berth(revised, True, saved["family_id"])
        self.assertEqual(2, published["version"])
        self.assertEqual(saved, next(p for p in self.service.berth_versions() if p["id"] == saved["id"]))
        self.service.transition([extra], "On Board")
        with self.assertRaises(PersonnelError):
            self.service.save_berth(revised, True, saved["family_id"])

    def test_complete_revision_retains_done_actions_and_rejects_reordered_route(self):
        identifier = self.passenger(source="CHA", target="Niger")
        saved = self.service.save_berth(self.generate(), published=True)
        self.service.transition([identifier], "On Board")
        self.service.transition([identifier], "Done")
        pending = self.generate(mode="Remaining", start_stop_id=self.s3)
        revised = revise_complete(saved, pending, read_cargo(self.cargo_path), self.service.passengers(self.plan))
        self.assertEqual(saved["stops"], revised["stops"])
        changed = read_cargo(self.cargo_path)
        changed["stops"].reverse()
        with self.assertRaises(PersonnelError):
            revise_complete(saved, pending, changed, self.service.passengers(self.plan))

    def test_complete_version_preserves_finished_actions_and_original_port(self):
        identifier = self.passenger()
        saved = self.service.save_berth(self.generate(), published=True)
        self.service.transition([identifier], "On Board")
        self.service.transition([identifier], "Done")
        self.service.set_setting("port_name", "New Port")
        copied = generate_plan(read_cargo(self.cargo_path), [], [], saved["date"], "New Port", base_plan=saved)
        self.assertEqual("xxPort", copied["port_name"])
        self.assertEqual(saved["stops"], copied["stops"])
        remaining = self.generate(mode="Remaining", start_stop_id=self.s1)
        self.assertFalse(remaining["stops"][0]["drop_passengers"])
        self.assertEqual(saved, self.service.berth_versions()[0])

    def test_changed_personnel_and_cargo_block_publication(self):
        identifier = self.passenger()
        draft = self.generate()
        self.service.transition([identifier], "On Board")
        with self.assertRaises(PersonnelError):
            self.service.save_berth(draft, published=True)
        draft = self.generate()
        self.command("add_stop", {"vessel_id": self.a})
        with self.assertRaises(PersonnelError):
            self.service.save_berth(draft, published=True)
        self.assertTrue(self.service.save_berth(draft)["id"])

    def test_draft_contains_unassigned_issues_and_cannot_be_published(self):
        self.passenger(source="INN", target="CHA")
        draft = self.generate()
        with self.assertRaises(PersonnelError):
            self.service.save_berth(draft, published=True)
        saved = self.service.save_berth(draft)
        path = self.root / "draft.xlsx"
        export_plan(saved, path)
        book = load_workbook(path)
        self.assertIn("DRAFT", " ".join(str(c.value or "") for row in book.active for c in row))
        book.close()

    def test_excel_snapshot_counts_structure_dates_and_literal_text(self):
        self.passenger(name="=1+1")
        numbers = [self.command("create_transfer", {"description": desc, "from_location_id": self.a,
                  "to_location_id": self.b, "physical_id": "", "comment": ""})["no"] for desc in ("Material A", "Material B")]
        saved = self.service.save_berth(self.generate(cargo=numbers), published=True)
        path = self.root / "20261006_Fleet Personnel Transfer Plan.xlsx"
        export_plan(saved, path)
        book = load_workbook(path)
        self.assertEqual(["Berthing Plan", "Transfer Details", "Transfer Print"], book.sheetnames)
        body = [str(cell.value or '') for row in book.active for cell in row]
        self.assertIn("Pick up 2 materials: 1. Material A to Niger; 2. Material B to Niger", body)
        self.assertFalse(any("departure" in value.casefold() for value in body))
        self.assertEqual(6, book["Transfer Details"].max_row)
        self.assertEqual("E2", book["Transfer Details"].freeze_panes)
        self.assertTrue(book["Transfer Details"].auto_filter.ref)
        self.assertFalse(any(c.data_type == "f" for sheet in book for row in sheet for c in row))
        self.assertTrue(any(c.is_date for row in book["Transfer Details"] for c in row))
        self.assertEqual(saved, self.service.berth_versions()[0])
        book.close()

    def test_readonly_cargo_integration_leaves_database_bytes_unchanged(self):
        before = self.cargo_path.read_bytes()
        self.imported()
        saved = self.service.save_berth(self.generate(), published=True)
        export_plan(saved, self.root / "plan.xlsx")
        self.service.backup(self.root / "backup.sqlite3")
        self.assertEqual(before, self.cargo_path.read_bytes())

    def test_unknown_alias_and_database_change_require_remapping(self):
        record = self.service.bind_record({"name": "One", "position": "QC", "date": "2026-10-06",
            "direction": "Sign-on", "vessel": "Charlie"}, self.snapshot)
        self.assertTrue(record["issues"])
        self.service.set_alias(self.snapshot, "Charlie", self.a)
        resolved = self.service.bind_record(record, self.snapshot)
        self.assertFalse(resolved["issues"])
        changed = dict(self.snapshot, source_path="different.sqlite3")
        self.assertTrue(self.service.route_issues(resolved, changed))
        changed["locations"] = [dict(r, name="Different") if r["id"] == self.a else r for r in changed["locations"]]
        self.assertTrue(self.service.route_issues(resolved, changed))

    def test_standard_single_table_and_manual_mapping(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Name", "Position", "From", "To", "Date", "Comment", "Passport"])
        sheet.append(["One", "QC", "CHA", "Niger", "2026-10-06", "Keep", "SECRET"])
        path = self.root / "standard.xlsx"
        book.save(path)
        book.close()
        parsed = read_personnel_excel(path)
        self.assertEqual("CHA", parsed["records"][0]["from_name"])
        self.assertNotIn("SECRET", json.dumps(parsed))
        mapped = read_personnel_excel(path, mapping={"name": 1, "position": 2, "from_name": 3, "to_name": 4, "date": 5}, header_row=1)
        self.assertEqual("One", mapped["records"][0]["name"])

    def test_bilingual_standard_headers(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["姓名\nName", "职位\nPosition", "From Vessel", "To Vessel", "日期\nDate"])
        sheet.append(["One", "QC", "CHA", "Niger", "2026-10-06"])
        path = self.root / "bilingual.xlsx"
        book.save(path)
        book.close()
        parsed = read_personnel_excel(path)
        self.assertEqual("CHA", parsed["records"][0]["from_name"])
        self.assertEqual("Niger", parsed["records"][0]["to_name"])

    def test_bilingual_departure_date_is_not_a_second_name_column(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["下船\nSign-off", "职位\nPosition", "船名\nVessel", "下船时间\nDate of departure",
                      "上船\nSign-on", "职位\nPosition", "船名\nVessel", "到船日期\nDate of arrival"])
        sheet.append(["Off", "QC", "CHA", "2026-10-06", "On", "Engineer", "INN", "2026-10-06"])
        path = self.root / "bilingual-crew.xlsx"
        book.save(path)
        book.close()
        parsed = read_personnel_excel(path)
        self.assertEqual(2, len(parsed["records"]))
        self.assertEqual(["2026-10-06"] * 2, [r["date"] for r in parsed["records"]])
        self.assertEqual("D2", parsed["records"][0]["source_cells"]["date"])

    def test_unambiguous_vessel_abbreviations_and_ambiguous_match(self):
        snapshot = dict(self.snapshot)
        names = {self.a: "BGP Challenger", self.b: "Stanford Niger", self.c: "BGP Innovator"}
        snapshot["locations"] = [dict(r, name=names.get(r["id"], r["name"])) for r in snapshot["locations"]]
        for label, identifier in (("CHA", self.a), ("Niger", self.b), ("INN", self.c)):
            self.assertEqual(identifier, self.service.resolve_location(snapshot, label)["id"])
        snapshot["locations"].append(dict(snapshot["locations"][0], id=999, name="BGP Chariot", kind="VESSEL"))
        self.assertIsNone(self.service.resolve_location(snapshot, "CHA"))

    def test_final_crew_grouping_does_not_require_execution_status_or_matching_date(self):
        sign_on = self.passenger(name="On person")
        sign_off = self.passenger(name="Off person", source="CHA", target="xxPort")
        self.service.transition([sign_off], "On Board")
        self.service.transition([sign_off], "Done")
        crew = self.service.crew_records(self.plan, self.snapshot)
        crew[0]["date"] = "2026-10-07"
        crew[1]["date"] = ""
        crew[1]["issues"] = ["Missing date"]
        source = read_cargo(self.cargo_path)
        source['stops'].insert(0, {'id': 'plan:departure-port', 'vessel_id': self.shore, 'vessel_name_snapshot': 'xxPort'})
        source['stops'].append({'id': 'plan:port', 'vessel_id': self.shore, 'vessel_name_snapshot': 'xxPort'})
        plan = generate_plan(source, crew, [], '2026-10-06', 'xxPort', mode='Final')
        self.assertFalse(plan["issues"])
        self.assertFalse(plan["warnings"])
        self.assertEqual([sign_on], [r["id"] for r in plan["stops"][0]["pick_passengers"]])
        self.assertEqual([sign_on], [r["id"] for r in plan["stops"][1]["drop_passengers"]])
        self.assertEqual([sign_off], [r["id"] for r in plan["stops"][1]["pick_passengers"]])
        self.assertEqual([sign_off], [r["id"] for r in plan["stops"][-1]["drop_passengers"]])
        self.assertEqual("Done", self.service.passenger(sign_off)["status"])

    def test_final_plan_has_independent_added_and_reordered_stops(self):
        identifier = self.passenger(target="Niger")
        source = read_cargo(self.cargo_path)
        before = source["source_hash"]
        source["stops"] = [source["stops"][2], source["stops"][0], dict(source["stops"][1], id="plan:added")]
        plan = generate_plan(source, self.service.crew_records(self.plan, source), [], "2026-10-06", "xxPort", mode="Final")
        self.assertFalse(plan["issues"])
        self.assertEqual(["INN", "CHA", "Niger"], [s["name"] for s in plan["stops"]])
        self.assertEqual(identifier, plan["stops"][2]["drop_passengers"][0]["id"])
        self.service.save_berth(plan, True)
        self.assertEqual(before, read_cargo(self.cargo_path)["source_hash"])

    def test_same_file_can_replace_crew_list_without_duplicate_active_entries(self):
        result, parsed = self.imported()
        old = self.service.passengers(self.plan)
        self.service.transition([old[0]["id"]], "On Board")
        self.service.transition([old[0]["id"]], "Done")
        preview = self.service.import_preview(parsed, self.snapshot, self.plan, mode="replace", crew_list=True)
        result = self.service.commit_import(preview, plan_date="2026-10-06", cancel_missing=True)
        self.assertEqual(5, result["cancelled"])
        self.assertEqual(5, len(self.service.crew_records(self.plan, self.snapshot)))
        self.assertEqual("2026-10-06", next(p for p in self.service.plans() if p["id"] == self.plan)["date"])
        self.assertTrue(all(not item["candidates"] for item in preview["items"]))

    def test_final_repeated_vessel_defaults_to_first_feasible_stop_and_can_be_changed(self):
        identifier = self.passenger()
        again = self.command("add_stop", {"vessel_id": self.a})["stop_id"]
        crew = self.service.crew_records(self.plan, read_cargo(self.cargo_path))
        first = self.generate(personnel=crew, mode="Final")
        self.assertFalse(first["issues"])
        self.assertEqual(1, len(first["stops"][0]["drop_passengers"]))
        moved = self.generate(personnel=crew, mode="Final", assignments={f"Passengers:{identifier}": {"pick": None, "drop": again}})
        self.assertEqual(1, len(moved["stops"][-1]["drop_passengers"]))

    def test_ignored_headers_never_become_name_or_date_even_in_manual_mapping(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["Name", "Company Name", "Position", "From", "To", "Date", "Passport Expiry Date", "Laptop Name"])
        sheet.append(["Person", "PRIVATE-COMPANY", "QC", "CHA", "Niger", "2026-10-06", "2030-01-01", "PRIVATE-LAPTOP"])
        path = self.root / "ignored-labels.xlsx"
        book.save(path)
        book.close()
        parsed = read_personnel_excel(path)
        self.assertEqual("Person", parsed["records"][0]["name"])
        self.assertEqual("2026-10-06", parsed["records"][0]["date"])
        self.assertNotIn("PRIVATE", json.dumps(parsed))
        self.assertNotIn("2030-01-01", json.dumps(parsed))
        with self.assertRaises(PersonnelError):
            read_personnel_excel(path, mapping={"name": 2, "position": 3, "from_name": 4, "to_name": 5, "date": 6}, header_row=1)

    def test_backup_restore_and_relocation_are_independent(self):
        identifier = self.passenger()
        backup = self.root / "backup.sqlite3"
        self.service.backup(backup)
        self.service.transition([identifier], "On Board")
        safety = self.service.restore(backup)
        self.assertTrue(Path(safety).is_file())
        self.assertEqual("Waiting", self.service.passenger(identifier)["status"])
        with self.assertRaises(PersonnelError):
            self.service.restore(self.cargo_path)
        moved = self.service.relocate(self.root / "new" / "personnel.sqlite3")
        try:
            self.assertEqual(1, len(moved.passengers(self.plan)))
            self.assertTrue(self.service.path.is_file())
        finally:
            moved.close()
        with self.assertRaises(PersonnelError):
            self.service.relocate(self.root / "new" / "personnel.sqlite3")


if __name__ == "__main__":
    unittest.main()
