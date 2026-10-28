"""Regressions for removed-leg history, explicit linking and printable exports."""
import copy
import json
import os
import unittest
from openpyxl import load_workbook
from unittest.mock import patch
import test_personnel_routes as fixtures
from cargo_track.berth_plan import export_plan
from cargo_track.personnel_service import PersonnelService
from cargo_track.personnel_import import PersonnelError
from cargo_track.personnel_routes import suggested_leg_order
from cargo_track.excel_layout import split_text, text_height

os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
from PySide6.QtWidgets import QApplication, QDialogButtonBox
from cargo_track.personnel_ui import CrewLegOrderDialog


class AuditReleaseFixTests(unittest.TestCase):
    def setUp(self):
        self.f=fixtures.PersonnelRouteTests();self.f.setUp()
    def tearDown(self):self.f.tearDown()

    def test_removed_descendant_is_historical_after_parent_relinked(self):
        f=self.f;s=f.service
        root=f.add('INN','CHA');first=f.add('CHA','Niger')
        child=s.add_next_leg(first,f.record('Niger','INN'),f.source(),0)
        s.remove_crew_entries({child:0})
        before=s.passenger(child);history=copy.deepcopy(s.history(child))
        s.link_crew_legs([root,first],{root:0,first:0})
        self.assertEqual(before,s.passenger(child));self.assertEqual(history,s.history(child))
        backup=f.f.root/'relinked.sqlite3';s.backup(backup);PersonnelService.validate_file(backup)
        stored=s.path.read_bytes();reopened=PersonnelService(s.path)
        try:
            self.assertEqual(root,reopened.passenger(first)['previous_task_id'])
            self.assertEqual(before,reopened.passenger(child))
        finally:reopened.close()
        self.assertEqual(stored,s.path.read_bytes())
        # Re-activating a removed, outdated leg (for example by import matching)
        # must roll back rather than make the next application open fail.
        with self.assertRaises(PersonnelError):
            with s.atomic():s.db.execute("UPDATE passengers SET status='Waiting' WHERE id=?",(child,))
        self.assertEqual(before,s.passenger(child))
        s.backup(backup)

    def test_active_broken_chain_is_still_rejected(self):
        first,second,_=self.f.chain();s=self.f.service
        payload=json.loads(s.db.execute('SELECT payload FROM passengers WHERE id=?',(second,)).fetchone()[0])
        payload['person_id']='different-person'
        s.db.execute('UPDATE passengers SET payload=? WHERE id=?',(json.dumps(payload),second))
        with self.assertRaises(PersonnelError):s.backup(self.f.f.root/'invalid.sqlite3')
        with self.assertRaises(PersonnelError):PersonnelService(s.path)

    def test_split_text_keeps_all_characters_and_bounds_each_row(self):
        source=('1. Engineering equipment 测试电缆; '*100)+'TAIL'
        parts=split_text(source,60,180,'(continued) ')
        rebuilt=parts[0]+''.join(p.removeprefix('(continued) ') for p in parts[1:])
        self.assertEqual(source,rebuilt)
        self.assertTrue(all(text_height(p,60)<=180 and len(p)<=32767 for p in parts))

    def test_large_roster_and_comments_export_with_continuations(self):
        names=[f'Fictional Passenger {i:03d} Long English Full Name' for i in range(120)]
        for name in names:self.f.add('CHA','INN',name)
        plan=self.f.plan();self.assertEqual([],plan['issues'])
        comment=('Safety briefing and transfer notes 测试备注. '*100)+'COMMENT-END'
        plan['stops'][0]['comment']=comment
        saved=self.f.service.save_berth(plan,True);frozen=copy.deepcopy(saved)
        path=self.f.f.root/'long-plan.xlsx';export_plan(saved,path)
        book=load_workbook(path)
        try:
            sheet=book['Berthing Plan'];body=' '.join(str(c.value or '') for cells in sheet for c in cells)
            self.assertTrue(all(name in body for name in names))
            self.assertTrue(any('(continued)' in str(c.value) for c in sheet['B']))
            self.assertTrue(all((r.height or 15)<=409 for r in sheet.row_dimensions.values()))
            comments=''.join(str(c.value or '') for c in sheet['D'][1:])
            self.assertEqual(comment,comments)
            self.assertEqual(241,book['Transfer Details'].max_row)
            self.assertIn('Transfer Print',book.sheetnames)
            self.assertTrue(book['Transfer Print'].print_title_rows)
        finally:book.close()
        self.assertEqual(frozen,saved);self.assertEqual(frozen,self.f.service.berth_versions()[0])

    def test_cargo_long_comment_has_complete_print_continuations(self):
        f=self.f.f;comment=('Transfer note '*300)+'END'
        f.command('create_transfer',{'description':'Test equipment','from_location_id':f.a,'to_location_id':f.b,'comment':comment})
        path=f.root/'cargo.xlsx';f.cargo.export_to(path,'all',{'range':'all'},'xlsx')
        book=load_workbook(path)
        try:
            self.assertEqual(comment,book['Transfers']['G2'].value)
            self.assertEqual('Cargo Print',book.active.title)
            self.assertEqual(comment,''.join(str(c.value or '') for c in book['Cargo Print']['G'][1:]))
            self.assertTrue(all((r.height or 15)<=409 for r in book.active.row_dimensions.values()))
        finally:book.close()

    def test_cargo_print_settings_keep_all_nine_columns(self):
        f=self.f.f;f.command('create_transfer',{'description':'Test','from_location_id':f.a,'to_location_id':f.b})
        path=f.root/'cargo.xlsx';f.cargo.export_to(path,'all',{'range':'all'},'xlsx')
        book=load_workbook(path)
        try:
            s=book['Transfers'];self.assertEqual('landscape',s.page_setup.orientation)
            self.assertEqual(100,s.page_setup.scale);self.assertEqual('$1:$1',s.print_title_rows)
            self.assertIn('$I$2',str(s.print_area));self.assertEqual(9,s.max_column)
        finally:book.close()


class CrewLegOrderDialogTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):cls.app=QApplication.instance() or QApplication([])
    @staticmethod
    def task(identifier,a,b):return {'id':identifier,'name':'Fictional','from_id':a,'to_id':b,'from_name':str(a),'to_name':str(b)}

    def test_reverse_record_order_suggests_valid_transfer_order(self):
        tasks=[self.task('second',2,3),self.task('first',1,2)]
        dialog=CrewLegOrderDialog(tasks)
        try:
            self.assertEqual(['first','second'],[t['id'] for t in dialog.ordered_tasks()])
            self.assertTrue(dialog.validate_order())
            dialog.order.setCurrentRow(1);dialog.move(-1)
            self.assertFalse(dialog.validate_order())
            self.assertFalse(dialog.buttons.button(QDialogButtonBox.StandardButton.Ok).isEnabled())
            dialog.move(1);self.assertTrue(dialog.validate_order())
        finally:dialog.deleteLater();self.app.processEvents()

    def test_cycle_order_can_be_changed_explicitly(self):
        tasks=[self.task('out',1,2),self.task('back',2,1)]
        dialog=CrewLegOrderDialog(tasks)
        try:
            dialog.order.setCurrentRow(1);dialog.move(-1)
            self.assertEqual('back',dialog.ordered_tasks()[0]['id'])
            self.assertTrue(dialog.validate_order())
        finally:dialog.deleteLater();self.app.processEvents()

    def test_disconnected_records_cannot_be_confirmed(self):
        tasks=[self.task('a',1,2),self.task('b',3,4)]
        self.assertEqual(tasks,suggested_leg_order(tasks))
        dialog=CrewLegOrderDialog(tasks)
        try:self.assertFalse(dialog.validate_order())
        finally:dialog.deleteLater();self.app.processEvents()
