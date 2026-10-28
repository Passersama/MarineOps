"""Command handling, cargo routes and atomic file operations."""

from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import sqlite3
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from .domain import AppError, Command, Receipt, Status, business_date, normalized_name, now_utc
from .repository import LedgerRepository
from .schema import check_integrity, connect, initialize
from .coordination import workspace_files, ProtectedFileError
from .text_case import normalize

SORT_COLUMNS={0:"t.no",1:"c.cargo_tag",2:"t.description",3:"t.from_name_snapshot",4:"t.to_name_snapshot",5:"t.status",6:"t.comment",7:"CASE WHEN t.completed_date IS NULL THEN 1 ELSE 0 END, t.completed_date"}
MAX_DESCRIPTION=1000
MAX_COMMENT=5000
MAX_IDENTIFIER=250


def _json(value):
    return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",":"),default=str)


class ApplicationService:
    """Application service owned by exactly one database-worker thread."""
    def __init__(self,database_path):
        self.path=Path(database_path)
        workspace_files.register_database(self.path)
        self.path.parent.mkdir(parents=True,exist_ok=True)
        self.connection=None
        self.repository=None
        self.dataset_epoch=str(uuid4())

    def open(self):
        self.connection=connect(str(self.path))
        try:
            initialize(self.connection)
            failures=check_integrity(str(self.path))
            if failures:raise AppError("DATABASE_INVALID","The local ledger did not pass its integrity check.",{"failures":failures})
            self.repository=LedgerRepository(self.connection)
        except Exception:
            self.close()
            raise

    def close(self):
        if self.connection is not None:self.connection.close()
        self.connection=None
        self.repository=None

    @property
    def db(self):
        if self.connection is None:raise AppError("DATABASE_UNAVAILABLE","The local database is not open.")
        return self.connection

    @property
    def repo(self):
        if self.repository is None:raise AppError("DATABASE_UNAVAILABLE","The local database is not open.")
        return self.repository

    def _timezone(self):
        row=self.db.execute("SELECT value FROM settings WHERE key='business_timezone'").fetchone()
        return row[0] if row else "Asia/Riyadh"

    def _revision(self):return int(self.db.execute("SELECT data_revision FROM app_state WHERE id=1").fetchone()[0])

    def _audit(self,command_id,entity_type,entity_id,event_type,before=None,after=None,reason=""):
        self.db.execute("INSERT INTO audit_events(entity_type,entity_id,command_id,event_type,occurred_utc,timezone_name,reason,data_json) VALUES(?,?,?,?,?,?,?,?)",
            (entity_type,entity_id,command_id,event_type,now_utc(),self._timezone(),str(reason or "").strip(),_json({"before":before,"after":after})))

    @staticmethod
    def _required(value,label,maximum):
        result=str(value if value is not None else "").strip()
        if not result:raise AppError("VALIDATION",f"{label} is required.",{"field":label})
        if len(result)>maximum:raise AppError("VALIDATION",f"{label} must be {maximum} characters or fewer.",{"field":label,"maximum":maximum})
        return result

    @staticmethod
    def _optional(value,label,maximum):
        result=str(value if value is not None else "").strip()
        if len(result)>maximum:raise AppError("VALIDATION",f"{label} must be {maximum} characters or fewer.",{"field":label,"maximum":maximum})
        return result

    def _location(self,location_id,allow_inactive=False):
        try:identifier=int(location_id)
        except (ValueError,TypeError) as exc:raise AppError("VALIDATION","Choose a valid From and To location.") from exc
        row=self.db.execute("SELECT * FROM locations WHERE id=?",(identifier,)).fetchone()
        if row is None or row["kind"] not in ("VESSEL","SHORE") or (not row["active"] and not allow_inactive):
            raise AppError("LOCATION_UNAVAILABLE","Choose an active vessel or the configured port.")
        return row

    def _chain(self,cargo_id):return self.repo.cargo_legs(cargo_id,include_void=False)

    def _descendants(self,number):
        return self.db.execute("""WITH RECURSIVE d(no) AS (SELECT no FROM transfer_legs WHERE predecessor_no=? AND is_void=0
          UNION ALL SELECT t.no FROM transfer_legs t JOIN d ON t.predecessor_no=d.no WHERE t.is_void=0)
          SELECT t.* FROM transfer_legs t JOIN d ON d.no=t.no ORDER BY t.no""",(number,)).fetchall()

    def _expected(self,row,expected):
        try:version=int(expected)
        except (ValueError,TypeError) as exc:raise AppError("STALE_RECORD","Reload this record before saving changes.") from exc
        if row["version"]!=version:raise AppError("STALE_RECORD",f"Transfer #{row['no']} has changed. Reload it to review the latest values.",{"current_version":row["version"]})

    def handle(self,command:Command)->Receipt:
        if not command.command_id or len(command.command_id)>100:raise AppError("INVALID_COMMAND","This action has no valid command identifier.")
        digest=hashlib.sha256(_json(command.payload).encode("utf-8")).hexdigest()
        db=self.db
        try:
            db.execute("BEGIN IMMEDIATE")
            old=db.execute("SELECT * FROM command_receipts WHERE command_id=?",(command.command_id,)).fetchone()
            if old:
                db.rollback()
                if old["operation"]!=command.operation or old["request_hash"]!=digest:raise AppError("COMMAND_ID_REUSED","This action identifier was already used for a different request.")
                return Receipt(command.command_id,command.operation,json.loads(old["result_json"]),int(old["revision"]))
            function=getattr(self,"_cmd_"+command.operation,None)
            if function is None:raise AppError("UNKNOWN_COMMAND",f"Unknown application command: {command.operation}")
            result=function(command.command_id,command.payload)
            revision=self._revision()+1
            db.execute("UPDATE app_state SET data_revision=? WHERE id=1",(revision,))
            db.execute("INSERT INTO command_receipts VALUES(?,?,?,?,?,?)",(command.command_id,command.operation,digest,_json(result),revision,now_utc()))
            db.commit()
            return Receipt(command.command_id,command.operation,result,revision)
        except AppError:
            if db.in_transaction:db.rollback()
            raise
        except sqlite3.IntegrityError as exc:
            if db.in_transaction:db.rollback()
            error=str(exc)
            if "one_open_transfer_per_cargo" in error:raise AppError("ACTIVE_TRANSFER_EXISTS","This cargo already has an active transfer.") from exc
            if "locations.name_key" in error:raise AppError("DUPLICATE_VESSEL","A vessel with this name already exists.") from exc
            raise AppError("DATA_CONFLICT","This change conflicts with existing ledger data.",{"detail":error}) from exc
        except (sqlite3.Error,OSError) as exc:
            if db.in_transaction:db.rollback()
            raise AppError("SAVE_FAILED","The change could not be saved. Previous records are intact.",{"detail":str(exc)}) from exc
        except Exception:
            if db.in_transaction:db.rollback()
            raise

    def _description(self, p, previous=None):
        raw=self._required(p.get('description'), 'Description', MAX_DESCRIPTION)
        if p.get('preserve_text_case') or raw == previous:
            return raw
        return self._required(normalize(raw, 'cargo')[0], 'Description', MAX_DESCRIPTION)

    def _cmd_create_transfer(self,command_id,p):
        cargo_id=p.get('cargo_id')
        if cargo_id is not None:
            try:cargo_id=int(cargo_id)
            except (TypeError,ValueError) as exc:raise AppError('CARGO_NOT_FOUND','Choose an existing cargo record.') from exc
        previous = self.db.execute('SELECT default_description FROM cargo WHERE id=?', (cargo_id,)).fetchone()
        description=self._description(p, previous[0] if previous else None)
        comment=self._optional(p.get("comment"),"Comment",MAX_COMMENT)
        physical=self._optional(p.get("physical_id"),"Physical ID",MAX_IDENTIFIER)
        source=self._location(p.get("from_location_id"));dest=self._location(p.get("to_location_id"))
        predecessor=None
        if p.get("cargo_id") is None:
            now=now_utc()
            cur=self.db.execute("INSERT INTO cargo(default_description,physical_id,created_utc,updated_utc) VALUES(?,?,?,?)",(description,physical or None,now,now))
            cargo_id=cur.lastrowid
            tag=f"C-{cargo_id:06d}"
            self.db.execute("UPDATE cargo SET cargo_tag=? WHERE id=?",(tag,cargo_id))
            self._audit(command_id,"cargo",cargo_id,"CARGO_CREATED",after={"cargo_tag":tag,"description":description,"physical_id":physical})
        else:
            try:cargo_id=int(p["cargo_id"])
            except (ValueError,TypeError) as exc:raise AppError("CARGO_NOT_FOUND","Choose an existing cargo record.") from exc
            cargo=self.db.execute("SELECT * FROM cargo WHERE id=?",(cargo_id,)).fetchone()
            if cargo is None:raise AppError("CARGO_NOT_FOUND","This cargo record no longer exists.")
            if physical and cargo["physical_id"] and physical.casefold()!=cargo["physical_id"].casefold():raise AppError("PHYSICAL_ID_MISMATCH","The physical identifier differs from this cargo record.")
            chain=self._chain(cargo_id)
            if not chain or chain[-1]["status"]!="Done":raise AppError("PREDECESSOR_NOT_DONE","Deliver the previous transfer before creating its next segment.")
            predecessor=chain[-1]
            if predecessor["to_location_id"]!=source["id"] and not str(p.get("route_reason","")).strip():
                raise AppError("ROUTE_DISCONTINUITY",f"The previous transfer ended at {predecessor['to_name_snapshot']}. Add a reason for this pickup location.")
        now=now_utc()
        if predecessor and 'preserve_text_case' not in p and description==cargo['default_description']:
            p={**p,'preserve_text_case':self._preserved_description(predecessor['no'])}
        cur=self.db.execute("""INSERT INTO transfer_legs(cargo_id,predecessor_no,from_location_id,from_name_snapshot,to_location_id,
            to_name_snapshot,description,comment,status,created_utc,updated_utc) VALUES(?,?,?,?,?,?,?,?,'Waiting',?,?)""",
            (cargo_id,predecessor["no"] if predecessor else None,source["id"],source["name"],dest["id"],dest["name"],description,comment,now,now))
        no=cur.lastrowid
        self._audit(command_id,"transfer",no,"TRANSFER_CREATED",after={"cargo_id":cargo_id,"from":source["name"],"to":dest["name"],"description":description,"raw_description":p.get('description'),"preserve_text_case":bool(p.get('preserve_text_case')),"status":"Waiting"},reason=p.get("route_reason",""))
        return {"no":no,"cargo_id":cargo_id}

    def _active_leg(self,number):
        try:no=int(number)
        except (TypeError,ValueError) as exc:raise AppError("TRANSFER_NOT_FOUND","Select a valid transfer.") from exc
        row=self.repo.leg(no)
        if row is None:raise AppError("TRANSFER_NOT_FOUND","This transfer no longer exists.")
        if row["is_void"]:raise AppError("TRANSFER_VOIDED",f"Transfer #{no} is voided. Restore it before changing it.")
        return row

    def _cmd_edit_transfer(self,command_id,p):
        row=self._active_leg(p.get("no"));self._expected(row,p.get("expected_version"))
        p={**p,'preserve_text_case':p.get('preserve_text_case',self._preserved_description(row['no']))}
        description=self._description(p, row['description'])
        comment=self._optional(p.get("comment"),"Comment",MAX_COMMENT)
        physical=self._optional(p.get("physical_id"),"Physical ID",MAX_IDENTIFIER)
        if any(p.get(key) is not None and int(p[key])!=row[column] for key,column in (("from_location_id","from_location_id"),("to_location_id","to_location_id"))):
            raise AppError("CORRECTION_REQUIRED","Changing a route requires Correct record.")
        cargo=self.db.execute("SELECT * FROM cargo WHERE id=?",(row["cargo_id"],)).fetchone()
        before={"description":row["description"],"comment":row["comment"],"physical_id":cargo["physical_id"] or ""}
        after={"description":description,"comment":comment,"physical_id":physical}
        if before==after and bool(p.get('preserve_text_case'))==self._preserved_description(row['no']):return {"no":row["no"],"unchanged":True}
        now=now_utc()
        self.db.execute("UPDATE transfer_legs SET description=?,comment=?,version=version+1,updated_utc=? WHERE no=?",(description,comment,now,row["no"]))
        if before["physical_id"]!=physical:
            self.db.execute("UPDATE cargo SET physical_id=?,updated_utc=? WHERE id=?",(physical or None,now,cargo["id"]))
            self._audit(command_id,"cargo",cargo["id"],"PHYSICAL_ID_CHANGED",before=before["physical_id"],after=physical)
        if before['description'] != description:
            self.db.execute('UPDATE cargo SET default_description=?,updated_utc=? WHERE id=? AND default_description=?',
                            (description,now,row['cargo_id'],before['description']))
        self._audit(command_id,"transfer",row["no"],"TRANSFER_EDITED",before=before,after={**after,'raw_description':p.get('description'),'preserve_text_case':bool(p.get('preserve_text_case'))})
        return {"no":row["no"],"cargo_id":row["cargo_id"]}

    def _transition(self,command_id,p,target:Status,event):
        row=self._active_leg(p.get("no"));self._expected(row,p.get("expected_version"))
        source_status=Status.WAITING if target==Status.ON_BOARD else Status.ON_BOARD
        if row["status"]!=source_status.value:raise AppError("INVALID_TRANSITION",f"{event.title()} requires status {source_status.value}.")
        state=self.repo.current_state()
        endpoint=row["from_location_id"] if target==Status.ON_BOARD else row["to_location_id"]
        at_current=state["vessel_id"] is not None and state["vessel_id"]==endpoint
        if not at_current and not bool(p.get("confirmed_location")):
            name=self._location_label(endpoint,row["from_name_snapshot"] if target==Status.ON_BOARD else row["to_name_snapshot"])
            raise AppError("LOCATION_CONFIRMATION_REQUIRED",f"Confirm that the cargo was physically {('picked up at' if target==Status.ON_BOARD else 'delivered to')} {name}.",{"location":name})
        now=now_utc();tz=None;date=None
        if target==Status.DONE:
            tz=self._timezone();date=business_date(now,tz)
            self.db.execute("UPDATE transfer_legs SET status='Done',completed_utc=?,completed_date=?,completed_tz=?,version=version+1,updated_utc=? WHERE no=?",(now,date,tz,now,row["no"]))
        else:self.db.execute("UPDATE transfer_legs SET status='On Board',version=version+1,updated_utc=? WHERE no=?",(now,row["no"]))
        self._audit(command_id,"transfer",row["no"],event.upper(),before={"status":row["status"]},after={"status":target.value,"completed_date":date,"location":row["from_name_snapshot"] if target==Status.ON_BOARD else row["to_name_snapshot"]})
        return {"no":row["no"],"cargo_id":row["cargo_id"],"status":target.value,"completed_date":date}

    def _cmd_pick_up(self,command_id,p):return self._transition(command_id,p,Status.ON_BOARD,"pickup")
    def _cmd_deliver(self,command_id,p):return self._transition(command_id,p,Status.DONE,"deliver")

    def _cmd_correct_transfer(self,command_id,p):
        row=self._active_leg(p.get("no"));self._expected(row,p.get("expected_version"))
        p={**p,'preserve_text_case':p.get('preserve_text_case',self._preserved_description(row['no']))}
        reason=self._required(p.get("reason"),"Correction reason",1000)
        if self._descendants(row["no"]):raise AppError("HAS_FOLLOWING_TRANSFER","Resolve later transfers before correcting this record.")
        chain=self._chain(row["cargo_id"])
        if not chain or chain[-1]["no"]!=row["no"]:raise AppError("NOT_CHAIN_TAIL","Correct the latest nonvoid transfer first.")
        source=self._location(p.get("from_location_id",row["from_location_id"]),allow_inactive=True)
        dest=self._location(p.get("to_location_id",row["to_location_id"]),allow_inactive=True)
        status=str(p.get("status",row["status"]))
        if status not in [s.value for s in Status]:raise AppError("VALIDATION","Choose Waiting, On Board or Done.")
        if row["predecessor_no"] is not None:
            prev=self.repo.leg(row["predecessor_no"],include_void=False)
            if prev is None or prev["status"]!="Done":raise AppError("INVALID_CHAIN","The preceding transfer must remain delivered.")
        description=self._description({**p,'description':p.get('description',row['description'])}, row['description'])
        comment=self._optional(p.get("comment",row["comment"]),"Comment",MAX_COMMENT)
        before={"status":row["status"],"from":row["from_name_snapshot"],"to":row["to_name_snapshot"],"description":row["description"],"comment":row["comment"],"completed_date":row["completed_date"]}
        completed=(None,None,None)
        if status==Status.DONE.value:
            if row["status"]==Status.DONE.value:completed=(row["completed_utc"],row["completed_date"],row["completed_tz"])
            else:
                utc=now_utc();tz=self._timezone();completed=(utc,business_date(utc,tz),tz)
        now=now_utc()
        self.db.execute("""UPDATE transfer_legs SET from_location_id=?,from_name_snapshot=?,to_location_id=?,to_name_snapshot=?,
          description=?,comment=?,status=?,completed_utc=?,completed_date=?,completed_tz=?,version=version+1,updated_utc=? WHERE no=?""",
          (source["id"],source["name"],dest["id"],dest["name"],description,comment,status,*completed,now,row["no"]))
        after={"status":status,"from":source["name"],"to":dest["name"],"description":description,"comment":comment,"completed_date":completed[1]}
        if row['description'] != description:
            self.db.execute('UPDATE cargo SET default_description=?,updated_utc=? WHERE id=? AND default_description=?',
                            (description,now,row['cargo_id'],row['description']))
        self._audit(command_id,"transfer",row["no"],"TRANSFER_CORRECTED",before=before,after={**after,'raw_description':p.get('description'),'preserve_text_case':bool(p.get('preserve_text_case'))},reason=reason)
        return {"no":row["no"],"cargo_id":row["cargo_id"],"status":status}

    def _preserved_description(self, number):
        return self._preserved_descriptions([number]).get(number,False)

    def _preserved_descriptions(self, numbers):
        numbers=list(dict.fromkeys(numbers));result={}
        for start in range(0,len(numbers),500):
            group=numbers[start:start+500]
            placeholders=','.join('?' for _ in group)
            for event in self.db.execute(f"SELECT entity_id,data_json FROM audit_events WHERE entity_type='transfer' AND entity_id IN ({placeholders}) ORDER BY id DESC",group):
                number=event['entity_id']
                if number in result:continue
                after=json.loads(event['data_json']).get('after')
                if isinstance(after,dict) and 'description' in after:
                    result[number]=bool(after.get('preserve_text_case'))
        return result

    def text_case_preview(self):
        changes = []
        protected = set()
        rows=list(self.db.execute('SELECT no,cargo_id,description,version FROM transfer_legs WHERE is_void=0 ORDER BY no'))
        preserved=self._preserved_descriptions(row['no'] for row in rows)
        for row in rows:
            if preserved.get(row['no'],False):
                protected.add((row['cargo_id'], row['description']))
                continue
            after, review = normalize(row['description'], 'cargo')
            if after != row['description']:
                changes.append({'kind':'transfer','id':row['no'],'version':row['version'],
                    'label':f"Transfer #{row['no']}",'before':row['description'],'after':after,'review':review})
        for row in self.db.execute('SELECT id,cargo_tag,default_description,updated_utc FROM cargo WHERE EXISTS(SELECT 1 FROM transfer_legs t WHERE t.cargo_id=cargo.id AND t.is_void=0)'):
            if (row['id'], row['default_description']) in protected:
                continue
            after, review = normalize(row['default_description'], 'cargo')
            if after != row['default_description']:
                changes.append({'kind':'cargo','id':row['id'],'version':row['updated_utc'],
                    'label':f"{row['cargo_tag']} / next-transfer default",'before':row['default_description'],
                    'after':after,'review':review})
        return changes

    def _cmd_normalize_text_case(self, command_id, p):
        changes = p.get('changes', [])
        if not changes:
            raise AppError('NO_SELECTION','Select at least one spelling change.')
        available = {(c['kind'], c['id']):c for c in self.text_case_preview()}
        seen = set()
        for change in changes:
            token = (change.get('kind'), change.get('id'))
            if token in seen or available.get(token) != change:
                raise AppError('STALE_RECORD','Cargo text changed. Reopen the spelling preview.')
            seen.add(token)
        for change in changes:
            if change['kind'] == 'transfer':
                self.db.execute('UPDATE transfer_legs SET description=?,version=version+1,updated_utc=? WHERE no=?',
                                (change['after'],now_utc(),change['id']))
            else:
                self.db.execute('UPDATE cargo SET default_description=?,updated_utc=? WHERE id=?',
                                (change['after'],now_utc(),change['id']))
            self._audit(command_id,change['kind'],change['id'],'TEXT_CASE_NORMALIZED',
                before={'description':change['before']},after={'description':change['after'],'raw_description':change['before']})
        return {'changed':len(changes)}

    def _cmd_void_transfer(self,command_id,p):
        row=self._active_leg(p.get("no"));self._expected(row,p.get("expected_version"))
        reason=self._required(p.get("reason"),"Void reason",1000)
        if row["status"]==Status.ON_BOARD.value:raise AppError("ON_BOARD_MUST_BE_RESOLVED","Use Correct record to confirm the cargo's actual location before voiding it.")
        if self._descendants(row["no"]):raise AppError("HAS_FOLLOWING_TRANSFER","Resolve later transfers before voiding this one.")
        chain=self._chain(row["cargo_id"])
        if not chain or chain[-1]["no"]!=row["no"]:raise AppError("NOT_CHAIN_TAIL","Only the latest nonvoid transfer can be voided.")
        self.db.execute("UPDATE transfer_legs SET is_void=1,version=version+1,updated_utc=? WHERE no=?",(now_utc(),row["no"]))
        self._audit(command_id,"transfer",row["no"],"TRANSFER_VOIDED",before={"status":row["status"]},after={"voided":True},reason=reason)
        return {"no":row["no"],"cargo_id":row["cargo_id"],"voided":True}

    def _cmd_restore_transfer(self,command_id,p):
        try:no=int(p["no"]);expected=int(p["expected_version"])
        except (ValueError,TypeError,KeyError) as exc:raise AppError("VALIDATION","Choose a voided transfer to restore.") from exc
        row=self.repo.leg(no)
        if row is None or not row["is_void"]:raise AppError("TRANSFER_NOT_VOIDED","This transfer is not voided.")
        self._expected(row,expected);reason=self._required(p.get("reason"),"Restore reason",1000)
        later=self.db.execute("SELECT no FROM transfer_legs WHERE cargo_id=? AND no>? AND is_void=0 ORDER BY no",(row["cargo_id"],no)).fetchone()
        if later:raise AppError("REPLACEMENT_CHAIN_EXISTS",f"Transfer #{later['no']} already replaced this route; the old branch cannot be restored.")
        if row["predecessor_no"] is not None:
            prev=self.repo.leg(row["predecessor_no"],include_void=False)
            if prev is None or prev["status"]!="Done":raise AppError("INVALID_CHAIN","The preceding transfer must be active and delivered.")
        self.db.execute("UPDATE transfer_legs SET is_void=0,version=version+1,updated_utc=? WHERE no=?",(now_utc(),no))
        self._audit(command_id,"transfer",no,"TRANSFER_RESTORED",before={"voided":True,"status":row["status"]},after={"voided":False,"status":row["status"]},reason=reason)
        return {"no":no,"cargo_id":row["cargo_id"],"restored":True}

    def _cmd_create_vessel(self,command_id,p):
        name=self._required(p.get("name"),"Vessel name",120);key=normalized_name(name)
        old=self.db.execute("SELECT * FROM locations WHERE name_key=?",(key,)).fetchone()
        if old:
            if not old["active"] and p.get("reactivate"):
                return self._cmd_set_vessel_active(command_id,{"id":old["id"],"active":True,"reason":p.get("reason","")})
            raise AppError("DUPLICATE_VESSEL",f"Vessel {name} already exists. Reactivate its saved record instead.")
        now=now_utc();cur=self.db.execute("INSERT INTO locations(kind,name,name_key,active,created_utc,updated_utc) VALUES('VESSEL',?,?,1,?,?)",(name,key,now,now))
        self._audit(command_id,"location",cur.lastrowid,"VESSEL_CREATED",after={"name":name})
        return {"id":cur.lastrowid,"name":name}

    def _cmd_rename_vessel(self,command_id,p):
        row=self._location(p.get("id"),allow_inactive=True)
        if row["kind"]!="VESSEL":raise AppError("LOCATION_IMMUTABLE","Rename the port in Settings.")
        if not row["active"]:raise AppError("VESSEL_INACTIVE","Reactivate this vessel before renaming it.")
        name=self._required(p.get("name"),"Vessel name",120)
        if name==row["name"]:return {"id":row["id"],"name":name,"unchanged":True}
        in_route=self.db.execute("SELECT 1 FROM voyage_stops s JOIN voyages v ON v.id=s.voyage_id WHERE s.vessel_id=? AND s.removed=0 AND v.is_current=1",(row["id"],)).fetchone()
        unfinished=self.db.execute("SELECT no FROM transfer_legs WHERE is_void=0 AND status IN ('Waiting','On Board') AND (from_location_id=? OR to_location_id=?) LIMIT 1",(row["id"],row["id"])).fetchone()
        if in_route or unfinished:raise AppError("VESSEL_IN_USE","Remove this vessel from the active voyage and resolve its unfinished transfers before renaming it.")
        self.db.execute("UPDATE locations SET name=?,name_key=?,updated_utc=? WHERE id=?",(name,normalized_name(name),now_utc(),row["id"]))
        self._audit(command_id,"location",row["id"],"VESSEL_RENAMED",before={"name":row["name"]},after={"name":name})
        return {"id":row["id"],"name":name}

    def _cmd_set_vessel_active(self,command_id,p):
        row=self._location(p.get("id"),allow_inactive=True)
        if row["kind"]!="VESSEL":raise AppError("LOCATION_IMMUTABLE","The port cannot be removed from Fleet.")
        if self.db.execute("SELECT 1 FROM audit_events WHERE entity_type='location' AND entity_id=? AND event_type='VESSEL_DELETED'",(row["id"],)).fetchone():raise AppError("LOCATION_NOT_FOUND","This vessel was deleted.")
        active=bool(p.get("active"))
        if active==bool(row["active"]):return {"id":row["id"],"active":active,"unchanged":True}
        if not active:
            stop=self.db.execute("SELECT s.id FROM voyage_stops s JOIN voyages v ON v.id=s.voyage_id WHERE s.vessel_id=? AND s.removed=0 AND v.is_current=1 LIMIT 1",(row["id"],)).fetchone()
            leg=self.db.execute("SELECT no FROM transfer_legs WHERE is_void=0 AND status IN ('Waiting','On Board') AND (from_location_id=? OR to_location_id=?) LIMIT 1",(row["id"],row["id"])).fetchone()
            if stop or leg:raise AppError("VESSEL_IN_USE",f"Remove active stop #{stop['id']} or resolve transfer #{leg['no']} before deactivating this vessel." if stop and leg else f"Remove active stop #{stop['id']} first." if stop else f"Resolve transfer #{leg['no']} first.")
        self.db.execute("UPDATE locations SET active=?,updated_utc=? WHERE id=?",(int(active),now_utc(),row["id"]))
        self._audit(command_id,"location",row["id"],"VESSEL_REACTIVATED" if active else "VESSEL_DEACTIVATED",before={"active":bool(row["active"])},after={"active":active},reason=p.get("reason",""))
        return {"id":row["id"],"active":active}

    def _cmd_delete_vessel(self,command_id,p):
        row=self._location(p.get("id"),allow_inactive=True)
        if row["kind"]!="VESSEL":raise AppError("LOCATION_IMMUTABLE","The port cannot be deleted.")
        if self.db.execute("SELECT 1 FROM audit_events WHERE entity_type='location' AND entity_id=? AND event_type='VESSEL_DELETED'",(row["id"],)).fetchone():raise AppError("LOCATION_NOT_FOUND","This vessel was already deleted.")
        stop=self.db.execute("SELECT 1 FROM voyage_stops s JOIN voyages v ON v.id=s.voyage_id WHERE s.vessel_id=? AND s.removed=0 AND v.is_current=1",(row["id"],)).fetchone()
        leg=self.db.execute("SELECT 1 FROM transfer_legs WHERE is_void=0 AND status!='Done' AND (from_location_id=? OR to_location_id=?)",(row["id"],row["id"])).fetchone()
        if stop or leg:raise AppError("VESSEL_IN_USE","Resolve unfinished transfers and remove this vessel from the current voyage before deleting it.")
        self.db.execute("UPDATE locations SET active=0,name_key=?,updated_utc=? WHERE id=?",("deleted:"+uuid4().hex,now_utc(),row["id"]))
        self._audit(command_id,"location",row["id"],"VESSEL_DELETED",before=dict(row),after={"deleted":True})
        return {"id":row["id"]}

    def _cmd_create_voyage(self,command_id,p):
        name=self._required(p.get("name"),"Voyage name",120);current=self.repo.current_voyage();now=now_utc()
        if current:self.db.execute("UPDATE voyages SET is_current=0,archived=1,updated_utc=? WHERE id=?",(now,current["id"]))
        cur=self.db.execute("INSERT INTO voyages(name,is_current,created_utc,updated_utc) VALUES(?,1,?,?)",(name,now,now))
        self.db.execute("UPDATE app_state SET current_voyage_id=?,current_stop_id=NULL WHERE id=1",(cur.lastrowid,))
        self._audit(command_id,"voyage",cur.lastrowid,"VOYAGE_CREATED",after={"name":name,"stops":[]})
        return {"voyage_id":cur.lastrowid,"name":name}

    def _cmd_copy_voyage(self,command_id,p):
        name=self._required(p.get("name"),"Voyage name",120)
        if p.get("voyage_id"):
            try:source_id=int(p["voyage_id"])
            except (ValueError,TypeError) as exc:raise AppError("NO_VOYAGE","Choose an existing voyage to copy.") from exc
            old=self.db.execute("SELECT * FROM voyages WHERE id=?",(source_id,)).fetchone()
        else:old=self.repo.current_voyage() or self.db.execute("SELECT * FROM voyages v WHERE NOT EXISTS(SELECT 1 FROM audit_events e WHERE e.entity_type='voyage' AND e.entity_id=v.id AND e.event_type='VOYAGE_DELETED') ORDER BY id DESC LIMIT 1").fetchone()
        if old is None or self.db.execute("SELECT 1 FROM audit_events WHERE entity_type='voyage' AND entity_id=? AND event_type='VOYAGE_DELETED'",(old["id"],)).fetchone():raise AppError("NO_VOYAGE","Create a voyage before copying one.")
        stops=self.db.execute("SELECT * FROM voyage_stops WHERE voyage_id=? AND removed=0 ORDER BY position,id",(old["id"],)).fetchall()
        current=self.repo.current_voyage();now=now_utc()
        if current:self.db.execute("UPDATE voyages SET is_current=0,archived=1,updated_utc=? WHERE id=?",(now,current["id"]))
        cur=self.db.execute("INSERT INTO voyages(name,is_current,created_utc,updated_utc) VALUES(?,1,?,?)",(name,now,now))
        for position,stop in enumerate(stops,1):
            self.db.execute("INSERT INTO voyage_stops(voyage_id,vessel_id,vessel_name_snapshot,position,created_utc) VALUES(?,?,?,?,?)",(cur.lastrowid,stop["vessel_id"],stop["vessel_name_snapshot"],position,now))
        self.db.execute("UPDATE app_state SET current_voyage_id=?,current_stop_id=NULL WHERE id=1",(cur.lastrowid,))
        self._audit(command_id,"voyage",cur.lastrowid,"VOYAGE_COPIED",after={"name":name,"source_id":old["id"],"stops":[s["vessel_name_snapshot"] for s in stops]})
        return {"voyage_id":cur.lastrowid,"name":name,"stops":len(stops)}

    def _cmd_delete_voyage(self,command_id,p):
        identifier=int(p.get("voyage_id",0))
        voyage=self.db.execute("SELECT * FROM voyages WHERE id=?",(identifier,)).fetchone()
        if voyage is None or self.db.execute("SELECT 1 FROM audit_events WHERE entity_type='voyage' AND entity_id=? AND event_type='VOYAGE_DELETED'",(identifier,)).fetchone():
            raise AppError("NO_VOYAGE","That voyage no longer exists.")
        if voyage["is_current"]:
            self.db.execute("UPDATE app_state SET current_voyage_id=NULL,current_stop_id=NULL WHERE id=1")
        self.db.execute("UPDATE voyages SET is_current=0,archived=1,updated_utc=? WHERE id=?",(now_utc(),identifier))
        self._audit(command_id,"voyage",identifier,"VOYAGE_DELETED",before=dict(voyage),after={"deleted":True})
        return {"voyage_id":identifier}

    def _cmd_set_current_voyage(self,command_id,p):
        try:identifier=int(p.get("voyage_id"))
        except (ValueError,TypeError) as exc:raise AppError("NO_VOYAGE","Choose a voyage.") from exc
        voyage=self.db.execute("SELECT * FROM voyages WHERE id=?",(identifier,)).fetchone()
        if voyage is None or self.db.execute("SELECT 1 FROM audit_events WHERE entity_type='voyage' AND entity_id=? AND event_type='VOYAGE_DELETED'",(identifier,)).fetchone():raise AppError("NO_VOYAGE","That voyage no longer exists.")
        if voyage["archived"]:raise AppError("VOYAGE_ARCHIVED","Historical voyages are read-only. Copy this route to create a new working voyage.")
        state=self.repo.current_state()
        if state["current_voyage_id"]==identifier:return {"voyage_id":identifier,"unchanged":True}
        current=self.repo.current_voyage();now=now_utc()
        if current:self.db.execute("UPDATE voyages SET is_current=0,archived=1,updated_utc=? WHERE id=?",(now,current["id"]))
        self.db.execute("UPDATE voyages SET is_current=1,updated_utc=? WHERE id=?",(now,identifier))
        self.db.execute("UPDATE app_state SET current_voyage_id=?,current_stop_id=NULL WHERE id=1",(identifier,))
        self._audit(command_id,"voyage",identifier,"VOYAGE_ACTIVATED",before={"voyage_id":state["current_voyage_id"],"stop_id":state["current_stop_id"]},after={"voyage_id":identifier,"stop_id":None})
        return {"voyage_id":identifier}

    def _cmd_add_stop(self,command_id,p):
        state=self.repo.current_state()
        if state["current_voyage_id"] is None:raise AppError("NO_VOYAGE","Create a voyage before adding a stop.")
        vessel=self._location(p.get("vessel_id"))
        if vessel["kind"]!="VESSEL":raise AppError("INVALID_STOP","The port is the terminal plan stop, not a Fleet vessel stop.")
        position=int(self.db.execute("SELECT COALESCE(MAX(position),0)+1 FROM voyage_stops WHERE voyage_id=? AND removed=0",(state["current_voyage_id"],)).fetchone()[0])
        cur=self.db.execute("INSERT INTO voyage_stops(voyage_id,vessel_id,vessel_name_snapshot,position,created_utc) VALUES(?,?,?,?,?)",(state["current_voyage_id"],vessel["id"],vessel["name"],position,now_utc()))
        self._audit(command_id,"stop",cur.lastrowid,"VOYAGE_STOP_ADDED",after={"vessel_id":vessel["id"],"name":vessel["name"],"position":position})
        return {"stop_id":cur.lastrowid,"position":position}

    def _cmd_remove_stop(self,command_id,p):
        state=self.repo.current_state()
        try:identifier=int(p["stop_id"])
        except (KeyError,ValueError,TypeError) as exc:raise AppError("STOP_NOT_FOUND","Choose a stop on the active voyage.") from exc
        stop=self.db.execute("SELECT * FROM voyage_stops WHERE id=? AND voyage_id=? AND removed=0",(identifier,state["current_voyage_id"])).fetchone()
        if stop is None:raise AppError("STOP_NOT_FOUND","This stop is not on the active voyage.")
        was_current=state["current_stop_id"]==identifier
        if was_current:self.db.execute("UPDATE app_state SET current_stop_id=NULL WHERE id=1")
        self.db.execute("UPDATE voyage_stops SET removed=1 WHERE id=?",(identifier,))
        tail=self.db.execute("SELECT id FROM voyage_stops WHERE voyage_id=? AND removed=0 AND position>? ORDER BY position",(state["current_voyage_id"],stop["position"])).fetchall()
        for row in tail:self.db.execute("UPDATE voyage_stops SET position=-id WHERE id=?",(row["id"],))
        for position,row in enumerate(tail,start=stop["position"]):self.db.execute("UPDATE voyage_stops SET position=? WHERE id=?",(position,row["id"]))
        self._audit(command_id,"stop",identifier,"VOYAGE_STOP_REMOVED",before={"name":stop["vessel_name_snapshot"],"position":stop["position"]},after={"removed":True})
        return {"stop_id":identifier,"current_stop_cleared":was_current}

    def _cmd_reorder_stops(self,command_id,p):
        state=self.repo.current_state()
        try:desired=[int(item) for item in p.get("stop_ids",[])]
        except (ValueError,TypeError) as exc:raise AppError("ROUTE_CONFLICT","The route selection is invalid.") from exc
        current=[r["id"] for r in self.db.execute("SELECT id FROM voyage_stops WHERE voyage_id=? AND removed=0 ORDER BY position,id",(state["current_voyage_id"],))]
        if len(desired)!=len(set(desired)) or set(current)!=set(desired):raise AppError("ROUTE_CONFLICT","The route changed. Reload it before reordering.")
        for identifier in current:self.db.execute("UPDATE voyage_stops SET position=-id WHERE id=?",(identifier,))
        for position,identifier in enumerate(desired,1):self.db.execute("UPDATE voyage_stops SET position=? WHERE id=?",(position,identifier))
        self._audit(command_id,"voyage",state["current_voyage_id"],"VOYAGE_REORDERED",before=current,after=desired)
        return {"stop_ids":desired,"current_stop_id":state["current_stop_id"]}

    def _cmd_choose_stop(self,command_id,p):
        state=self.repo.current_state()
        try:identifier=int(p.get("stop_id"))
        except (ValueError,TypeError) as exc:raise AppError("INVALID_STOP","Choose a stop on the active voyage.") from exc
        stop=self.db.execute("SELECT * FROM voyage_stops WHERE id=? AND voyage_id=? AND removed=0",(identifier,state["current_voyage_id"])).fetchone()
        if stop is None:raise AppError("INVALID_STOP","Choose a stop on the active voyage.")
        self.db.execute("UPDATE app_state SET current_stop_id=? WHERE id=1",(identifier,))
        self._audit(command_id,"stop",identifier,"CURRENT_STOP_CHANGED",before={"stop_id":state["current_stop_id"]},after={"stop_id":identifier,"vessel_id":stop["vessel_id"]})
        return {"stop_id":identifier,"vessel_id":stop["vessel_id"]}

    def _cmd_clear_current_stop(self,command_id,p):
        state=self.repo.current_state()
        current=state['current_stop_id']
        if current is None:return {'stop_id':None,'vessel_id':None}
        if current is not None and p.get('expected_stop_id')!=current:
            raise AppError('STALE_STOP','The current stop changed. Refresh before clearing it.')
        self.db.execute('UPDATE app_state SET current_stop_id=NULL WHERE id=1')
        self._audit(command_id,'voyage',state['current_voyage_id'],'CURRENT_STOP_CLEARED',
                    before={'stop_id':current},after={'stop_id':None})
        return {'stop_id':None,'vessel_id':None}

    def _cmd_set_setting(self,command_id,p):
        key=str(p.get("key",""))
        if key not in {"business_timezone","transfer_vessel_name","port_name"}:raise AppError("INVALID_SETTING","This setting cannot be changed.")
        value=(self._required(p.get("value"),"Business time zone" if key=="business_timezone" else "Port name",120)
               if key in {"business_timezone","port_name"} else self._optional(p.get("value"),"Transfer vessel name",120))
        if key=="business_timezone":
            from zoneinfo import ZoneInfo,ZoneInfoNotFoundError
            try:ZoneInfo(value)
            except (ZoneInfoNotFoundError,ValueError) as exc:raise AppError("INVALID_TIMEZONE",f"Unknown time zone: {value}") from exc
        old=self.db.execute("SELECT value FROM settings WHERE key=?",(key,)).fetchone()[0]
        if old==value:return {"key":key,"value":value,"unchanged":True}
        self.db.execute("UPDATE settings SET value=?,updated_utc=? WHERE key=?",(value,now_utc(),key))
        if key=="port_name":
            self.db.execute("UPDATE locations SET name=?,name_key=?,updated_utc=? WHERE kind='SHORE'",(value,normalized_name(value),now_utc()))
        self._audit(command_id,"settings",1,"SETTING_CHANGED",before={key:old},after={key:value})
        return {"key":key,"value":value}

    def _cmd_update_settings(self,command_id,p):
        # handle() supplies one transaction: a later invalid field rolls back every setting.
        values=p.get("values",{})
        if not isinstance(values,dict) or not values:
            raise AppError("INVALID_SETTING","Choose the settings to save.")
        return {"settings":{key:self._cmd_set_setting(command_id,{"key":key,"value":value})["value"]
                            for key,value in values.items()}}

    def _location_label(self, location_id, snapshot):
        location=self.db.execute("SELECT kind,name FROM locations WHERE id=?",(location_id,)).fetchone()
        return location["name"] if location and location["kind"]=="SHORE" else snapshot

    def _cargo_location(self,cargo_id):
        leg=self.db.execute("SELECT * FROM transfer_legs WHERE cargo_id=? AND is_void=0 ORDER BY no DESC LIMIT 1",(cargo_id,)).fetchone()
        if leg is None:return {"label":"Not confirmed","confirmed":False,"at_utc":None}
        if leg["status"]=="On Board":
            name=self.db.execute("SELECT value FROM settings WHERE key='transfer_vessel_name'").fetchone()[0].strip()
            return {"label":name or "On my vessel","confirmed":True,"at_utc":leg["updated_utc"]}
        if leg["status"]=="Done":return {"label":self._location_label(leg["to_location_id"],leg["to_name_snapshot"]),"confirmed":True,"at_utc":leg["completed_utc"]}
        if leg["predecessor_no"] is not None:
            prev=self.repo.leg(leg["predecessor_no"],include_void=False)
            if prev is not None and prev["status"]=="Done":return {"label":self._location_label(prev["to_location_id"],prev["to_name_snapshot"]),"confirmed":True,"at_utc":prev["completed_utc"]}
        return {"label":"Not confirmed","confirmed":False,"at_utc":None}

    def _row_dicts(self,rows):
        results=[dict(row) for row in rows]
        if not results:return results
        preserved=self._preserved_descriptions(row['no'] for row in results)
        cargo_ids=list({int(row["cargo_id"]) for row in results})
        placeholders=",".join("?" for _ in cargo_ids)
        latest={int(row["cargo_id"]):row for row in self.db.execute(f"""SELECT t.* FROM transfer_legs t JOIN (
          SELECT cargo_id,max(no) AS no FROM transfer_legs WHERE is_void=0 AND cargo_id IN ({placeholders}) GROUP BY cargo_id
          ) x ON x.no=t.no""",cargo_ids)}
        predecessor_ids=[row["predecessor_no"] for row in latest.values() if row["status"]=="Waiting" and row["predecessor_no"] is not None]
        predecessors={int(row["no"]):row for row in self.db.execute(
            f"SELECT * FROM transfer_legs WHERE no IN ({','.join('?' for _ in predecessor_ids)}) AND is_void=0 AND status='Done'",predecessor_ids
        )} if predecessor_ids else {}
        state=self.repo.current_state();vessel_name=self.db.execute("SELECT value FROM settings WHERE key='transfer_vessel_name'").fetchone()[0].strip()
        shore=self.db.execute("SELECT id,name FROM locations WHERE kind='SHORE'").fetchone()
        for result in results:
            result['preserve_text_case']=preserved.get(result['no'],False)
            if result["from_location_id"]==shore["id"]:result["from_name"]=shore["name"]
            if result["to_location_id"]==shore["id"]:result["to_name"]=shore["name"]
            result["is_void"]=bool(result["is_void"]);result["from_active"]=bool(result["from_active"]);result["to_active"]=bool(result["to_active"])
            leg=latest.get(int(result["cargo_id"]))
            location={"label":"Not confirmed","confirmed":False,"at_utc":None}
            if leg and leg["status"]=="On Board":location={"label":vessel_name or "On my vessel","confirmed":True,"at_utc":leg["updated_utc"]}
            elif leg and leg["status"]=="Done":location={"label":shore["name"] if leg["to_location_id"]==shore["id"] else leg["to_name_snapshot"],"confirmed":True,"at_utc":leg["completed_utc"]}
            elif leg and leg["predecessor_no"] in predecessors:
                previous=predecessors[leg["predecessor_no"]];location={"label":shore["name"] if previous["to_location_id"]==shore["id"] else previous["to_name_snapshot"],"confirmed":True,"at_utc":previous["completed_utc"]}
            result["cargo_location"]=location
            at_from=state["vessel_id"] is not None and result["from_location_id"]==state["vessel_id"]
            at_to=state["vessel_id"] is not None and result["to_location_id"]==state["vessel_id"]
            result["current_relation"]="both" if at_from and at_to else "from" if at_from else "to" if at_to else ""
            result["allowed_actions"]=[action for status,action in (("Waiting","pick_up"),("On Board","deliver")) if result["status"]==status]
        return results

    def _row_dict(self,row):return self._row_dicts([row])[0]

    def _where(self,view,filters):
        state=self.repo.current_state();clauses=[];args=[]
        if view=="current_stop":
            if state["vessel_id"] is None:clauses.append("0=1")
            else:clauses.append("t.is_void=0 AND t.status!='Done' AND (t.from_location_id=? OR t.to_location_id=?)");args.extend((state["vessel_id"],state["vessel_id"]))
        elif view=="on_my_vessel":clauses.append("t.is_void=0 AND t.status='On Board'")
        elif view=="voided":clauses.append("t.is_void=1")
        else:
            date_range=filters.get("range","open")
            clauses.append({"open":"t.is_void=0 AND t.status IN ('Waiting','On Board')","completed":"t.is_void=0 AND t.status='Done'","voided":"t.is_void=1","all":"t.is_void=0"}.get(date_range,"t.is_void=0"))
        if view=="current_stop":
            action=filters.get("action","")
            if action=="pick_up":clauses.append("t.status='Waiting' AND t.from_location_id=?");args.append(state["vessel_id"])
            elif action=="deliver":clauses.append("t.status='On Board' AND t.to_location_id=?");args.append(state["vessel_id"])
            elif action=="other":clauses.append("NOT ((t.status='Waiting' AND t.from_location_id=?) OR (t.status='On Board' AND t.to_location_id=?))");args.extend((state["vessel_id"],state["vessel_id"]))
        status=filters.get("status","")
        if status and view!="on_my_vessel":clauses.append("t.status=?");args.append(status)
        vessel=filters.get("vessel_id")
        if vessel:clauses.append("(t.from_location_id=? OR t.to_location_id=?)");args.extend((int(vessel),int(vessel)))
        query=str(filters.get("query","")).strip()
        if query:
            clauses.append("(instr(lower(c.cargo_tag),lower(?))>0 OR instr(lower(t.description),lower(?))>0 OR instr(lower(t.comment),lower(?))>0 OR instr(lower(COALESCE(c.physical_id,'')),lower(?))>0 OR CAST(t.no AS TEXT)=?)")
            no=query.removeprefix("#")
            args.extend((query,query,query,query,no))
        return (" AND ".join(clauses) or "1=1"),args

    def workspace_snapshot(self):
        db=self.db;db.execute("BEGIN")
        try:
            state=self.repo.current_state()
            vessels=[dict(r) for r in db.execute("SELECT id,name,active FROM locations l WHERE kind='VESSEL' AND NOT EXISTS(SELECT 1 FROM audit_events e WHERE e.entity_type='location' AND e.entity_id=l.id AND e.event_type='VESSEL_DELETED') ORDER BY active DESC,name COLLATE NOCASE")]
            locations=[dict(r) for r in db.execute("SELECT id,name,active,kind FROM locations l WHERE NOT EXISTS(SELECT 1 FROM audit_events e WHERE e.entity_type='location' AND e.entity_id=l.id AND e.event_type='VESSEL_DELETED') ORDER BY kind,name COLLATE NOCASE")]
            voyages=[dict(r) for r in db.execute("SELECT id,name,is_current,archived,created_utc FROM voyages v WHERE NOT EXISTS(SELECT 1 FROM audit_events e WHERE e.entity_type='voyage' AND e.entity_id=v.id AND e.event_type='VOYAGE_DELETED') ORDER BY id DESC")]
            stops=[]
            if state["current_voyage_id"] is not None:
                stops=[dict(r) for r in db.execute("SELECT s.id,s.vessel_id,s.vessel_name_snapshot,s.position,l.name AS current_name FROM voyage_stops s JOIN locations l ON l.id=s.vessel_id WHERE s.voyage_id=? AND s.removed=0 ORDER BY s.position,s.id",(state["current_voyage_id"],))]
            settings={r["key"]:r["value"] for r in db.execute("SELECT key,value FROM settings")}
            pickup=[];delivery=[];other=[]
            if state["vessel_id"] is not None:
                rows=db.execute("SELECT no,status,from_location_id,to_location_id FROM transfer_legs WHERE is_void=0 AND status!='Done' AND (from_location_id=? OR to_location_id=?) ORDER BY no",(state["vessel_id"],state["vessel_id"])).fetchall()
                for row in rows:
                    if row["status"]=="Waiting" and row["from_location_id"]==state["vessel_id"]:pickup.append(row["no"])
                    elif row["status"]=="On Board" and row["to_location_id"]==state["vessel_id"]:delivery.append(row["no"])
                    else:other.append(row["no"])
            onboard=[int(r[0]) for r in db.execute("SELECT no FROM transfer_legs WHERE is_void=0 AND status='On Board' ORDER BY no")]
            total=int(db.execute("SELECT count(*) FROM transfer_legs WHERE is_void=0").fetchone()[0])
            open_total=int(db.execute("SELECT count(*) FROM transfer_legs WHERE is_void=0 AND status IN ('Waiting','On Board')").fetchone()[0])
            result={"dataset_epoch":self.dataset_epoch,"revision":int(state["data_revision"]),"current_voyage_id":state["current_voyage_id"],"current_stop_id":state["current_stop_id"],"voyage_name":state["voyage_name"],"current_vessel_id":state["vessel_id"],"current_vessel_name":state["vessel_current_name"],"current_stop_position":state["position"],"vessels":vessels,"locations":locations,"shore_location_id":next((l["id"] for l in locations if l["kind"]=="SHORE"),None),"voyages":voyages,"stops":stops,"settings":settings,"ready_pickup":pickup,"ready_delivery":delivery,"other_related":other,"onboard_ids":onboard,"counts":{"all":total,"open":open_total,"pickup":len(pickup),"delivery":len(delivery),"other":len(other),"onboard":len(onboard)}}
            result['ledger_identity']=db.execute('SELECT identity FROM ledger_meta WHERE id=1').fetchone()[0]
            db.commit();return result
        except Exception:db.rollback();raise

    def query_transfers(self,view,filters=None,*,order_column=0,descending=False,limit=200,offset=0):
        if view not in {"all","current_stop","on_my_vessel","voided"}:raise AppError("INVALID_VIEW","Choose a valid ledger view.")
        filters=filters or {};order=self._sort_order(order_column,descending)
        where,args=self._where(view,filters)
        total=int(self.db.execute(f"SELECT count(*) FROM transfer_legs t JOIN cargo c ON c.id=t.cargo_id WHERE {where}",args).fetchone()[0])
        limit=max(1,min(int(limit),1000));offset=max(0,int(offset))
        rows=self.repo.rows(where,args,order,limit,offset)
        return {"epoch":self.dataset_epoch,"revision":self._revision(),"total":total,"offset":offset,"limit":limit,"rows":self._row_dicts(rows)}

    def health_report(self):
        integrity=[r[0] for r in self.db.execute('PRAGMA quick_check')]
        foreign_keys=[list(r) for r in self.db.execute('PRAGMA foreign_key_check')]
        failures=check_integrity(str(self.path))
        return {'integrity':'ok' if integrity==['ok'] else 'failed',
                'foreign_keys':'ok' if not foreign_keys else 'failed',
                'business_validation':'ok' if not failures else 'failed',
                'integrity_details':integrity,'foreign_key_details':foreign_keys,
                'failures':failures,'epoch':self.dataset_epoch}

    @staticmethod
    def _sort_order(order_column,descending):
        order=SORT_COLUMNS.get(int(order_column),SORT_COLUMNS[0])
        if order_column==7:order+=f" {'DESC' if descending else 'ASC'}, t.no ASC"
        else:order+=f" {'DESC' if descending else 'ASC'}"+("" if order_column==0 else ", t.no ASC")
        return order

    def _simple_leg(self,row):
        return {"no":row["no"],"predecessor_no":row["predecessor_no"],"from":self._location_label(row["from_location_id"],row["from_name_snapshot"]),"to":self._location_label(row["to_location_id"],row["to_name_snapshot"]),"description":row["description"],"status":row["status"],"completed_date":row["completed_date"],"is_void":bool(row["is_void"])}

    def transfer_detail(self,number,include_void=False):
        try:no=int(number)
        except (TypeError,ValueError) as exc:raise AppError("TRANSFER_NOT_FOUND","Choose a valid transfer.") from exc
        row=self.repo.leg(no)
        if row is None or (row["is_void"] and not include_void):raise AppError("TRANSFER_NOT_FOUND","This transfer is unavailable.")
        cargo=self.db.execute("SELECT * FROM cargo WHERE id=?",(row["cargo_id"],)).fetchone()
        history=self.repo.cargo_legs(row["cargo_id"])
        events=self.db.execute("SELECT * FROM audit_events WHERE (entity_type='transfer' AND entity_id=?) OR (entity_type='cargo' AND entity_id=?) ORDER BY id DESC LIMIT 100",(no,cargo["id"])).fetchall()
        base=self.repo.rows("t.no=?",[no],"t.no",1,0)[0]
        location=self._cargo_location(row["cargo_id"])
        return {"row":self._row_dict(base),"cargo":{"id":cargo["id"],"cargo_tag":cargo["cargo_tag"],"physical_id":cargo["physical_id"] or "","created_utc":cargo["created_utc"]},"location":location,"legs":[self._simple_leg(item) for item in history],"events":[{"id":e["id"],"type":e["event_type"],"occurred_utc":e["occurred_utc"],"timezone":e["timezone_name"],"reason":e["reason"],"data":json.loads(e["data_json"])} for e in events],"descendants":[int(r["no"]) for r in self._descendants(no)]}

    def list_vessels(self):
        return [dict(r) for r in self.db.execute("""SELECT l.id,l.name,l.active,
          (SELECT count(*) FROM transfer_legs t WHERE t.is_void=0 AND t.status IN ('Waiting','On Board') AND (t.from_location_id=l.id OR t.to_location_id=l.id)) open_transfers,
          (SELECT count(*) FROM voyage_stops s JOIN voyages v ON v.id=s.voyage_id WHERE s.vessel_id=l.id AND s.removed=0 AND v.is_current=1) current_stops
          FROM locations l WHERE l.kind='VESSEL' AND NOT EXISTS(SELECT 1 FROM audit_events e WHERE e.entity_type='location' AND e.entity_id=l.id AND e.event_type='VESSEL_DELETED') ORDER BY l.active DESC,l.name COLLATE NOCASE""")]

    def voyage_history(self):
        voyages=[]
        for voyage in self.db.execute("SELECT id,name,is_current,archived,created_utc FROM voyages v WHERE NOT EXISTS(SELECT 1 FROM audit_events e WHERE e.entity_type='voyage' AND e.entity_id=v.id AND e.event_type='VOYAGE_DELETED') ORDER BY id DESC"):
            stops=[dict(row) for row in self.db.execute("SELECT s.id,s.vessel_id,s.vessel_name_snapshot,s.position,s.removed FROM voyage_stops s WHERE s.voyage_id=? ORDER BY s.position,s.id",(voyage["id"],))]
            voyages.append({**dict(voyage),"stops":stops})
        return voyages

    def cargo_choices(self):
        rows=self.db.execute("""SELECT c.id AS cargo_id,c.cargo_tag,c.default_description,c.physical_id,t.no AS last_no,
          t.status,t.to_location_id,t.to_name_snapshot
          FROM cargo c JOIN transfer_legs t ON t.cargo_id=c.id AND t.is_void=0
          WHERE t.no=(SELECT max(last.no) FROM transfer_legs last WHERE last.cargo_id=c.id AND last.is_void=0)
          AND t.status='Done' ORDER BY c.id DESC""")
        result=[dict(r) for r in rows]
        preserved=self._preserved_descriptions(row['last_no'] for row in result)
        shore=self.db.execute("SELECT id,name FROM locations WHERE kind='SHORE'").fetchone()
        for row in result:
            row['preserve_text_case']=preserved.get(row['last_no'],False)
            if row["to_location_id"]==shore["id"]:row["to_name_snapshot"]=shore["name"]
        return result

    def export_rows(self,view,filters,include_completed=False,*,order_column=0,descending=False):
        selected=dict(filters)
        if view=="all" and include_completed and selected.get("range","open")=="open":selected["range"]="all"
        where,args=self._where(view,selected)
        rows=self._row_dicts(self.repo.rows(where,args,self._sort_order(order_column,descending),1000000,0))
        return [{**row,"display_no":i} for i,row in enumerate(rows,1)]

    def relocated_service(self,destination):
        target=Path(destination).resolve()
        if target==self.path.resolve():return self
        self._check_output(target)
        target.parent.mkdir(parents=True,exist_ok=True)
        if target.exists():raise AppError("DATABASE_TARGET_EXISTS","Choose a folder without cargo-track.sqlite3. Existing data will not be overwritten.")
        prepared=target.with_name(target.name+"."+uuid4().hex+".relocate")
        replacement=None;created=False
        try:
            self.create_backup(prepared)
            # Exclusive creation refuses existing/concurrently created targets on all supported volumes.
            with target.open("xb") as out:
                created=True
                with prepared.open("rb") as source:shutil.copyfileobj(source,out)
                out.flush();os.fsync(out.fileno())
            replacement=ApplicationService(target)
            replacement.open()
            return replacement
        except Exception:
            if replacement is not None:replacement.close()
            if created:self._remove_temporary(target)
            raise
        finally:self._remove_temporary(prepared)

    def create_backup(self,destination):
        target=Path(destination);target.parent.mkdir(parents=True,exist_ok=True)
        self._check_output(target)
        temp=target.with_name(target.name+f".{uuid4().hex}.tmp")
        try:
            out=sqlite3.connect(str(temp))
            try:self.db.backup(out);out.commit()
            finally:out.close()
            failures=check_integrity(str(temp))
            if failures:raise AppError("BACKUP_INVALID","The backup failed integrity checks.",{"failures":failures})
            self._check_output(target)
            os.replace(temp,target)
        except (OSError,sqlite3.Error) as exc:
            raise AppError("BACKUP_FAILED","The backup could not be saved. The current ledger is unchanged.",{"detail":str(exc)}) from exc
        finally:self._remove_temporary(temp)
        return {"module":"Cargo","path":str(target),"transfers":self._count_file(target,"transfer_legs"),"created_utc":now_utc(),
                "identity":self.db.execute("SELECT identity FROM ledger_meta WHERE id=1").fetchone()[0]}

    def _check_output(self,path):
        try:workspace_files.check_output(path,self.path)
        except ProtectedFileError as exc:raise AppError("BACKUP_TARGET_ACTIVE_DATABASE",str(exc)) from exc

    @staticmethod
    def _remove_temporary(path):
        try:Path(path).unlink(missing_ok=True);return True
        except OSError:return False

    @staticmethod
    def _count_file(path,table):
        if table not in {"transfer_legs","cargo","locations","voyages"}:raise ValueError(table)
        conn=sqlite3.connect(Path(path).resolve().as_uri()+"?mode=ro",uri=True)
        try:return int(conn.execute(f"SELECT count(*) FROM {table}").fetchone()[0])
        finally:conn.close()

    def inspect_backup(self,candidate):
        path=Path(candidate)
        if not path.is_file() or path.resolve()==self.path.resolve():raise AppError("BACKUP_NOT_FOUND","Choose a separate existing backup file.")
        try:failures=check_integrity(str(path))
        except (sqlite3.Error,OSError) as exc:raise AppError("BACKUP_INVALID","The selected file is not a readable MarineOps database; the current ledger is unchanged.",{"detail":str(exc)}) from exc
        if failures:raise AppError("BACKUP_INVALID","The backup failed validation; the current ledger is unchanged.",{"failures":failures})
        conn=sqlite3.connect(Path(path).resolve().as_uri()+"?mode=ro",uri=True)
        try:
            counts={name:int(conn.execute(f"SELECT count(*) FROM {name}").fetchone()[0]) for name in ("cargo","transfer_legs","locations","voyages")}
            state=conn.execute("SELECT current_voyage_id,current_stop_id FROM app_state WHERE id=1").fetchone()
            return {"module":"Cargo","path":str(path.resolve()),"modified_utc":datetime.fromtimestamp(path.stat().st_mtime,UTC).isoformat(timespec="seconds"),"counts":counts,"current_voyage_id":state[0],"current_stop_id":state[1]}
        finally:conn.close()

    def restore_backup(self,candidate):
        preview=self.inspect_backup(candidate);stamp=datetime.now().strftime("%Y%m%d-%H%M%S")+"-"+uuid4().hex[:8]
        safety=self.path.with_name(f"{self.path.stem}-before-restore-{stamp}{self.path.suffix}")
        self.create_backup(str(safety))
        prepared=self.path.with_name(f".{self.path.name}.{uuid4().hex}.restore")
        old=self.path.with_name(f".{self.path.name}.{uuid4().hex}.rollback")
        try:
            shutil.copy2(candidate,prepared)
            if check_integrity(str(prepared)):raise AppError("BACKUP_INVALID","The selected file failed its final check.")
        except Exception as exc:
            self._remove_temporary(prepared)
            if isinstance(exc,AppError):raise
            raise AppError("RESTORE_PREPARE_FAILED","The restore could not be prepared. The current ledger is unchanged.",{"detail":str(exc),"safety_copy":str(safety)}) from exc
        self.close();moved=False
        try:
            os.replace(self.path,old);moved=True;os.replace(prepared,self.path);self.open()
        except Exception as exc:
            self.close()
            try:
                # Replace atomically; never delete the current path before rollback is available.
                if moved:os.replace(old,self.path)
                self.open()
            except Exception as recovery_error:
                self.close()
                raise AppError("RESTORE_RECOVERY_REQUIRED",
                    "The restore and automatic rollback both failed. Data copies were preserved; close MarineOps and recover the previous ledger from the listed files.",
                    {"detail":str(exc),"recovery_error":str(recovery_error),"safety_copy":str(safety),"previous_ledger":str(old if moved and old.exists() else self.path),"prepared_backup":str(prepared),"current_path":str(self.path)}) from recovery_error
            self._remove_temporary(prepared)
            raise AppError("RESTORE_FAILED","The restore failed. The previous ledger was reopened.",{"detail":str(exc),"safety_copy":str(safety)}) from exc
        removed=self._remove_temporary(old);self.dataset_epoch=str(uuid4());workspace_files.invalidate(self.path)
        return {"counts":preview["counts"],"safety_copy":str(safety),"restored_from":preview["path"],"revision":self._revision(),"retained_rollback":None if removed else str(old)}

    def export_to(self,destination,view,filters,file_format,include_completed=False,*,order_column=0,descending=False):
        records=self.export_rows(view,filters,include_completed,order_column=order_column,descending=descending)
        path=Path(destination);path.parent.mkdir(parents=True,exist_ok=True)
        self._check_output(path)
        columns=[("NO","display_no"),("Cargo ID","cargo_tag"),("Description","description"),("From","from_name"),("To","to_name"),("Status","status"),("Comment","comment"),("Completed Date","completed_date"),("Physical ID","physical_id")]
        temp=path.with_name(path.name+f".{uuid4().hex}.tmp"+(".xlsx" if file_format.lower()=="xlsx" else ""))
        if file_format.lower()=="csv":
            try:
                with temp.open("w",encoding="utf-8-sig",newline="") as stream:
                    writer=csv.writer(stream);writer.writerow([name for name,_ in columns])
                    for row in records:
                        values=[]
                        for _,key in columns:
                            value="" if row.get(key) is None else str(row[key])
                            if key!="display_no" and value.lstrip().startswith(("=","+","-","@")):value="'"+value
                            values.append(value)
                        writer.writerow(values)
                self._check_output(path);os.replace(temp,path)
            except Exception:temp.unlink(missing_ok=True);raise
        elif file_format.lower()=="xlsx":
            from .excel_layout import configure_print,text_height,populate_print_table
            from openpyxl import Workbook
            from openpyxl.styles import Alignment,Font,PatternFill
            from openpyxl.utils import get_column_letter
            book=Workbook();sheet=book.active;sheet.title="Transfers";sheet.append([name for name,_ in columns])
            for cell in sheet[1]:
                cell.font=Font(name='Calibri',size=10,bold=True,color="FFFFFF")
                cell.fill=PatternFill("solid",fgColor="173A63")
                cell.alignment=Alignment(vertical='center',wrap_text=True)
            for record in records:
                sheet.append([record.get(key) if record.get(key) is not None else "" for _,key in columns])
                for cell in sheet[sheet.max_row]:
                    if cell.column!=1 and isinstance(cell.value,str):cell.data_type="s"
                    cell.font=Font(name='Calibri',size=10)
                    cell.alignment=Alignment(vertical="top",wrap_text=True)
            for i,width in enumerate([5,12,30,13,13,11,25,13,13],1):sheet.column_dimensions[get_column_letter(i)].width=width
            sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
            sheet.row_dimensions[1].height=32
            overflow=False
            for cells in sheet.iter_rows(min_row=2):
                height=max(text_height(c.value or '',sheet.column_dimensions[c.column_letter].width,20) for c in cells)
                overflow=overflow or height>400
                sheet.row_dimensions[cells[0].row].height=min(409,height)
            configure_print(sheet,'I','landscape')
            if overflow:
                printable=book.create_sheet('Cargo Print')
                populate_print_table(printable,[name for name,_ in columns],
                    [[r.get(key) if r.get(key) is not None else '' for _,key in columns] for r in records],
                    (5,12,30,13,13,11,25,13,13),repeat_columns=(0,1,3,4,5,7,8))
                book.active=1
            try:book.save(temp);self._check_output(path);os.replace(temp,path)
            except Exception:temp.unlink(missing_ok=True);raise
        else:raise AppError("INVALID_EXPORT","Choose CSV or Excel format.")
        return {"path":str(path),"transfers":len(records),"format":file_format.lower()}
