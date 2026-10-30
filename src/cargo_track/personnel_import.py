"""Read crew workbooks without retaining ignored personal fields."""
from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from pathlib import Path
from itertools import islice
from contextlib import contextmanager

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .text_case import normalize, personnel_record
from .xlsx_safety import workbook_snapshot, WorkbookInputError


class PersonnelError(ValueError):
    """A recoverable personnel workflow error."""


def key(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     separators=(",", ":"), default=str).encode()).hexdigest()


def text(value):
    return str(value).strip() if value is not None else ""


def personnel_title(value, field='name'):
    """Use consistent title case for imported names and positions."""
    return normalize(value, field)[0]


def iso_date(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    value = text(value)
    for fmt in ("%Y-%m-%d", "%Y%m%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(value, fmt).date().isoformat()
        except ValueError:
            pass
    return ""


HEADERS = {
    "name": ("name", "姓名", "sign-off", "sign off", "sign-on", "sign on", "下船", "上船"),
    "position": ("position", "职位", "职务"),
    "nationality": ("nationality", "国籍"),
    "vessel": ("vessel", "船名"),
    "from_name": ("from", "来源", "from vessel"),
    "to_name": ("to", "目的地", "to vessel"),
    "date": ("date", "日期", "date of departure", "date of arrival", "离船时间", "到船日期"),
    "comment": ("comment", "comments", "备注"),
    "task_id": ("task id", "task_id", "任务id"),
}


def ignored_header(value):
    return bool(re.search(r"passport|iqama|visa|company|laptop|\bmark\b|\bid\s*(?:no\.?|number)\b|national\s*id|sr\s*no|serial|护照|证件|签证|公司|笔记本|标记", key(value)))


def header_field(value):
    normalized = key(value)
    if ignored_header(value):
        return None
    for field, labels in HEADERS.items():
        if normalized in labels:
            return field
    # A bilingual departure-date label contains 下船, which must not become a Name column.
    if re.search(r"(?:^|\s)date(?:$|\s)", normalized) or any(v in normalized for v in (
            "日期", "时间", "時間")):
        return "date"
    matches = [(len(label), field) for field, labels in HEADERS.items() for label in labels
               if label in normalized and (any('\u4e00' <= c <= '\u9fff' for c in label) or
                   (len(label) > 3 and label not in ("name", "date", "from", "to")))]
    if matches:
        return max(matches)[1]
    if re.search(r"(?:^|\s)name(?:$|\s)", normalized):
        return "name"
    return None


@contextmanager
def _open_workbook(path, *, data_only):
    try:
        with workbook_snapshot(path) as (frozen, file_hash):
            book = load_workbook(frozen, read_only=True, data_only=data_only)
            try:yield book, file_hash
            finally:book.close()
    except WorkbookInputError as exc:
        raise PersonnelError(str(exc)) from exc


def workbook_sheets(path):
    with _open_workbook(path, data_only=True) as (book, _):
        return book.sheetnames


def read_personnel_excel(path, sheet_name=None, mapping=None, header_row=None):
    """Return whitelisted records only. Mapping values are 1-based column numbers."""
    path = Path(path)
    with _open_workbook(path, data_only=False) as (book, file_hash):
        sheet = book[sheet_name] if sheet_name else book.active
        if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 200:
            raise PersonnelError("This workbook is too large. Use a sheet with at most 20,000 rows and 200 columns.")
        max_row,max_column=sheet.max_row or 20000,sheet.max_column or 200
        rows = list(islice(sheet.iter_rows(values_only=True),40))
        blocks = []
        if mapping is not None:
            blocks = [{"direction": mapping.get("direction", "Transfer"),
                       "columns": {f: int(c) for f, c in mapping.items() if f in HEADERS and c}}]
            header = int(header_row or 1)
            if not 1 <= header <= 20000 or not blocks[0]['columns'].get('name') or any(not 1<=c<=200 for c in blocks[0]['columns'].values()):
                raise PersonnelError("Use a header row up to 20,000 and column numbers between 1 and 200; Name is required.")
            header_values=next(sheet.iter_rows(min_row=header,max_row=header,values_only=True),())
            for field, column in blocks[0]["columns"].items():
                if column <= len(header_values) and ignored_header(header_values[column - 1]):
                    raise PersonnelError(f"Column {get_column_letter(column)} is an ignored personal field. Choose the {field} column instead.")
        else:
            header = 0
            for number, row in enumerate(rows[:40], 1):
                names = [(col, text(v)) for col, v in enumerate(row, 1)
                         if header_field(v) == "name"]
                if not names:
                    continue
                detected = []
                for offset, (name_col, label) in enumerate(names):
                    end = names[offset + 1][0] - 1 if offset + 1 < len(names) else len(row)
                    columns = {header_field(row[c - 1]): c for c in range(name_col, end + 1)
                               if header_field(row[c - 1]) is not None}
                    columns["name"] = name_col
                    direction = "Sign-off" if "sign-off" in key(label) or "sign off" in key(label) or "下船" in label else (
                        "Sign-on" if "sign-on" in key(label) or "sign on" in key(label) or "上船" in label else "Transfer")
                    endpoint = 'to_name' if direction == 'Sign-on' else 'from_name' if direction == 'Sign-off' else None
                    if "position" in columns and ("vessel" in columns or {"from_name", "to_name"} <= columns.keys()
                                                  or endpoint in columns):
                        detected.append({"direction": direction, "columns": columns})
                if detected:
                    header, blocks = number, detected
                    break
            if not blocks:
                raise PersonnelError("Could not recognize Name, Position and Vessel / From / To. Use manual column mapping.")
        result = []
        needed_column=max(c for block in blocks for c in block['columns'].values())
        for row_number, row in enumerate(sheet.iter_rows(min_row=header+1,max_row=max_row,max_col=needed_column,values_only=True),header+1):
            if row_number>20000:
                raise PersonnelError("This worksheet exceeds the 20,000-row limit.")
            for block in blocks:
                columns = block["columns"]
                values = {field: row[col - 1] if col <= len(row) else None for field, col in columns.items()}
                if not any(text(values.get(f)) for f in ("name", "position", "vessel", "from_name", "to_name", "date", "task_id")):
                    continue
                # Header and laptop subheader rows have no person name in the sample.
                if header_field(values["name"]) == "name":
                    continue
                record = {field: text(values.get(field)) for field in (
                    "name", "position", "nationality", "vessel", "from_name", "to_name", "comment", "task_id")}
                record = personnel_record(record)
                record.update(direction=block["direction"], date=iso_date(values.get("date")),
                              source_date=iso_date(values.get("date")), source_row=row_number,
                              source_sheet=sheet.title,
                              source_cells={f: f"{get_column_letter(c)}{row_number}" for f, c in columns.items()},
                              original_values={f: text(v) for f, v in values.items()})
                record["issues"] = [f"Missing {f}" for f in ("name", "position", "date") if not record[f]]
                if any(isinstance(v, str) and v.startswith("=") for v in values.values()):
                    record["issues"].append("Formula in an imported field; replace with a confirmed value")
                result.append(record)
        if not result:
            raise PersonnelError("No personnel records were found in this sheet.")
        canonical = [{f: r.get(f, "") for f in ("name", "position", "nationality", "vessel", "from_name",
                      "to_name", "date", "direction", "comment", "task_id")} for r in result]
        return {"file_name": path.name, "file_hash": file_hash,
                "content_hash": digest(canonical), "sheet": sheet.title, "header_row": header,
                "parser_version":1,"mapping_hash":digest({'header':header,'blocks':blocks}),
                "blocks": blocks, "records": result}
