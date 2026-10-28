"""Generate event-faithful berth schedules and export immutable versions."""
from __future__ import annotations

import copy
import os
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

from .personnel_import import PersonnelError, digest, iso_date, key
from .personnel_cargo import dependency_spec, dependency_hash
from .coordination import workspace_files, ProtectedFileError
from .personnel_routes import transfer_name, transfer_origin, person_identity, leg_order
from .excel_layout import text_height, split_text, configure_print, populate_print_table

GROUPS = (("drop_passengers", "Drop", "passengers"), ("pick_passengers", "Pick up", "passengers"),
          ("drop_materials", "Drop", "Materials"), ("pick_materials", "Pick up", "Materials"))


def task_key(task):
    return f"{task['kind']}:{task['id']}"


def plan_review_groups(plan):
    """One actionable row per record, without merging different people by name."""
    groups = {}
    for issue in plan.get("issues", []):
        identifier = issue.get("task", "")
        group = groups.setdefault(identifier, {"task": identifier, "name": issue.get("name") or "Voyage", "messages": []})
        if issue["message"] not in group["messages"]:
            group["messages"].append(issue["message"])
    return list(groups.values())


def plan_task_selection(plan):
    if "selected_tasks" in plan:
        return set(plan["selected_tasks"])
    return set(plan.get("included", [])) | {i["task"] for i in plan.get("issues", []) if i.get("task")}


def plan_assignments(plan):
    result = copy.deepcopy(plan.get("assignments", {}))
    for stop in plan["stops"]:
        for group, side in (("pick_passengers", "pick"), ("drop_passengers", "drop"),
                            ("pick_materials", "pick"), ("drop_materials", "drop")):
            for task in stop[group]:
                result.setdefault(task_key(task), {"pick": None, "drop": None})[side] = stop["id"]
    return result


def cargo_tasks(snapshot, numbers):
    wanted = set(numbers)
    found = set()
    result = []
    for row in snapshot["cargo"]:
        if row["no"] not in wanted:
            continue
        found.add(row["no"])
        result.append({"kind": "Materials", "id": str(row["no"]), "cargo_id": row["cargo_id"],
                       "cargo_tag": row["cargo_tag"], "name": row["description"], "position": "",
                       "from_id": row["from_location_id"], "to_id": row["to_location_id"],
                       "from_name": row["from_name_snapshot"], "to_name": row["to_name_snapshot"],
                       "status": row["status"], "is_void": bool(row["is_void"]), "comment": row["comment"],
                       "pickup_utc": row.get("pickup_utc", ""), "issues": []})
    if found != wanted:
        raise PersonnelError("A selected cargo transfer disappeared. Refresh the cargo selection.")
    return result


def route_options(snapshot, task, mode="Complete", start_stop_id=None, followup=False):
    stops = snapshot["stops"]
    if mode == 'Final' and snapshot.get('transport_contract'):
        from .berth_optimizer import route_options as strict_options
        return strict_options(snapshot['transport_contract'], task, stops)
    if mode == "Remaining":
        start_stop_id = start_stop_id or snapshot["current_stop_id"]
        found = next((i for i, s in enumerate(stops) if s["id"] == start_stop_id), None)
        if found is None:
            return []
        stops = stops[found:]
    shore = next((r["id"] for r in snapshot["locations"] if r["kind"] == "SHORE"), None)
    ports = [s for s in stops if s['vessel_id'] == shore]
    if (mode == 'Final' and snapshot.get('require_personnel_port_calls') and
            task.get('kind', 'Passengers') == 'Passengers' and task.get('from_id') == shore and not ports):
        return []
    # A real port call is the pickup boundary for first sign-on legs too.
    # Preserve implicit departure only for old routes with no port at all.
    if mode == 'Final' and task.get('kind', 'Passengers') == 'Passengers' and ports:
        order = {s['id']: i for i, s in enumerate(stops)}
        if task.get('from_id') == shore and not task.get('previous_task_id'):
            return [{'pick': p['id'], 'drop': d['id']} for p in ports for d in stops
                    if d['vessel_id'] == task['to_id'] and order[p['id']] < order[d['id']]
                    and not any(order[p['id']] < order[q['id']] < order[d['id']] for q in ports)]
        if task.get('to_id') == shore:
            return [{'pick': s['id'], 'drop': p['id']} for s in stops for p in ports
                    if s['vessel_id'] == task['from_id'] and order[s['id']] < order[p['id']]
                    and not any(order[s['id']] < order[q['id']] < order[p['id']] for q in ports)]
    if followup:
        return [{"pick": "PORT", "drop": None}] if task["from_id"] == shore and task["status"] == "Waiting" else []
    targets = [s["id"] for s in stops if s["vessel_id"] == task["to_id"]]
    if mode != 'Final' and task["to_id"] == shore:
        targets = ["PORT"]
    later_leg = mode == 'Final' and task.get('kind', 'Passengers') == 'Passengers' and task.get('previous_task_id')
    if not later_leg and (task["status"] == "On Board" or task["from_id"] == shore or transfer_origin(task, snapshot)):
        return [{"pick": None, "drop": target} for target in targets]
    sources = [s["id"] for s in stops if s["vessel_id"] == task["from_id"]]
    if mode == "Final" and task.get("kind") == "Materials" and task["status"] == "Waiting" and not targets:
        # A vessel call can collect materials even when delivery is on a later route.
        return [{"pick": source, "drop": None} for source in sources]
    order = {s["id"]: i for i, s in enumerate(stops)} | {"PORT": len(stops)}
    return [{"pick": source, "drop": target} for source in sources for target in targets if order[source] < order[target]]


def leg_choices(snapshot, task, resolved, selected, mode='Final', start_stop_id=None, followup=False):
    choices = route_options(snapshot, task, mode, start_stop_id, followup)
    previous = task.get('previous_task_id')
    if mode != 'Final' or not previous:
        return choices
    if 'Passengers:' + previous not in selected:
        return []
    earlier = resolved.get('Passengers:' + previous)
    if not earlier or earlier.get('drop') is None:
        return []
    order = {s['id']: i for i, s in enumerate(snapshot['stops'])}
    return [choice for choice in choices if choice['pick'] is not None
            and order[choice['pick']] > order[earlier['drop']]]


def generate_plan(snapshot, personnel, selected_cargo, plan_date, port_name, mode="Complete", start_stop_id=None,
                  assignments=None, followups=None, confirmed_conflicts=False, base_plan=None, personnel_plan_id=None,
                  selections=None, transport_contract=None):
    if transport_contract is not None:
        snapshot = {**snapshot, 'transport_contract': copy.deepcopy(transport_contract)}
    if not iso_date(plan_date) or not port_name.strip():
        raise PersonnelError("Enter a valid plan date and Port name.")
    if mode not in {"Complete", "Remaining", "Final"}:
        raise PersonnelError("Choose Complete or Remaining plan.")
    if base_plan is not None:
        if base_plan["mode"] != "Complete":
            raise PersonnelError("Choose an original Complete plan as the base.")
        draft = copy.deepcopy(base_plan)
        for field in ("id", "version", "status"):
            draft.pop(field, None)
        draft["base_version_id"] = base_plan["id"]
        return draft
    assignments, followups = assignments or {}, set(followups or [])
    all_stops = snapshot["stops"]
    issues, warnings = [], []
    if not snapshot["voyage_id"]:
        issues.append({"task": "", "message": "No Voyage selected in MarineOps"})
    if not all_stops and not snapshot.get('transport_contract'):
        issues.append({"task": "", "message": "Voyage has no vessel stops"})
    if snapshot.get("archived"):
        issues.append({"task": "", "message": "An archived Voyage cannot be used for new work"})
    route = all_stops
    if mode == "Remaining":
        start_stop_id = start_stop_id or snapshot["current_stop_id"]
        start = next((i for i, s in enumerate(all_stops) if s["id"] == start_stop_id), None)
        if start is None:
            issues.append({"task": "", "message": "Choose the actual starting stop for Remaining plan"})
            route = []
        else:
            route = all_stops[start:]
    elif mode == "Complete" and all_stops and snapshot["current_stop_id"] and snapshot["current_stop_id"] != all_stops[0]["id"]:
        issues.append({"task": "", "message": "Voyage is in progress. Load its saved Complete version or use Remaining plan"})
    numbers = {s["id"]: i + 1 for i, s in enumerate(all_stops)}
    stops = [{"id": s["id"], "number": numbers[s["id"]], "name": s["vessel_name_snapshot"],
              "vessel_id": s["vessel_id"], **{g: [] for g, _, _ in GROUPS}} for s in route]
    shore = next((r["id"] for r in snapshot["locations"] if r["kind"] == "SHORE"), None)
    if mode != 'Final':
        stops.append({"id": "PORT", "number": len(all_stops) + 1, "name": port_name.strip(), "vessel_id": shore,
                      **{g: [] for g, _, _ in GROUPS}})
    by_id = {s["id"]: s for s in stops}
    tasks = [{**copy.deepcopy(r), "kind": "Passengers"} for r in personnel] + cargo_tasks(snapshot, selected_cargo)
    tasks.sort(key=leg_order)
    task_map = {t['id']: t for t in tasks if t['kind'] == 'Passengers'}
    selected = {task_key(t) for t in tasks if not selections or selections.get(task_key(t), {}).get('state', 'included') == 'included'}
    resolved = {}
    included, excluded = [], []
    names = {}
    for task in tasks:
        identifier = task_key(task)
        selection=(selections or {}).get(identifier)
        if selection and selection['state']!='included':
            excluded.append({'task':identifier,'name':task['name'],'status':task['status'],
                             'selection':selection['state'],'reason':selection.get('reason','')})
            if selection['state']=='unavailable':
                reasons=task.get('issues') or [selection.get('reason') or 'No feasible route; review this entry or explicitly exclude it.']
                issues.extend({'task':identifier,'name':task['name'],'message':p} for p in reasons)
            continue
        if mode == "Final" and task["kind"] == "Passengers":
            task["status"] = "Waiting"
        if task["status"] in {"Done", "Cancelled"} or task.get("is_void"):
            excluded.append({"task": identifier, "name": task["name"], "status": task["status"]})
            if mode == "Complete" and task["status"] == "Done":
                issues.append({"task": identifier, "message": "Completed task selected without original Complete version"})
            continue
        problems = list(task.get("issues", []))
        if mode == "Final" and task["kind"] == "Passengers":
            problems = [p for p in problems if p != "Missing date"]
        if not task.get("name") or (task["kind"] == "Passengers" and not task.get("position")):
            problems.append("Missing name or position")
        onboard_source = transfer_origin(task, snapshot)
        if (not task.get("from_id") and not onboard_source) or not task.get("to_id") or task.get("from_id") == task.get("to_id"):
            problems.append("Missing or identical From / To")
        previous = task_map.get(task.get('previous_task_id'))
        if task.get('previous_task_id'):
            if not previous or person_identity(previous) != person_identity(task) or leg_order(previous) + 1 != leg_order(task):
                problems.append('The preceding personnel leg is missing or its identity changed; review this person.')
            elif previous.get('to_id') != task.get('from_id'):
                problems.append('This leg must start at the preceding destination; review From / To.')
            elif mode == 'Final' and 'Passengers:' + task['previous_task_id'] not in selected:
                problems.append('Include the preceding personnel leg before planning this next leg.')
        if task["kind"] == "Passengers":
            if mode != "Final" and task.get("date") != iso_date(plan_date) and identifier not in followups:
                problems.append("Passenger transfer date does not match this plan")
            name_key = key(task["name"])
            if name_key in names and not confirmed_conflicts and mode != "Final":
                problems.append("Repeated passenger name; confirm distinct people or resolve conflicting tasks")
            names[name_key] = identifier
        if mode == "Remaining" and task.get("from_id") == shore and task["status"] == "Waiting" and identifier not in followups:
            problems.append("Port-origin passenger / cargo is still Waiting; verify actual pickup before Remaining plan")
        if problems:
            issues.extend({"task": identifier, "message": p} for p in dict.fromkeys(problems))
            continue
        choices = leg_choices(snapshot, task, resolved, selected, mode, start_stop_id, identifier in followups)
        assignment = assignments.get(identifier)
        if assignment is not None and assignment not in choices:
            issues.append({"task": identifier, "message": "Saved stop assignment is no longer valid"})
            continue
        if assignment is None and len(choices) == 1:
            assignment = choices[0]
        if assignment is None and choices and mode == "Final":
            assignment = choices[0]
        if assignment is None:
            if not choices and mode == "Final":
                visited = {s["vessel_id"] for s in route}
                missing = []
                if task['kind'] == 'Passengers' and task['from_id'] == shore:
                    missing.append('Add or move Port before this passenger destination to collect Sign-on passengers.')
                if not onboard_source and task["from_id"] != shore and task["status"] != "On Board" and task["from_id"] not in visited:
                    missing.append(f"Add source vessel {task['from_name']} to this plan before its destination.")
                if task["to_id"] not in visited:
                    missing.append(f"Add destination vessel {task['to_name']} to this plan.")
                if task.get('previous_task_id') and 'Passengers:' + task['previous_task_id'] in selected:
                    missing.append('Pick this person up on a later visit after the preceding drop; add or choose the next visit.')
                issues.extend({"task": identifier, "message": reason} for reason in
                    (missing or ["Move the source vessel before the destination, or add another visit."]))
            else:
                issues.append({"task": identifier, "message": "Select a specific repeated stop" if choices else "Route cannot pick up before drop / destination is outside this route"})
            continue
        if mode == "Complete" and task["status"] == "On Board" and task["from_id"] != shore and task.get("pickup_utc") and task["pickup_utc"] >= snapshot.get("voyage_created_utc", ""):
            issues.append({"task": identifier, "message": "Task already picked up in this Voyage; use saved Complete version or Remaining plan"})
            continue
        # Persist only useful task fields, never runtime DB payloads or ignored personal fields.
        item = {f: task.get(f, "") for f in ("id", "kind", "name", "position", "from_id", "to_id", "from_name",
                "to_name", "cargo_id", "cargo_tag", "status", "comment", "date", 'direction', 'person_id', 'leg_index', 'previous_task_id')}
        item["task_key"] = identifier
        if task['kind']=='Passengers':item['leg_index']=leg_order(task)
        item["followup"] = identifier in followups
        for side in ("from", "to"):
            if item[side + "_id"] == shore:
                item[side + "_name"] = port_name.strip()
        suffix = "passengers" if task["kind"] == "Passengers" else "materials"
        if assignment["pick"] is not None:
            by_id[assignment["pick"]]["pick_" + suffix].append(copy.deepcopy(item))
        if assignment["drop"] is not None:
            by_id[assignment["drop"]]["drop_" + suffix].append(copy.deepcopy(item))
        included.append(identifier)
        resolved[identifier] = assignment
        if mode != "Final" and task["from_id"] == shore and task["status"] == "Waiting" and identifier not in followups:
            warnings.append({"task": identifier, "message": "Port pickup must be confirmed before actual Drop"})
    labels = {task_key(t): t["name"] for t in tasks}
    for problem in issues + warnings:
        problem["name"] = labels.get(problem["task"], "Voyage")
    dependencies=dependency_spec(snapshot,personnel,selected_cargo)
    dependencies['mode']=mode
    personnel_dependencies=dependency_spec(snapshot,personnel,[])
    personnel_dependencies['mode']=mode
    result = {"date": iso_date(plan_date), "mode": mode, "start_stop_id": start_stop_id,
            "personnel_plan_id": personnel_plan_id, "voyage_id": snapshot["voyage_id"], "voyage_name": snapshot["voyage_name"],
            "port_name": port_name.strip(), "transfer_vessel": transfer_name(snapshot),
            "source_path": snapshot["source_path"], "source_hash": snapshot["source_hash"], "source_revision": snapshot["revision"],
            "source_identity":snapshot.get('ledger_identity'),"source_generation":snapshot.get('source_generation',0),
            "source_dependencies":dependencies,"dependency_hash":dependency_hash(snapshot,dependencies),
            "personnel_dependencies":personnel_dependencies,"personnel_dependency_hash":dependency_hash(snapshot,personnel_dependencies),
            "personnel_hash": digest(personnel), "stops": stops, "issues": issues, "warnings": warnings,
            "personnel_versions": {r['id']: r['version'] for r in personnel},
            "included": included, "excluded": excluded,
            "selected_tasks": [task_key(t) for t in tasks if not selections or selections.get(task_key(t),{}).get('state','included')=='included'],
            "task_selections":copy.deepcopy(selections or {}),
            "assignments": assignments, "followups": sorted(followups),
            "plan_route_hash": digest(stops)}
    if snapshot.get('transport_contract'):
        from .berth_optimization_validation import validate_saved_plan
        result['transport_contract'] = copy.deepcopy(snapshot['transport_contract'])
        try:validate_saved_plan(result)
        except PersonnelError as exc:result['issues'].append({'task': '', 'name': 'Complete transport', 'message': str(exc)})
    if snapshot.get('require_personnel_port_calls'):
        result['require_personnel_port_calls'] = True
    return result


def revise_complete(base, pending, snapshot, live_personnel):
    """Keep evidenced past actions; replace only future actions with reviewed selections."""
    if base["mode"] != "Complete" or not base.get("id"):
        raise PersonnelError("Choose a saved Complete version.")
    if base["source_path"] != snapshot["source_path"] or base["voyage_id"] != snapshot["voyage_id"]:
        raise PersonnelError("The saved version belongs to another Cargo ledger or Voyage.")
    if [s["id"] for s in base["stops"] if s["id"] != "PORT"] != [s["id"] for s in snapshot["stops"]]:
        raise PersonnelError("Voyage stop order changed. Keep the original Complete version and generate a Remaining plan for the new route.")
    if pending["mode"] != "Remaining" or pending["date"] != base["date"]:
        raise PersonnelError("Review pending actions for the original Complete plan date.")
    live = {task_key({**r, "kind": "Passengers"}): r for r in live_personnel}
    live.update({task_key(r): r for r in cargo_tasks(snapshot, [r["no"] for r in snapshot["cargo"]])})
    result = copy.deepcopy(base)
    for field in ("id", "version", "status"):
        result.pop(field, None)
    retained = set()
    for stop in result["stops"]:
        for group, _, _ in GROUPS:
            keep = []
            for item in stop[group]:
                current = live.get(item["task_key"])
                if current is None:
                    raise PersonnelError(f"Original task is missing: {item['name']}. Restore its source before revising the Complete plan.")
                if current["status"] == "Done" or (current["status"] == "On Board" and group.startswith("pick_")):
                    keep.append(item)
                    retained.add(item["task_key"])
            stop[group] = keep
    by_id = {s["id"]: s for s in result["stops"]}
    for stop in pending["stops"]:
        for group, _, _ in GROUPS:
            by_id[stop["id"]][group].extend(copy.deepcopy(stop[group]))
    # Historical item labels stay frozen; the new terminal heading uses reviewed settings.
    by_id["PORT"]["name"] = pending["port_name"]
    for field in ("source_hash", "source_revision", "personnel_hash", "issues", "warnings", "assignments", "followups", "excluded", "port_name", "transfer_vessel", "source_dependencies", "dependency_hash", "personnel_dependencies", "personnel_dependency_hash", "source_identity", "source_generation"):
        result[field] = copy.deepcopy(pending[field])
    result["included"] = sorted(retained | set(pending["included"]))
    result["personnel_versions"] = dict(pending["personnel_versions"])
    for identifier in retained:
        if identifier.startswith("Passengers:"):
            record = live[identifier]
            result["personnel_versions"][record["id"]] = record["version"]
    result.update(base_version_id=base["id"], revision_reviewed=True, pending_start_stop_id=pending["start_stop_id"])
    return result


def item_text(item, action):
    if item["kind"] == "Passengers":
        value = f"{item['name']}({item['position']})"
    else:
        side = "from" if action == "Drop" else "to"
        value = f"{item['name']} {side} {item[side + '_name']}"
    return value + (" [Next voyage / follow-up]" if item.get("followup") else "")


def plan_text(plan,compact=False,personnel_only=False):
    parts = [f"{plan.get('status', 'Unsaved Draft')} · {plan['mode']} · {plan['date']} · {plan['voyage_name']}"]
    for stop in plan["stops"]:
        parts.append(f"\n{stop['number']:02d}. {stop['name']}")
        for group, action, noun in GROUPS:
            if personnel_only and group.endswith('materials'):continue
            items = stop[group]
            entries = [f"{i}. {item_text(item, action)}" for i, item in enumerate(items, 1)]
            parts.append(f"{action} {len(items)} {noun}:" + (" " + "; ".join(entries) if entries else ""))
    if plan["issues"]:
        parts.append("\nUnassigned / requires review:")
        parts.extend(f"{p.get('name', p['task'])}: {p['message']}" for p in plan["issues"] if not personnel_only or not p['task'].startswith('Materials:'))
        if personnel_only and any(p['task'].startswith('Materials:') for p in plan['issues']):
            parts.append('Linked Cargo route requires review before export. Review the source route in Cargo.')
    if plan.get("warnings"):
        parts.append("\nExecution prerequisites:")
        parts.extend(f"{p.get('name', p['task'])}: {p['message']}" for p in plan["warnings"] if not personnel_only or not p['task'].startswith('Materials:'))
    if plan.get('excluded'):
        parts.append('\nScope exclusions:')
        parts.extend(f"{p['name']}: {p.get('selection',p.get('status',''))} — {p.get('reason','')}" for p in plan['excluded'] if not personnel_only or p.get('kind')=='Passengers')
    return "\n".join(parts)


def export_plan(plan, destination):
    from .berth_optimization_validation import validate_saved_plan
    validate_saved_plan(plan)
    if not plan.get("id"):
        raise PersonnelError("Save a plan version before exporting.")
    if plan.get("status") == "Published" and plan["issues"]:
        raise PersonnelError("Published plans cannot contain unresolved issues.")
    path = Path(destination).resolve()
    try:workspace_files.check_output(path,plan.get('source_path'),plan.get('personnel_source_path'))
    except ProtectedFileError as exc:raise PersonnelError(str(exc)) from exc
    if path.suffix.casefold() != ".xlsx":
        path = path.with_suffix(".xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.title = "Berthing Plan"
    detail = book.create_sheet("Transfer Details")
    printable = book.create_sheet("Transfer Print")
    for ws in (sheet, detail, printable):
        ws.sheet_view.showGridLines = False
    for col, width in zip('ABCD', (4, 20, 60, 16)):
        sheet.column_dimensions[col].width = width
    row = 2
    border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                    top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))

    def style(cell, bold=False, center=False):
        cell.font = Font(name='Times New Roman', size=10, bold=bold, color='000000')
        cell.alignment = Alignment(horizontal='center' if center else 'left', vertical='center', wrap_text=True)
        cell.border = border
        if isinstance(cell.value, str):
            cell.data_type = 's'

    def write(value, fill=None, bold=False):
        nonlocal row, page_height
        width = sum(sheet.column_dimensions[c].width for c in 'ABCD')
        for part in split_text(value, width, 180):
            height = text_height(part, width, 20, 15)
            if page_height + height > 760:
                sheet.row_breaks.append(Break(id=row-1));page_height = header_height
            sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            cell = sheet.cell(row, 1, part);cell.data_type = 's'
            cell.font = Font(name="Arial", size=11 if bold else 10, bold=bold, color="24262B")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:cell.fill = PatternFill("solid", fgColor=fill)
            sheet.row_dimensions[row].height = height
            row += 1;page_height += height

    from datetime import date
    header_values = ('', date.fromisoformat(plan['date']), (plan.get('transfer_vessel') or plan['voyage_name']) + ' plan', 'Comments')
    for column, value in enumerate(header_values, 1):
        cell = sheet.cell(1, column, value)
        style(cell, bold=True, center=True)
        cell.fill = PatternFill('solid', fgColor='F58220')
    sheet['B1'].number_format = '[$-409]dd mmmm yyyy'
    header_height = max(23, *(text_height(str(v), sheet.column_dimensions[c].width) for c,v in zip('BCD',header_values[1:])))
    sheet.row_dimensions[1].height = header_height
    sheet.freeze_panes = 'C2'
    page_height = header_height
    for stop in plan["stops"]:
        segments, current, current_height = [], [], 0
        for group, action, noun in GROUPS:
            items = stop[group]
            entries = '; '.join(f'{i}. {item_text(item, action)}' for i, item in enumerate(items, 1))
            label = 'passengers' if noun == 'passengers' else ('material' if len(items) == 1 else 'materials')
            text = f'{action} {len(items)} {label}:' + (' ' + entries if entries else '')
            for part in split_text(text, sheet.column_dimensions['C'].width, 200, f'{action} {label} (continued): '):
                height = text_height(part, sheet.column_dimensions['C'].width)
                if current and current_height + height > 720:
                    segments.append(current);current = [];current_height = 0
                current.append((part,height));current_height += height
        if current:segments.append(current)
        comments = split_text(stop.get('comment',''), sheet.column_dimensions['D'].width, 700)
        while len(segments) < len(comments):segments.append([('',21)] * 4)
        for index, segment in enumerate(segments):
            heading = 'Full speed to ' + stop['name'] + (' (continued)' if index else '')
            comment = comments[index] if index < len(comments) else ''
            block_height = max(sum(h for _,h in segment), text_height(heading,sheet.column_dimensions['B'].width),
                               text_height(comment,sheet.column_dimensions['D'].width))
            extra = block_height - sum(h for _,h in segment)
            # Merged comments may need more room; keep individual Excel rows
            # below the 409-point limit rather than clipping or rejecting a stop.
            for i,(text,height) in enumerate(segment):
                growth = min(extra,200-height);segment[i] = (text,height+growth);extra -= growth
            while extra > 0:
                height = min(200,extra);segment.append(('',height));extra -= height
            if (index or page_height + block_height > 760) and page_height > header_height:
                sheet.row_breaks.append(Break(id=row-1));page_height = header_height
            start = row
            for value,height in segment:
                cell = sheet.cell(row,3,value);style(cell)
                sheet.row_dimensions[row].height = height;row += 1
            for column,value in ((1,stop['number']),(2,heading),(4,comment)):
                sheet.merge_cells(start_row=start,start_column=column,end_row=row-1,end_column=column)
                cell = sheet.cell(start,column,value);style(cell,center=column != 4)
                for cells in sheet.iter_rows(min_row=start,max_row=row-1,min_col=column,max_col=column):cells[0].border = border
            page_height += block_height
    if plan["issues"]:
        row += 1
        write("DRAFT — Unassigned tasks / requires review", fill="FFF2D6", bold=True)
        for issue in plan["issues"]:
            write(f"{issue.get('name', issue['task'])}: {issue['message']}")
    if plan.get("warnings"):
        write("Execution prerequisites", bold=True)
        for warning in plan["warnings"]:
            write(f"{warning.get('name', warning['task'])}: {warning['message']}")
    if plan.get('excluded'):
        write('Scope exclusions',bold=True)
        for excluded in plan['excluded']:
            write(f"{excluded['name']}: {excluded.get('selection',excluded.get('status',''))} — {excluded.get('reason','')}")
    configure_print(sheet,'D',last_row=row-1)
    sheet.oddFooter.left.text = f"{plan['status']} v{plan['version']}"
    headings = ["Stop No.", "Location", "Action", "Type", "Name / Description", "Position", "From", "To",
                "Task ID", "Cargo ID", "Comment", "Follow-up", "Plan date", "Mode", "Version", "Plan status",
                'Person ID', 'Leg', 'Previous task ID']
    detail.append(headings)
    print_rows = []
    for stop in plan["stops"]:
        for group, action, _ in GROUPS:
            for item in stop[group]:
                values = [stop["number"], stop["name"], action, item["kind"], item["name"], item["position"],
                          item["from_name"], item["to_name"], item["id"], item.get("cargo_tag", ""), item.get("comment", ""),
                          "Next voyage" if item.get("followup") else "", iso_date(plan["date"]), plan["mode"], plan["version"], plan["status"],
                          item.get('person_id', ''), item.get('leg_index', ''), item.get('previous_task_id', '')]
                detail.append(values)
                print_rows.append([values[i] for i in (0,1,2,3,4,5,6,7,10,17)])
                for cell, value in zip(detail[detail.max_row], values):
                    if isinstance(value, str):
                        cell.data_type = "s"
                # Store dates as real dates rather than display strings.
                from datetime import date
                detail.cell(detail.max_row, 13, date.fromisoformat(plan["date"])).number_format = "yyyy-mm-dd"
    detail.freeze_panes = "E2"
    detail.auto_filter.ref = detail.dimensions
    for col, width in zip("ABCDEFGHIJKLMNOPQRS", (6, 17, 10, 11, 38, 16, 17, 17, 38, 18, 40, 18, 15, 15, 10, 16, 38, 8, 38)):
        detail.column_dimensions[col].width = width
    for cells in detail:
        for cell in cells:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in detail[1]:
        cell.font = Font(name="Arial", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="EAF3FF")
    detail.row_dimensions[1].height = 30
    for cells in detail.iter_rows(min_row=2):
        detail.row_dimensions[cells[0].row].height = min(409,max(text_height(c.value or '',detail.column_dimensions[c.column_letter].width,20) for c in cells[:8]))
    configure_print(detail,'H','landscape')
    populate_print_table(printable,
        ['Stop','Location','Action','Type','Name / Description','Position','From','To','Comment','Leg'],
        print_rows,(5,12,9,11,30,14,13,13,22,5),repeat_columns=(0,1,2,3,9))
    printable.oddFooter.left.text = f"{plan['date']} / {plan['status']} v{plan['version']}"
    temp = path.with_name(path.stem + "." + uuid4().hex + ".tmp.xlsx")
    try:
        book.save(temp)
        try:workspace_files.check_output(path,plan.get('source_path'),plan.get('personnel_source_path'))
        except ProtectedFileError as exc:raise PersonnelError(str(exc)) from exc
        os.replace(temp, path)
    finally:
        book.close()
        temp.unlink(missing_ok=True)
    return str(path)
