"""Deterministic wrapping and physical A4 layouts for offline XLSX exports."""
import unicodedata


def text_height(value, width, minimum=21, line_points=13):
    capacity = max(1, int(width - 2))
    lines = 0
    def weight(word):
        return sum(2 if unicodedata.east_asian_width(c) in {'F','W'} or c in 'WM@' else 1 for c in word)
    for paragraph in str(value).expandtabs(4).split('\n'):
        used = 0;lines += 1
        for word in paragraph.split(' '):
            length = weight(word)
            if used and used + 1 + length > capacity:
                lines += 1;used = 0
            elif used:used += 1
            while length > capacity:
                lines += 1;length -= capacity
            used += length
    return max(minimum, lines * line_points + 8)


def split_text(value, width, max_height=200, continuation=''):
    """Keep every source character; prefer item/word boundaries for continuations."""
    remaining = str(value);parts = [];prefix = ''
    while remaining:
        if len(prefix + remaining) <= 30000 and text_height(prefix + remaining, width) <= max_height:
            parts.append(prefix + remaining);break
        low, high = 1, min(len(remaining), 30000-len(prefix))
        while low < high:
            mid = (low + high + 1)//2
            if text_height(prefix + remaining[:mid], width) <= max_height:low = mid
            else:high = mid-1
        cut = low
        boundary = remaining.rfind('; ', 0, cut)
        if boundary > cut//2:cut = boundary + 2
        else:
            boundary = max(remaining.rfind(' ', 0, cut), remaining.rfind('\n', 0, cut))
            if boundary > cut//2:cut = boundary+1
        parts.append(prefix + remaining[:cut]);remaining = remaining[cut:];prefix = continuation
    return parts or ['']


def configure_print(sheet, last_column, orientation='portrait', last_row=None):
    # In the verified Excel/PDF environment, reduced Zoom expands PDF MediaBox.
    # Fixed 100% plus bounded column widths gives real A4 without tiny fonts.
    sheet.page_setup.orientation = orientation
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.scale = 100
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = False
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = '1:1'
    sheet.print_area = f'A1:{last_column}{last_row or sheet.max_row}'
    sheet.page_margins.left = sheet.page_margins.right = .2
    sheet.page_margins.top = sheet.page_margins.bottom = .3
    sheet.oddFooter.center.text = 'Page &P of &N'


def populate_print_table(sheet, headings, rows, widths, repeat_columns=()):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.pagebreak import Break
    sheet.append(headings)
    for i,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width = width
    sheet.row_dimensions[1].height = 34
    page_height = 34
    for values in rows:
        chunks = [split_text('' if v is None else v,width,164) for v,width in zip(values,widths)]
        for part in range(max(len(parts) for parts in chunks)):
            data = [parts[part] if part < len(parts) else (parts[0] if i in repeat_columns else '') for i,parts in enumerate(chunks)]
            height = max(text_height(v,w,20) for v,w in zip(data,widths))
            if page_height + height > 530 and page_height > 34:
                sheet.row_breaks.append(Break(id=sheet.max_row));page_height = 34
            sheet.append(data);sheet.row_dimensions[sheet.max_row].height = height;page_height += height
    for cells in sheet:
        for cell in cells:
            cell.font = Font(name='Arial',size=10,bold=cell.row == 1)
            cell.alignment = Alignment(vertical='top',wrap_text=True)
            if isinstance(cell.value,str):cell.data_type = 's'
            if cell.row == 1:cell.fill = PatternFill('solid',fgColor='EAF3FF')
    sheet.freeze_panes = 'E2'
    sheet.sheet_view.showGridLines = False
    configure_print(sheet,get_column_letter(len(headings)),'landscape')
