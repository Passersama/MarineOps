"""Check import title case, collection-only materials and automatic selection in a visible window."""
from pathlib import Path
import hashlib
import json
import os
import sys
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))


def main():
    from PySide6.QtCore import QSettings, Qt
    from PySide6.QtWidgets import QApplication, QPushButton
    from PySide6.QtTest import QTest
    from openpyxl import Workbook, load_workbook
    from cargo_track import __version__
    from cargo_track.appearance import configure_application
    from cargo_track.domain import Command
    from cargo_track.service import ApplicationService
    from cargo_track.workspace import WorkspaceWindow
    from cargo_track.personnel_ui import ImportDialog, CrewPlanBuilder
    from cargo_track.personnel_import import read_personnel_excel
    from cargo_track.berth_plan import export_plan, plan_text
    from verify_crew_final_runtime import wait

    output = ROOT / "artifacts" / f"personnel-defaults-{__version__}"
    data = output / str(time.time_ns())
    data.mkdir(parents=True)
    os.environ["CARGO_TRACK_DATA_DIR"] = str(data)
    QSettings.setDefaultFormat(QSettings.Format.IniFormat)
    QSettings.setPath(QSettings.Format.IniFormat, QSettings.Scope.UserScope, str(data / "preferences"))
    app = QApplication([])
    app.setOrganizationName("Personnel default validation")
    app.setApplicationName("Personnel default validation")
    configure_application(app)
    cargo_path = data / "cargo-track.sqlite3"
    cargo = ApplicationService(cargo_path)
    cargo.open()
    sequence = 0

    def command(operation, payload):
        nonlocal sequence
        sequence += 1
        return cargo.handle(Command(f"defaults-{sequence}", operation, payload)).result

    dfkt, inn, hb6, victor = [command("create_vessel", {"name": name})["id"]
        for name in ("DFKT2", "BGP Innovator", "HB6", "BGP Victor")]
    command("create_voyage", {"name": "Default selection validation"})
    for vessel in (dfkt, inn):
        command("add_stop", {"vessel_id": vessel})
    materials = [("220V-24v电源转换器", dfkt, hb6), ("GPS电缆", dfkt, victor),
                 ("OBS Hock", dfkt, inn), ("Hard Disk", hb6, inn), ("Return equipment", inn, dfkt)]
    for name, source, target in materials:
        command("create_transfer", {"description": name, "from_location_id": source,
                "to_location_id": target, "physical_id": "", "comment": "Keep UPPER Comment"})
    cargo.close()
    cargo_hash = hashlib.sha256(cargo_path.read_bytes()).hexdigest()
    book = Workbook()
    sheet = book.active
    sheet.append(["Name", "Position", "From", "To", "Date"])
    for name, position in (("ALEX MORGAN", "NAV QC"), ("sam cARTER", "GUNNER")):
        sheet.append([name, position, "xxPort", "BGP Innovator", "2026-10-06"])
    workbook = data / "mixed-case.xlsx"
    book.save(workbook)
    book.close()
    workbook_hash = hashlib.sha256(workbook.read_bytes()).hexdigest()
    window = WorkspaceWindow(cargo_path, data)
    builder = None
    try:
        window.show()
        window.resize(1366, 768)
        wait(app, lambda: window.snapshot is not None and not window._busy)
        QTest.mouseClick(window.personnel_button, Qt.MouseButton.LeftButton)
        wait(app, lambda: window.personnel_page is not None and window.personnel_page.source is not None)
        page = window.personnel_page
        dialog = ImportDialog(page, read_personnel_excel(workbook))
        dialog.show()
        app.processEvents()
        assert [(r["data"]["name"], r["data"]["position"]) for r in dialog.preview["items"]] == [
            ("Alex Morgan", "Nav Qc"), ("Sam Carter", "Gunner")]
        imported = page.service.commit_import(dialog.preview, "Case validation", "2026-10-06")
        dialog.reject()
        page.refresh_plans(imported["plan_id"])
        assert page.passenger_table.item(0, 1).text() == "Alex Morgan"
        assert page.passenger_table.item(0, 2).text() == "Nav Qc"
        window.grab().save(str(output / "crew-title-case.png"))
        builder = CrewPlanBuilder(page)
        builder.show()
        app.processEvents()
        rows = {t["name"]: (t, check, combo) for t, check, combo in
                zip(builder.tasks, builder.includes, builder.assignments)}
        for name in ("Alex Morgan", "Sam Carter", "220V-24v电源转换器", "GPS电缆", "OBS Hock"):
            assert rows[name][1].isChecked(), name
        for name in ("Hard Disk", "Return equipment"):
            assert not rows[name][1].isChecked(), name
        for name in ("220V-24v电源转换器", "GPS电缆"):
            assert rows[name][2].currentData()["drop"] is None
            assert "not called" in rows[name][2].currentText()
        builder.grab().save(str(output / "automatic-materials.png"))
        draft = builder.build()
        assert not draft["issues"]
        assert len(draft["stops"][0]["pick_materials"]) == 3
        assert not any(s["name"] in ("HB6", "BGP Victor") for s in draft["stops"])
        saved = page.service.save_berth(draft, published=True)
        export = output / "title-case-collection-only.xlsx"
        export_plan(saved, export)
        checked_book = load_workbook(export, read_only=True)
        try:
            printed = " ".join(str(c.value) for s in checked_book for row in s for c in row if c.value is not None)
            for text in ("Alex Morgan", "Nav Qc", "220V-24v电源转换器 to HB6", "GPS电缆 to BGP Victor"):
                assert text in printed, text
        finally:
            checked_book.close()
        # Manual exclusion survives route changes; newly feasible rows select themselves.
        rows["220V-24v电源转换器"][1].setChecked(False)
        builder.vessels.setCurrentIndex(builder.vessels.findData(dfkt))
        QTest.mouseClick(next(b for b in builder.findChildren(QPushButton) if b.text() == "Add stop"), Qt.MouseButton.LeftButton)
        assert rows["Return equipment"][1].isChecked()
        assert not rows["220V-24v电源转换器"][1].isChecked()
        builder.vessels.setCurrentIndex(builder.vessels.findData(hb6))
        builder.add_stop()
        builder.route.setCurrentRow(builder.route.count() - 1)
        while builder.route.currentRow() > 0:
            builder.move_stop(-1)
        assert rows["Hard Disk"][1].isChecked()
        edited = builder.build()
        assert not edited["issues"]
        builder.reject()
        reopened = CrewPlanBuilder(page, base_plan=edited)
        try:
            assert not next(c for t, c in zip(reopened.tasks, reopened.includes) if t["name"] == "220V-24v电源转换器").isChecked()
            inline = CrewPlanBuilder(page, base_plan=edited, inline=True)
            try:
                assert all(t["kind"] == "Passengers" for t in inline.tasks)
                assert "GPS电缆 to BGP Victor" in plan_text(inline.build())
            finally:
                inline.close()
        finally:
            reopened.close()
        assert cargo_hash == hashlib.sha256(cargo_path.read_bytes()).hexdigest()
        assert workbook_hash == hashlib.sha256(workbook.read_bytes()).hexdigest()
        report = {"version": __version__, "native_visible": window.isVisible(), "data_directory": str(data),
            "title_case_preview_and_list": True, "title_case_export": True, "collection_only_materials_exported": True,
            "feasible_rows_checked_automatically": True, "newly_feasible_rows_checked": True,
            "manual_exclusions_preserved": True, "reopened_exclusions_preserved": True,
            "inline_materials_readonly": True, "cargo_hash_unchanged": True, "workbook_hash_unchanged": True,
            "export": str(export)}
        (output / "native-visible.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(json.dumps(report, indent=2))
    finally:
        if builder is not None:
            builder.close()
        window.close()
        app.processEvents()


if __name__ == "__main__":
    main()
