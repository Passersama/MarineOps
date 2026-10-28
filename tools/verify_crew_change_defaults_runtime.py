"""Native workbook-to-plan acceptance with fictional sign-on/off rows only."""
from pathlib import Path
import hashlib,json,os,sys,time
from uuid import uuid4
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'src'))
OUT=ROOT/'output/crew-change-defaults-20261021';OUT.mkdir(parents=True,exist_ok=True)
data=OUT/('fixture-'+str(time.time_ns()));data.mkdir();os.environ['CARGO_TRACK_DATA_DIR']=str(data)
os.environ.pop('QT_QPA_PLATFORM',None)
from openpyxl import Workbook,load_workbook
from PySide6.QtCore import QSettings
from PySide6.QtWidgets import QApplication
from cargo_track.domain import Command
from cargo_track.service import ApplicationService
from cargo_track.personnel_service import PersonnelService
from cargo_track.personnel_import import read_personnel_excel
from cargo_track.personnel_ui import CrewPlanBuilder
from cargo_track.berth_plan import export_plan,plan_assignments
from cargo_track.appearance import configure_application
from cargo_track.workspace import WorkspaceWindow

cargo=ApplicationService(data/'cargo-track.sqlite3');cargo.open()
def command(name,payload):return cargo.handle(Command(str(uuid4()),name,payload)).result
fleet={name:command('create_vessel',{'name':name})['id'] for name in ('BGP Victor','HB6','BGP Innovator')}
command('create_voyage',{'name':'Fictional crew change 20261021'})
for name in ('BGP Innovator','BGP Victor','HB6','BGP Innovator'):command('add_stop',{'vessel_id':fleet[name]})
command('set_setting',{'key':'port_name','value':'Lagos Port'})
command('choose_stop',{'stop_id':cargo.workspace_snapshot()['stops'][0]['id']});cargo.close()
book=Workbook();sheet=book.active;sheet.title='Crew rotation'
sheet.append(['Sign-off','Position','From','To','Date','Sign-on','Position','From','To','Date'])
for i in range(8):
    vessel=('BGP Victor','HB6','BGP Innovator')[i%3]
    sheet.append([f'Off Crew {i+1}','Engineer',vessel,None,'2026-10-21',f'On Crew {i+1}','Technician',None,vessel,'2026-10-21'])
workbook=OUT/'fictional-rotation.xlsx';book.save(workbook);book.close();original_workbook=workbook.read_bytes()
from cargo_track.personnel_cargo import read_cargo
service=PersonnelService(data/'personnel-transfer.sqlite3');roster=service.create_plan('October rotation','2026-10-21')
preview=service.import_preview(read_personnel_excel(workbook),read_cargo(data/'cargo-track.sqlite3'),roster,crew_list=True)
assert len(preview['items'])==16 and all(not item['data']['issues'] for item in preview['items'])
service.commit_import(preview);service.close()
before=(data/'cargo-track.sqlite3').read_bytes()
QSettings.setDefaultFormat(QSettings.Format.IniFormat);QSettings.setPath(QSettings.Format.IniFormat,QSettings.Scope.UserScope,str(data/'preferences'))
app=QApplication([]);QApplication.setOrganizationName('Crew defaults validation');QApplication.setApplicationName('Isolated validation');configure_application(app)
w=WorkspaceWindow(data/'cargo-track.sqlite3',data)
def wait(predicate):
    until=time.monotonic()+20
    while time.monotonic()<until:
        app.processEvents()
        if predicate():return
        time.sleep(.01)
    raise RuntimeError('Native operation timed out')
def settle():
    for _ in range(12):app.processEvents();time.sleep(.01)
report={'fictional_data':True,'data_directory':str(data),'platform':app.platformName(),'scenarios':[]}
try:
    w.show();wait(lambda:w.snapshot is not None and not w._refresh_wait)
    w.open_personnel();wait(lambda:w.personnel_page and w.personnel_page.source and not w.personnel_page._jobs)
    p=w.personnel_page;p.plan_combo.setCurrentIndex(p.plan_combo.findData(roster))
    initial=CrewPlanBuilder(p);plan=initial.build();initial.deleteLater();assert not plan['issues'],plan['issues']
    p.draft=plan;p.show_plan(plan);editor=p.live_builder
    for controls in editor.endpoint_controls:
        assert all('Review' not in control.currentText() for control in controls)
    assert all(include.isChecked() for include in editor.includes)
    draft=editor.build();assert not draft['issues']
    port=draft['stops'][-1];assert port['name']=='Lagos Port' and len(port['drop_passengers'])==8
    assert sum(len(s['drop_passengers']) for s in draft['stops'][:-1])==8
    assert sum(len(s['pick_passengers']) for s in draft['stops'])==8
    for theme in ('light','dark'):
        for width,height in ((1366,768),(1000,700)):
            QSettings().setValue('appearance/theme',theme);configure_application(app);w.resize(width,height);settle()
            p.grab().save(str(OUT/f'personnel-{theme}-{width}.png'))
            report['scenarios'].append({'theme':theme,'width':width,'rows':len(editor.tasks),'review_assignments':0})
    saved=p.service.save_berth(editor.build(),True);export=OUT/'crew-change-plan.xlsx';export_plan(saved,export)
    reopened=CrewPlanBuilder(p,base_plan=saved)
    try:
        assert len(reopened.build()['stops'])==len(saved['stops'])
        assert plan_assignments(reopened.build())==plan_assignments(saved)
    finally:reopened.deleteLater()
    excel=load_workbook(export,read_only=True)
    try:assert 'Lagos Port' in ' '.join(str(c.value or '') for row in excel['Berthing Plan'] for c in row)
    finally:excel.close()
    assert before==(data/'cargo-track.sqlite3').read_bytes()
    assert original_workbook==workbook.read_bytes()
    report.update(result='PASS',sign_on=8,sign_off=8,port_returns=8,review_assignments=0,
                  saved_and_reopened=True,excel_export=True,cargo_unchanged=True,original_workbook_unchanged=True)
    (OUT/'native-checks.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report,indent=2))
finally:
    if w.personnel_page:w.personnel_page.view_state.dirty=False
    w.close();app.processEvents()
